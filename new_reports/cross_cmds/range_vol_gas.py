import pandas as pd
import numpy as np
import datetime as dt
import time
import sys
from ecm.cmds.config import root_path, output_path, html_path
sys.path.append(f"{root_path}\\autoreports\\reports\\cross_cmds")
import range_vol_am
import os
import math
import base64
from dateutil.relativedelta import relativedelta, FR
import ecm.cmds.bbg as bbg
import ecm.cmds.pyg as pyg
import ecm.cmds.talib as talib
import ecm.cmds.ticker as tk
import ecm.cmds.table as table
import ecm.cmds.utils as ut
import ecm.cmds.time_series as ts
import ecm.cmds.market_scan as ms
from ecm.cmds._email import send_email
from ecm.cmds.cdr import today, getworkingdays, CDR, now_ldn
from ecm.cmds.pyg import get_data
if sys.platform.startswith("win"):
    import excel2img
    from win32com.client import Dispatch
    from ecm.atom.wintask.scheduler import ECMWinTask
    from ecm.atom.wintask.utils import Days
from xlsx2html import xlsx2html
from copy import copy
import openpyxl
import plotly.graph_objects as go
from pandas.tseries.offsets import BDay

send_to = ["jmcphillips@elementcapital.com", "rzhao@elementcapital.com", "ltrindade", "caitcheson@elementcapital.com"]  # source35 recipient tail clipped
report_name = "Range-Vol and Divergence - Gas"
file_name = "range_vol_gas"  # without .py
file_path = f"{root_path}\\autoreports\\reports\\cross_cmds\\{file_name}.py"
cc_csv_folder = f"{output_path}\\csvs\\cross_cmds"
cc_json_folder = f"{output_path}\\json\\cross_cmds"
cc_pdf_folder = f"{output_path}\\pdf\\cross_cmds"
folder_path = f"{cc_csv_folder}\\market_scan\\"
folder_path_gas = f"{cc_csv_folder}\\market_scan\\gas\\"


def _photo_gap(source_line, visible_prefix):
    raise NotImplementedError(f"Unrecovered range_vol_gas source line {source_line}: {visible_prefix}")


def add_schedule():
    sys.path.append(f"{root_path}\\autoreports\\schedule")
    from tasks import WinTaskPoC
    win_task = ECMWinTask(
        **_photo_gap(55, 'unphotographed schedule arguments55–58'),
        task_name=report_name,
        task=WinTaskPoC(name=f'"{report_name}"', filepath=f'"{file_path}"'),
        background_task=True,
        python_excecutable=r"C:\local\python\miniconda\envs\ecm_cmds\pythonw.exe",
    )
    win_task.create_task()


name_dict = {
    'NGA Comdty': ('US_GAS', 'TICKER', 'FLAT', 1, 1, False),
    'NG2 Comdty': ('US_GAS', 'CONTRACT', 'FLAT', 1, 1, False),
    'NG3 Comdty': ('US_GAS', 'CONTRACT', 'FLAT', 1, 1, False),
    'NG4 Comdty': ('US_GAS', 'CONTRACT', 'FLAT', 1, 1, False),
    'NG5 Comdty': ('US_GAS', 'CONTRACT', 'FLAT', 0, 1, False),
    'NG6 Comdty': ('US_GAS', 'CONTRACT', 'FLAT', 0, 1, False),
    'S:NGNG 2-6 Comdty': ('US_GAS', 'CALSPRD', 'SPRD', 1, 1, False),
    'NG 1st Spread': ('US_GAS', 'CALCULATE', 'SPRD', 1, 1, False),
    'NG 2nd Spread': ('US_GAS', 'CALCULATE', 'SPRD', 1, 1, False),
    'NG 3rd Spread': ('US_GAS', 'CALCULATE', 'SPRD', 1, 1, False),
    'NGMAR1 Comdty': ('US_GAS', 'FUT_CUR_GEN_TICKER', 'FLAT', 0, 1, False),
    'NGOCT1 Comdty': ('US_GAS', 'FUT_CUR_GEN_TICKER', 'FLAT', 0, 1, False),
    'NG Mar-Apr Spread': ('US_GAS', 'CALCULATE', 'SPRD', 1, 1, False),
    'NG Oct-Jan Spread': ('US_GAS', 'CALCULATE', 'SPRD', 1, 1, False),
    'NGH4NGJ4 Comdty': ('US_GAS', 'COPY', 'SPRD', 1, 1, False),
    'NGV4NGF5 Comdty': ('US_GAS', 'COPY', 'SPRD', 0, 1, False),
    'NG 1M 100 VOL BVOL Comdty': ('US_GAS', 'COPY', 'SPRD', 1, 0, False),
    'NG 3M 100 VOL BVOL Comdty': ('US_GAS', 'COPY', 'SPRD', 1, 0, False),
    'NG 6M 100 VOL BVOL Comdty': ('US_GAS', 'COPY', 'SPRD', 1, 0, False),
    '.NG1M Index': ('US_GAS', 'COPY', 'SPRD', 1, 0, False),
    '.NG3M Index': ('US_GAS', 'COPY', 'SPRD', 0, 0, False),
    '.NG6M Index': ('US_GAS', 'COPY', 'SPRD', 0, 0, False),
    'TZTA Comdty': ('EU_GAS', 'TICKER', 'FLAT', 1, 1, False),
    'TZT2 Comdty': ('EU_GAS', 'CONTRACT', 'FLAT', 1, 1, False),
    'TZT3 Comdty': ('EU_GAS', 'CONTRACT', 'FLAT', 0, 1, False),
    'TZT 1st spread': ('EU_GAS', 'CALCULATE', 'SPRD', 1, 1, False),
    'TZT 2nd spread': ('EU_GAS', 'CALCULATE', 'SPRD', 1, 1, False),
    'TZTH4J4 Comdty': ('EU_GAS', 'COPY', 'SPRD', 1, 1, False),
    'QZTA Comdty': ('EU_GAS', 'TICKER', 'FLAT', 1, 1, False),
    'QZT2 Comdty': ('EU_GAS', 'CONTRACT', 'FLAT', 1, 1, False),
    'QZT3 Comdty': ('EU_GAS', 'CONTRACT', 'FLAT', 0, 1, False),
    'QZT 1st spread': ('EU_GAS', 'CALCULATE', 'SPRD', 1, 1, False),
    'QZT 2nd spread': ('EU_GAS', 'CALCULATE', 'SPRD', 1, 1, False),
    'QQTA Comdty': ('EU_GAS', 'TICKER', 'FLAT', 1, 1, False),
    'QQT2 Comdty': ('EU_GAS', 'CONTRACT', 'FLAT', 0, 1, False),
    'QQT3 Comdty': ('EU_GAS', 'CONTRACT', 'FLAT', 0, 1, False),
    'QQT 1st spread': ('EU_GAS', 'CALCULATE', 'SPRD', 1, 1, False),
    'QQT 2nd spread': ('EU_GAS', 'CALCULATE', 'SPRD', 1, 1, False),
    'QTTA Comdty': ('EU_GAS', 'TICKER', 'FLAT', 1, 1, False),
    'QTT2 Comdty': ('EU_GAS', 'CONTRACT', 'FLAT', 1, 1, False),
    'QTT3 Comdty': ('EU_GAS', 'CONTRACT', 'FLAT', 1, 1, False),
    'QTT4 Comdty': ('EU_GAS', 'CONTRACT', 'FLAT', 1, 1, False),
    'QTT 1st spread': ('EU_GAS', 'CALCULATE', 'SPRD', 1, 1, False),
    'FJS 1M 100 VOL BVOL Comdty': ('EU_GAS', 'COPY', 'SPRD', 1, 0, False),
    'FJS 3M 100 VOL BVOL Comdty': ('EU_GAS', 'COPY', 'SPRD', 1, 0, False),
    'FJS 6M 100 VOL BVOL Comdty': ('EU_GAS', 'COPY', 'SPRD', 1, 0, False),
    '.FJS1MRR Index': ('EU_GAS', 'COPY', 'SPRD', 1, 0, False),
    'GERW1WYY Comdty': ('EU_POWER', 'WEEKLY', 'FLAT', 1, 1, False),
    'GERW2WYY Comdty': ('EU_POWER', 'WEEKLY', 'FLAT', 1, 1, False),
    'DETA Comdty': ('EU_POWER', 'TICKER', 'FLAT', 1, 1, False),
    'DET2 Comdty': ('EU_POWER', 'CONTRACT', 'FLAT', 0, 1, False),
    'DET 1st spread': ('EU_POWER', 'CALCULATE', 'SPRD', 0, 1, False),
    'JXTA Comdty': ('EU_POWER', 'TICKER', 'FLAT', 1, 1, False),
    'JXT2 Comdty': ('EU_POWER', 'CONTRACT', 'FLAT', 0, 1, False),
    'JXT 1st spread': ('EU_POWER', 'CALCULATE', 'SPRD', 0, 1, False),
    'JXYA Comdty': ('EU_POWER', 'TICKER', 'FLAT', 1, 1, False),
    'JXY2 Comdty': ('EU_POWER', 'CONTRACT', 'FLAT', 1, 1, False),
    'JXY3 Comdty': ('EU_POWER', 'CONTRACT', 'FLAT', 1, 1, False),
    'JXY 1st spread': ('EU_POWER', 'CALCULATE', 'SPRD', 0, 1, False),
    'DET 1M 100 VOL BVOL Comdty': ('EU_POWER', 'COPY', 'SPRD', 1, 0, False),
    'DET 3M 100 VOL BVOL Comdty': ('EU_POWER', 'COPY', 'SPRD', 1, 0, False),
    'FACA Comdty': ('EU_POWER', 'TICKER', 'FLAT', 1, 1, False),
    'FABA Comdty': ('EU_POWER', 'TICKER', 'FLAT', 1, 1, False),
    'FAAA Comdty': ('EU_POWER', 'TICKER', 'FLAT', 1, 1, False),
    'FAA2 Comdty': ('EU_POWER', 'TICKER', 'FLAT', 1, 1, False),
    'FAA3 Comdty': ('EU_POWER', 'TICKER', 'FLAT', 1, 1, False),
    'FAC 1M 100 VOL BVOL Comdty': ('EU_POWER', 'COPY', 'SPRD', 1, 0, False),
    'FAC 3M 100 VOL BVOL Comdty': ('EU_POWER', 'COPY', 'SPRD', 1, 0, False),
    'ELGBY 22 OECM Index': ('EU_POWER', 'ROLLYEAR', 'FLAT', 1, 0, False),
    'FSNGY 23 Index': ('US_GAS', 'ROLLYEAR', 'FLAT', 1, 0, False),
    'FSNGY 24 Index': ('US_GAS', 'ROLLYEAR+1', 'FLAT', 1, 0, False),
    'FSNGY 25 Index': ('US_GAS', 'ROLLYEAR+2', 'FLAT', 1, 0, False),
    'FSNGY 26 Index': ('US_GAS', 'ROLLYEAR+3', 'FLAT', 1, 0, False),
    'FSNGS SU23 Index': ('US_GAS', 'ROLLYEAR', 'FLAT', 1, 0, False),
    'FSNGS SU24 Index': ('US_GAS', 'ROLLYEAR+1', 'FLAT', 1, 0, False),
    'FSNGS SU25 Index': ('US_GAS', 'ROLLYEAR+2', 'FLAT', 1, 0, False),
    'FSNGS WI2223 Index': ('US_GAS', 'ROLLYEAR', 'FLAT', 1, 0, False),
    'FSNGS WI2324 Index': ('US_GAS', 'ROLLYEAR+1', 'FLAT', 1, 0, False),
    'FSNGS WI2425 Index': ('US_GAS', 'ROLLYEAR+2', 'FLAT', 1, 0, False),
    'MOA Comdty': ('EU_POWER', 'TICKER', 'FLAT', 1, 1, False),
    'CTIA Comdty': ('US_GAS', 'TICKER', 'FLAT', 1, 1, False),
    'XAA Comdty': ('COAL', 'TICKER', 'FLAT', 1, 1, False),
    'XEA Comdty': ('COAL', 'TICKER', 'FLAT', 1, 1, False),
    'TMA Comdty': ('COAL', 'TICKER', 'FLAT', 1, 1, False),
    'XNG Index': ('NG_STOCKS', 'COPY', 'FLAT', 1, 0, False),
    'EQNR NO Equity': ('NG_STOCKS', 'COPY', 'FLAT', 1, 1, False),
    'GAZP RM Equity': ('NG_STOCKS', 'COPY', 'FLAT', 1, 1, False),
    'LNG US Equity': ('NG_STOCKS', 'COPY', 'FLAT', 1, 1, False),
    '.DRY_E&P Index': ('NG_STOCKS', 'COPY', 'FLAT', 1, 1, False),
    'FNA Comdty': ('EU_GAS', 'TICKER', 'FLAT', 1, 1, True),
    'FN2 Comdty': ('EU_GAS', 'CONTRACT', 'FLAT', 1, 1, True),
    'FN3 Comdty': ('EU_GAS', 'CONTRACT', 'FLAT', 0, 1, True),
    'FN 1st spread': ('EU_GAS', 'CALCULATE', 'SPRD', 1, 1, True),
    'FN 2nd spread': ('EU_GAS', 'CALCULATE', 'SPRD', 1, 1, True),
    'FNMAR1 Comdty': ('EU_GAS', 'FUT_CUR_GEN_TICKER', 'FLAT', 0, 1, True),
    'FN Mar-Apr Spread': ('EU_GAS', 'CALCULATE', 'SPRD', 1, 1, True),
    'QRA Comdty': ('EU_GAS', 'TICKER', 'FLAT', 1, 1, True),
    'QR2 Comdty': ('EU_GAS', 'CONTRACT', 'FLAT', 1, 1, True),
    'QR3 Comdty': ('EU_GAS', 'CONTRACT', 'FLAT', 0, 1, True),
    'QR 1st spread': ('EU_GAS', 'CALCULATE', 'SPRD', 1, 1, True),
    'QR 2nd spread': ('EU_GAS', 'CALCULATE', 'SPRD', 1, 1, True),
    'SAA Comdty': ('EU_GAS', 'TICKER', 'FLAT', 1, 1, True),
    'SA2 Comdty': ('EU_GAS', 'CONTRACT', 'FLAT', 0, 1, True),
    'SA 1st spread': ('EU_GAS', 'CALCULATE', 'SPRD', 1, 1, True),
    'EEDATTF1 Index': ('EU_GAS', 'COPY', 'FLAT', 1, 1, True),
    'NBPGDAHD Index': ('EU_GAS', 'COPY', 'FLAT', 1, 1, True),
    'EGTHDAHD Index': ('EU_GAS', 'COPY', 'FLAT', 1, 1, True),
    'NCG1M Comdty': ('EU_GAS', 'TICKER', 'FLAT', 1, 1, True),
    'NCG1Q Comdty': ('EU_GAS', 'TICKER', 'FLAT', 1, 1, True),
    'NCG1S Comdty': ('EU_GAS', 'TICKER', 'FLAT', 1, 1, True),
    'NCG2S Comdty': ('EU_GAS', 'TICKER', 'FLAT', 1, 1, True),
    'NCG1Y Comdty': ('EU_GAS', 'TICKER', 'FLAT', 1, 1, True),
    'NCG2Y Comdty': ('EU_GAS', 'TICKER', 'FLAT', 1, 1, True),
    'COA Comdty': ('OIL', 'TICKER', 'FLAT', 1, 1, False),
    'CODEC1 Comdty': ('OIL', 'FUT_CUR_GEN_TICKER', 'FLAT', 1, 1, False),
    'CODEC2 Comdty': ('OIL', 'FUT_CUR_GEN_TICKER', 'FLAT', 1, 1, False),
    'CLA Comdty': ('OIL', 'TICKER', 'FLAT', 1, 1, False),
    'CLDEC1 Comdty': ('OIL', 'FUT_CUR_GEN_TICKER', 'FLAT', 1, 1, False),
    'CLDEC2 Comdty': ('OIL', 'FUT_CUR_GEN_TICKER', 'FLAT', 1, 1, False),
}
month_ticker_dict = {'JAN': 'F', 'FEB': 'G', 'MAR': 'H', 'APR': 'J', 'MAY': 'K', 'JUN': 'M', 'JUL': 'N', 'AUG': 'Q', 'SEP': 'U', 'OCT': 'V', 'NOV': 'X', 'DEC': 'Z'}
month_ticker_dict1 = {1: 'F', 2: 'G', 3: 'H', 4: 'J', 5: 'K', 6: 'M', 7: 'N', 8: 'Q', 9: 'U', 10: 'V', 11: 'X', 12: 'Z'}


def sr_3d_cob(intraday_price_chg, sdate_3d_cob, edate_3d_cob, instr, chg_type):
    if edate_3d_cob > intraday_price_chg.index[-1] or edate_3d_cob not in intraday_price_chg.index:
        three_d_chg_cob = intraday_price_chg[(intraday_price_chg.index >= sdate_3d_cob) & (intraday_price_chg.index <= edate_3d_cob)]
    else:
        three_d_chg_cob = intraday_price_chg.loc[sdate_3d_cob:edate_3d_cob]
    if len(three_d_chg_cob) > 10:
        if instr == "FI" and chg_type == "SPRD":
            weight = 1
            three_d_chg_cob = three_d_chg_cob * -1
            three_d_chg_cob = three_d_chg_cob.drop(three_d_chg_cob.index[0])
            chg_3d_cob = three_d_chg_cob.mean()
            stdev_3d_cob = three_d_chg_cob.std()
        elif instr != "FI" and chg_type == "SPRD":
            weight = .3
            three_d_chg_cob = three_d_chg_cob
            three_d_chg_cob = three_d_chg_cob.drop(three_d_chg_cob.index[0])
            chg_3d_cob = three_d_chg_cob.mean()
            stdev_3d_cob = three_d_chg_cob.std()
        else:
            weight = 1
            three_d_chg_cob = three_d_chg_cob.drop(three_d_chg_cob.index[0])
            chg_3d_cob = three_d_chg_cob.mean()
            stdev_3d_cob = three_d_chg_cob.std()
        try:
            sharpe_3d_cob = (chg_3d_cob / stdev_3d_cob * np.sqrt(len(three_d_chg_cob))) * weight
        except:
            sharpe_3d_cob = np.nan
    else:
        sharpe_3d_cob = np.nan
    return sharpe_3d_cob


def ticker_split(ticker='NGV21 Comdty'):
    '''Return ticker for instrument, month symbol, year, and asset class'''
    ticker_parts = ticker.split(' ')
    if len(ticker_parts) > 2:
        instrument = ticker_parts[0] + ' '
        month_symbol = ticker_parts[1][0]
        year = ticker_parts[1][1:]
        asset_class = ticker_parts[-1]
    else:
        i = -1
        tick = ticker.rpartition(' ')[0]
        while abs(i) < len(tick):
            if not tick[i].isdigit():
                month_symbol = tick[i]
                year = tick[i + 1:]
                instrument = tick[:i]
                break
            else:
                i -= 1
        asset_class = ticker_parts[-1]
    return instrument, month_symbol, year, asset_class


def ticker_split_gen(ticker='C 2 Comdty'):
    '''Return ticker for instrument, number of contract, and asset class'''
    ticker_parts = ticker.split(' ')
    if len(ticker_parts) > 2:
        instrument = ticker_parts[0] + ' '
        num_contract = ticker_parts[1]
        asset_class = ticker_parts[-1]
    else:
        if ticker_parts[0][-1] == 'A':
            instrument = ticker_parts[:-1]
            num_contract = 'A'
        elif ticker_parts[0][-1].isdigit():
            i = -1
            tick = ticker.rpartition(' ')[0]
            while abs(i) < len(tick):
                if not tick[i].isdigit():
                    instrument = tick[:i + 1]
                    num_contract = tick[i + 1:]
                    break
                else:
                    i -= 1
        else:
            i = -2
            tick = ticker.rpartition(' ')[0]
            while abs(i) < len(tick):
                if not tick[i].isdigit():
                    instrument = tick[:i + 1]
                    num_contract = tick[i + 1:]
                    break
                else:
                    i -= 1
        asset_class = ticker_parts[-1]
    return instrument, num_contract, asset_class


def get_bbg_ticker(name, ticker_type, **kwargs):
    cur_ticker = kwargs.get('cur_ticker', None)
    next = kwargs.get('next', None)
    name1 = kwargs.get('name1', None)
    name2 = kwargs.get('name2', None)
    if ticker_type in ['TICKER', 'FUT_CUR_GEN_TICKER']:
        ticker = bbg.bref(name, [ticker_type])
        if name.split(' ')[-2] == 'US':
            ticker = ticker[ticker_type][0] + ' US Equity'
        else:
            ticker = ticker[ticker_type][0] + ' ' + name.split(' ')[-1]
            if name == 'CTIA Comdty':
                ticker = f'{ticker[:3]}Z{ticker[-9:]}'
    elif ticker_type == 'COPY':
        ticker = name
    elif ticker_type == 'CONTRACT':
        month_symbol_list = bbg.bref(cur_ticker, ['FUT_GEN_MONTH']).values[0][0]
        month_symbol_list = [x for x in month_symbol_list]
        instrument, month_symbol, yr, asset_class = ticker_split(cur_ticker)
        next_ticker_loc = month_symbol_list.index(month_symbol) + next - 1
        if next_ticker_loc > len(month_symbol_list) - 1:
            next_ticker_symbol = month_symbol_list[next_ticker_loc % len(month_symbol_list)]
            yr = str(int(yr) + math.floor(next_ticker_loc / len(month_symbol_list)))
        else:
            next_ticker_symbol = month_symbol_list[next_ticker_loc]
        ticker = name.rpartition(' ')[0][:-1] + next_ticker_symbol + yr + ' ' + name.rpartition(' ')[-1]
    elif ticker_type == 'CALCULATE':
        instrument, month_symbol, yr, asset_class = ticker_split(name1)
        if instrument[-1] == ' ':
            if len(yr) == 1:
                ticker = name1.split(' ')[0] + '_' + name1.split(' ')[1] + name2.split(' ')[0] + '_' + name2.split(' ')[1] + ' ' + name1.split(' ')[-1]
            elif len(yr) == 2:
                ticker = name1.split(' ')[0] + '_' + name1.split(' ')[1][0] + name1.split(' ')[1][-1] + name2.split(' ')[0] + '_' + name2.split(' ')[1][0] + name2.split(' ')[1][-1] + ' ' + name1.split(' ')[-1]
        elif len(instrument) > 2:
            if len(yr) == 1:
                ticker = name1.split(' ')[0] + name2.split(' ')[0][-2:] + ' ' + name2.split(' ')[1]
            elif len(yr) == 2:
                ticker = name1.split(' ')[0][:-2] + name1.split(' ')[0][-1] + name2.split(' ')[0][-3] + name2.split(' ')[0][-1] + ' ' + name1.split(' ')[-1]
        else:
            if len(yr) == 1:
                ticker = name1.split(' ')[0] + name2.split(' ')[0] + ' ' + name2.split(' ')[1]
            elif len(yr) == 2:
                ticker = name1.split(' ')[0][:-2] + name1.split(' ')[0][-1] + name2.split(' ')[0][:-2] + name2.split(' ')[0][-1] + ' ' + name1.split(' ')[-1]
    elif ticker_type[:8] == 'ROLLYEAR':
        if '+' in ticker_type:
            year_str = str(dt.datetime.now().year + 1 + int(ticker_type[-1]))[-2:]
            year_str0 = str(dt.datetime.now().year + int(ticker_type[-1]))[-2:]
        else:
            year_str = str(dt.datetime.now().year + 1)[-2:]
            year_str0 = str(dt.datetime.now().year)[-2:]
        ticker_parts = name.split(' ')
        if ticker_parts[1].isdecimal():
            if len(ticker_parts[1]) > 2:
                ticker_parts = [x if idx != 1 else year_str0 + year_str for idx, x in enumerate(ticker_parts)]
            else:
                ticker_parts = [x if idx != 1 else year_str for idx, x in enumerate(ticker_parts)]
        else:
            if len(ticker_parts[1]) > 4:
                ticker_parts = [x if idx != 1 else x[:2] + year_str0 + year_str for idx, x in enumerate(ticker_parts)]
            else:
                ticker_parts = [x if idx != 1 else x[:2] + year_str for idx, x in enumerate(ticker_parts)]
        ticker = ' '.join(ticker_parts)
    elif ticker_type == "WEEKLY":
        add_week = int(name[4])
        week_num = "{0}{1:02}".format(*(today() + dt.timedelta(7 * add_week)).isocalendar())
        ticker = name[:3] + week_num[-2:] + name[5] + week_num[2:4] + " " + name.split(" ")[-1]
    return ticker


def get_tickers(name_dict=name_dict):
    ticker_dict = ms.get_ticker_dict(env="prod", force_update=False, gas_mode=True)
    return ticker_dict


def datetime_range(start, end, delta):
    current = start
    while current <= end:
        yield current
        current += delta


def get_basic_info(ticker_dict, use_bbg=True):
    output_list = []
    sdate = today() - dt.timedelta(days=35)
    edate = dt.datetime.now()
    next_friday = today() + relativedelta(weekday=FR(1))
    next_monthend = today() + relativedelta(day=31)
    fom_path = ut.convert_path_to_linux(_photo_gap(398, r'\\elementcapital.corp\ecns01\PM\Michel Kikano\Market Scan\F...'))
    fom_comments = pd.read_csv(fom_path)
    fom_comments.set_index('Unnamed: 0', inplace=True)
    fom_comments_ = fom_comments.copy()
    fom_comments_['days'] = 0
    for idx, row in fom_comments_.iterrows():
        if isinstance(row['FOM Dates'], str):
            if len(row['FOM Dates']) == 10:
                try:
                    fom_d0 = dt.datetime.strptime(row['FOM Dates'][:10], '%Y-%m-%d')
                except:
                    fom_d0 = dt.datetime.strptime(row['FOM Dates'][:10], '%m/%d/%Y')
                fom_comments_.loc[idx, 'days'] = int(getworkingdays(fom_d0, today()))
                fom_comments_.loc[idx, 'FOM Dates'] = str(dt.datetime.strftime(fom_d0, '%d/%b'))
            else:
                try:
                    fom_d0 = dt.datetime.strptime(row['FOM Dates'][:10], '%Y-%m-%d')
                except:
                    fom_d0 = dt.datetime.strptime(row['FOM Dates'][:10], '%m/%d/%Y')
                try:
                    fom_d1 = dt.datetime.strptime(row['FOM Dates'][-10:], '%Y-%m-%d')
                except:
                    fom_d1 = dt.datetime.strptime(row['FOM Dates'][-10:], '%m/%d/%Y')
                fom_comments_.loc[idx, 'days'] = int(getworkingdays(fom_d1, today()))
                fom_comments_.loc[idx, 'FOM Dates'] = dt.datetime.strftime(fom_d0, '%d/%b') + ' - ' + dt.datetime.strftime(fom_d1, '%d/%b')
        elif isinstance(row['FOM Dates'], dt.datetime):
            fom_comments_.loc[idx, 'days'] = int(getworkingdays(row['FOM Dates'], today()))
            fom_comments_.loc[idx, 'FOM Dates'] = str(dt.datetime.strftime(row['FOM Dates'], '%d/%b'))
    fom_comments_.to_csv(ut.convert_path_to_linux(f'{folder_path}\\fom_comments_gas.csv'))
    alert_email_all = []
    sharpe_ratio_list = []
    week_dates = pd.bdate_range(today() - dt.timedelta(days=7), today())
    for i in range(len(week_dates) - 2, -1, -1):
        if week_dates[i].weekday() > week_dates[i + 1].weekday():
            week_sdate = week_dates[i + 1]
    ticker_dict_rec = {x[1]: y for x, y in ticker_dict.items()}
    for (asset_class, key), val in ticker_dict.items():
        print(val[0])
        use_bbg = val[5]
        if val[3]:
            instr = val[1]
            ticker = val[0]
            chg_type = val[2]
            try:
                info = bbg.bref(ticker, ['FUT_LAST_TRADE_DT', 'PX_CLOSE_1D'])
                exp_date = info['FUT_LAST_TRADE_DT'].values[0]
                prev_close = info['PX_CLOSE_1D'].values[0]
            except:
                exp_date = np.nan
                prev_close = np.nan
            if isinstance(exp_date, dt.date):
                days_before_exp = (exp_date - today().date()).days
            else:
                days_before_exp = np.nan
            if ticker.split(' ')[-2] == 'BVOL':
                print("Skiping BVOL Ticker dont have permissions")
                continue
            elif ticker.split(' ')[0] in ['.CL1M', '.CO1M', '.NG1M', '.S1M', '.C1M', '.W1M', '.SB_JUL', *_photo_gap(470, '.F...; volatility index list tail')]:
                if ticker == '.SB_JUL Index':
                    ticker_vol = 'SBJUL1 Comdty'
                    tenor_vol = '1M'
                elif ticker.split(' ')[0] == '.FJS1MRR':
                    ticker_vol = ticker_dict_rec['TZT2 Comdty'][0]
                    tenor_vol = '1M'
                elif ticker.split(' ')[0][-2:] == 'RR':
                    ticker_vol = ticker_dict_rec[ticker.split(' ')[0][1:-4] + '2 Comdty'][0]
                    tenor_vol = ticker.split(' ')[0][-4:-2]
                else:
                    tenor_vol = ticker.split(' ')[0][-2:]
                    if len(ticker.split(' ')[0][1:-2]) < 2:
                        ticker_vol = ticker_dict_rec[ticker.split(' ')[0][1:-2] + ' 2 Comdty'][0]
                    else:
                        ticker_vol = ticker_dict_rec[ticker.split(' ')[0][1:-2] + '2 Comdty'][0]
                if os.path.exists(folder_path + 'price_daily\\{:s}.csv'.format(ticker_vol + '_PUT_CALL_{:s}'.format(tenor_vol))):
                    try:
                        daily_price_put_history = ts.read_csv(folder_path + 'price_daily\\{:s}.csv'.format(ticker_vol + '_PUT_CALL_{:s}'.format(tenor_vol)), index_name='date')
                        exist_sdate = daily_price_put_history.index[-2]
                        if use_bbg:
                            daily_price_put_new = bbg.bdh(ticker_vol, ['{:s}_PUT_IMP_VOL_25DELTA_DFLT'.format(tenor_vol), '{:s}_CALL_IMP_VOL_25DELTA_DFLT'.format(tenor_vol)], sdate=exist_sdate, edate=edate)
                            daily_price_put = pd.concat([daily_price_put_history.iloc[:-2, :], daily_price_put_new], axis=0)
                        else:
                            daily_price_put = daily_price_put_history
                    except:
                        daily_price_put = bbg.bdh(ticker_vol, ['{:s}_PUT_IMP_VOL_25DELTA_DFLT'.format(tenor_vol), '{:s}_CALL_IMP_VOL_25DELTA_DFLT'.format(tenor_vol)], sdate=sdate - dt.timedelta(days=182), edate=edate)
                else:
                    daily_price_put = bbg.bdh(ticker_vol, ['{:s}_PUT_IMP_VOL_25DELTA_DFLT'.format(tenor_vol), '{:s}_CALL_IMP_VOL_25DELTA_DFLT'.format(tenor_vol)], sdate=sdate - dt.timedelta(days=182), edate=edate)
                daily_price_put.to_csv(ut.convert_path_to_linux(folder_path + 'price_daily\\{:s}.csv'.format(ticker_vol + '_PUT_CALL_{:s}'.format(tenor_vol))))
                if len(daily_price_put) > 0:
                    daily_price = daily_price_put.iloc[:, 0] - daily_price_put.iloc[:, 1]
                    daily_price = daily_price.to_frame('PX_LAST')
                    weekly_price_put = bbg.bdh(ticker_vol, ['{:s}_PUT_IMP_VOL_25DELTA_DFLT'.format(tenor_vol), '{:s}_CALL_IMP_VOL_25DELTA_DFLT'.format(tenor_vol)],
                                               sdate=sdate, edate=next_friday, elms=[('periodicitySelection', 'WEEKLY')])
                    weekly_price = weekly_price_put.iloc[:, 0] - weekly_price_put.iloc[:, 1]
                    weekly_price = weekly_price.to_frame('PX_LAST')
                    monthly_price_put = bbg.bdh(ticker_vol, ['{:s}_PUT_IMP_VOL_25DELTA_DFLT'.format(tenor_vol), '{:s}_CALL_IMP_VOL_25DELTA_DFLT'.format(tenor_vol)],
                                                sdate=sdate - dt.timedelta(days=35), edate=next_monthend, elms=[('periodicitySelection', 'MONTHLY')])
                    monthly_price = monthly_price_put.iloc[:, 0] - monthly_price_put.iloc[:, 1]
                    monthly_price = monthly_price.to_frame('PX_LAST')
                    intraday_price = pd.DataFrame()
                else:
                    daily_price = pd.DataFrame()
            else:
                if os.path.exists(folder_path + 'price_daily\\{:s}.csv'.format(ticker)):
                    try:
                        daily_price_history = ts.read_csv(folder_path + 'price_daily\\{:s}.csv'.format(ticker), index_name='date')
                        exist_sdate = daily_price_history.index[-2]
                    except:
                        exist_sdate = sdate - dt.timedelta(days=182)
                    if use_bbg:
                        daily_price_new = bbg.bdh(ticker, ['PX_OPEN', 'PX_HIGH', 'PX_LOW', 'PX_LAST', 'PX_SETTLE', 'FUT_PX', 'VOLUME', '3MTH_IMPVOL_100.0%MNY_DF'], sdate=exist_sdate, edate=edate)
                        daily_price = pd.concat([daily_price_history.iloc[:-2, :], daily_price_new], axis=0)
                    else:
                        daily_price = daily_price_history
                else:
                    try:
                        daily_price = bbg.bdh(ticker, ['PX_OPEN', 'PX_HIGH', 'PX_LOW', 'PX_LAST', 'PX_SETTLE', 'FUT_PX', 'VOLUME', '3MTH_IMPVOL_100.0%MNY_DF'], sdate=sdate - dt.timedelta(days=182), edate=edate)
                    except:
                        print(f"Could not get data for {ticker}")
                        daily_price = pd.DataFrame()
                try:
                    daily_price.to_csv(ut.convert_path_to_linux(folder_path + 'price_daily\\{:s}.csv'.format(ticker)))
                except:
                    pass
                if ticker.split(' ')[0] in ['EURUSD', 'USDJPY', 'AUDUSD', 'USDCAD', 'USDCNH', 'USDINR', 'USDBRL', 'USDRUB']:
                    ticker_adj = ticker.split(' ')[0] + 'V3M' + ticker.split(' ')[1]
                    if os.path.exists(folder_path + 'price_daily\\{:s}.csv'.format(ticker_adj)):
                        try:
                            daily_vol_history = ts.read_csv(folder_path + 'price_daily\\{:s}.csv'.format(ticker_adj), index_name='date')
                            exist_sdate = daily_vol_history.index[-2]
                            daily_vol_new = bbg.bdh(ticker_adj, ['PX_LAST'], sdate=exist_sdate, edate=edate)
                            daily_vol = pd.concat([daily_vol_history.iloc[:-2, :], daily_vol_new], axis=0)
                        except:
                            daily_vol = bbg.bdh(ticker_adj, ['PX_LAST'], sdate=sdate - dt.timedelta(days=182), edate=edate)
                    else:
                        daily_vol = bbg.bdh(ticker_adj, ['PX_LAST'], sdate=sdate - dt.timedelta(days=182), edate=edate)
                    daily_vol.to_csv(ut.convert_path_to_linux(folder_path + 'price_daily\\{:s}.csv'.format(ticker_adj)))
                    daily_price['3MTH_IMPVOL_100.0%MNY_DF'] = daily_vol['PX_LAST']
                try:
                    weekly_price = bbg.bdh(ticker, ['PX_OPEN', 'PX_HIGH', 'PX_LOW', 'PX_LAST', 'PX_SETTLE', 'FUT_PX', 'VOLUME'],
                                           sdate=sdate, edate=next_friday, elms=[('periodicitySelection', 'WEEKLY')])
                except:
                    print(f'Could not get weekly price for {ticker}')
                try:
                    monthly_price = bbg.bdh(ticker, ['PX_OPEN', 'PX_HIGH', 'PX_LOW', 'PX_LAST', 'PX_SETTLE', 'FUT_PX', 'VOLUME'],
                                            sdate=sdate - dt.timedelta(days=35), edate=next_monthend, elms=[('periodicitySelection', 'MONTHLY')])
                except:
                    print(f'Could not get monthly price for {ticker}')
                if os.path.exists(folder_path + 'price_intraday\\{:s}.csv'.format(ticker)):
                    try:
                        intraday_price_history = ts.read_csv(folder_path + 'price_intraday\\{:s}.csv'.format(ticker), index_name='time')
                        exist_sdate = intraday_price_history.index[-2]
                        intraday_price_new = bbg.bdib(ticker, sdate=exist_sdate, edate=edate, interval=10)
                        intraday_price = pd.concat([intraday_price_history.iloc[:-2, :], intraday_price_new], axis=0)
                    except:
                        try:
                            intraday_price = bbg.bdib(ticker, sdate=sdate, edate=edate, interval=10)
                        except:
                            intraday_price = pd.DataFrame()
                else:
                    try:
                        intraday_price = bbg.bdib(ticker, sdate=sdate, edate=edate, interval=10)
                    except:
                        print(f"Could not load intraday price for {ticker}")
                try:
                    intraday_price.to_csv(ut.convert_path_to_linux(folder_path + 'price_intraday\\{:s}.csv'.format(ticker)))
                except:
                    pass
            if len(daily_price) > 10 and 'PX_LAST' in daily_price.columns:
                closep = daily_price['PX_LAST']
                try:
                    highp = daily_price['PX_HIGH']
                except:
                    highp = closep
                try:
                    lowp = daily_price['PX_LOW']
                except:
                    lowp = closep
                try:
                    imp_vol = daily_price['3MTH_IMPVOL_100.0%MNY_DF']
                except:
                    imp_vol = pd.DataFrame()
                if len(weekly_price) > 0:
                    closep_w = weekly_price['PX_LAST']
                else:
                    closep_w = pd.Series()
                try:
                    highp_w = weekly_price['PX_HIGH']
                    lowp_w = weekly_price['PX_LOW']
                except:
                    highp_w = closep_w
                    lowp_w = closep_w
                if len(monthly_price) > 0:
                    closep_m = monthly_price['PX_LAST']
                else:
                    closep_m = pd.Series()
                try:
                    highp_m = monthly_price['PX_HIGH']
                    lowp_m = monthly_price['PX_LOW']
                except:
                    highp_m = closep_m
                    lowp_m = closep_m
                if val[4] and len(intraday_price) > 10:
                    date_range_by_minute = pd.DatetimeIndex([dt for dt in datetime_range(intraday_price.index[0], intraday_price.index[-1], dt.timedelta(minutes=10))])
                    date_range_by_minute = date_range_by_minute[date_range_by_minute.dayofweek != 5]
                    date_range_by_minute = date_range_by_minute[date_range_by_minute.dayofweek != 6]
                    volume_by_minute = intraday_price['volume'].reindex(date_range_by_minute)
                    volume_by_minute.fillna(value=0, inplace=True)
                    volume_by_minute = volume_by_minute.to_frame('volume')
                    volume_by_minute['date'] = volume_by_minute.index.date
                    agg_volume = volume_by_minute.groupby(['date'])['volume'].cumsum()
                    agg_volume = agg_volume.to_frame('agg_volume')
                    agg_volume['time'] = volume_by_minute.index.time
                    volume_at_minute = agg_volume.loc[agg_volume['time'] == agg_volume['time'].iloc[-1], 'agg_volume']
                    volume_at_minute.drop(volume_at_minute.loc[volume_at_minute == 0].index, axis=0, inplace=True)
                    if len(volume_at_minute) > 20:
                        volume_spike_5d = volume_at_minute.iloc[-1] / volume_at_minute.rolling(5).mean().iloc[-2]
                        volume_spike_20d = volume_at_minute.iloc[-1] / volume_at_minute.rolling(20).mean().iloc[-2]
                    else:
                        volume_spike_5d = np.nan
                        volume_spike_5d_1 = np.nan
                        volume_spike_20d = np.nan
                    if daily_price.index[-1] >= today():
                        volume_spike_5d_1 = daily_price['VOLUME'].iloc[-2] / daily_price['VOLUME'].rolling(5).mean().iloc[-3]
                    else:
                        volume_spike_5d_1 = np.nan
                else:
                    volume_spike_5d = np.nan
                    volume_spike_5d_1 = np.nan
                    volume_spike_20d = np.nan
                last_price = closep[-1]
                price_change = closep[-1] - closep[-2]
                price_change1 = closep[-2] - closep[-3]
                if chg_type == 'FLAT':
                    price_change_pct = (closep[-1] - closep[-2]) / closep[-2]
                    price_change_pct1 = (closep[-2] - closep[-3]) / closep[-3]
                    daily_price_chg = closep.diff() / closep.shift(1)
                    chg_5d = (closep[-1] - closep[-5]) / closep[-5]
                    if len(intraday_price) > 0:
                        intraday_price_chg = intraday_price['close'].diff() / intraday_price['close'].shift(1)
                elif chg_type == 'SPRD':
                    price_change_pct = price_change
                    price_change_pct1 = price_change1
                    daily_price_chg = closep.diff()
                    chg_5d = closep[-1] - closep[-5]
                    if len(intraday_price) > 0:
                        intraday_price_chg = intraday_price['close'].diff()
                else:
                    raise ('ERROR!')
                try:
                    if np.isnan(daily_price['PX_SETTLE'].iloc[-1]):
                        post_close = daily_price['FUT_PX'].iloc[-2] - daily_price['PX_SETTLE'].iloc[-2]
                    else:
                        post_close = daily_price['FUT_PX'].iloc[-1] - daily_price['PX_SETTLE'].iloc[-1]
                except:
                    post_close = 0
                if isinstance(fom_comments.loc[val[1], 'FOM Dates'], str):
                    if len(row['FOM Dates']) == 10:
                        try:
                            fom_date = dt.datetime.strptime(fom_comments.loc[val[1], 'FOM Dates'], '%Y-%m-%d')
                            fom_date0 = fom_date
                            fom_date1 = today()
                        except:
                            fom_date = dt.datetime.strptime(fom_comments.loc[val[1], 'FOM Dates'], '%m/%d/%Y')
                            fom_date0 = fom_date
                            fom_date1 = today()
                    else:
                        try:
                            fom_date0 = dt.datetime.strptime(fom_comments.loc[val[1], 'FOM Dates'][:10], '%Y-%m-%d')
                        except:
                            fom_date0 = dt.datetime.strptime(fom_comments.loc[val[1], 'FOM Dates'][:10], '%m/%d/%Y')
                        try:
                            fom_date1 = dt.datetime.strptime(fom_comments.loc[val[1], 'FOM Dates'][-10:], '%Y-%m-%d')
                        except:
                            fom_date1 = dt.datetime.strptime(fom_comments.loc[val[1], 'FOM Dates'][-10:], '%m/%d/%Y')
                        fom_date = (fom_date0, fom_date1)
                elif isinstance(fom_comments.loc[val[1], 'FOM Dates'], dt.datetime):
                    fom_date = fom_comments.loc[val[1], 'FOM Dates']
                vwap_track_date_range = dict(sdate=fom_date0, edate=fom_date1)
                try:
                    vwap_overrides = [
                        ("VWAP_START_TIME", "00:00:00"), ("VWAP_END_TIME", "23:00:00"),
                        ("VWAP_START_DT", f"{vwap_track_date_range.get('sdate').strftime('%Y%m%d')}"),
                        ("VWAP_END_DT", f"{vwap_track_date_range.get('edate').strftime('%Y%m%d')}"),
                    ]
                    fom_vwap = bbg.bdh(ticker, ["EQY_WEIGHTED_AVG_PX"], ovrds=vwap_overrides, **vwap_track_date_range)
                    _photo_gap(760, 'possible result selection after Bloomberg VWAP call is beyond photo edge')
                    fom_delta = intraday_price.close[-1] - fom_vwap
                except:
                    fom_vwap = np.nan
                    fom_delta = np.nan
                try:
                    if isinstance(fom_date, tuple):
                        fom_high = highp[(highp.index >= fom_date[0]) & (highp.index <= fom_date[1])].max()
                        fom_low = lowp[(lowp.index >= fom_date[0]) & (lowp.index <= fom_date[1])].min()
                        fom_date = dt.datetime.strftime(fom_date[0], '%Y-%m-%d') + ' - ' + dt.datetime.strftime(fom_date[1], '%Y-%m-%d')
                    else:
                        fom_high = highp[highp.index >= fom_date].iloc[0]
                        fom_low = lowp[lowp.index >= fom_date].iloc[0]
                    if closep[-1] > fom_high:
                        fom_flag = 'A+'
                    elif closep[-1] < fom_low:
                        fom_flag = 'A-'
                    else:
                        fom_flag = None
                except:
                    fom_flag = None
                if len(highp.dropna().values) < 13:
                    print("Not enough high and low price values" f" for {ticker} to do stochastic calculations")
                    continue
                stok, stod, stods = talib.stochastic(highp.values, lowp.values, closep.values, 13, 3, 3)
                if stod[-1] > 80 and stods[-1] > 80:
                    if stod[-2] > stods[-2] and stod[-1] < stods[-1]:
                        sto_ind = 'OB CROSS'
                    else:
                        sto_ind = 'OB'
                elif stod[-1] < 20 and stods[-1] < 20:
                    if stod[-2] < stods[-2] and stod[-1] > stods[-1]:
                        sto_ind = 'OS CROSS'
                    else:
                        sto_ind = 'OS'
                else:
                    sto_ind = None
                sharpe_22d = daily_price_chg.iloc[-23:-1].mean() / daily_price_chg.iloc[-23:-1].std() * np.sqrt(22)
                if len(intraday_price) > 0:
                    upload = True
                    sdate_3d = daily_price_chg.index[-3]
                    edate_3d = daily_price_chg.index[-1] + dt.timedelta(hours=23, minutes=59)
                    sdate_3d_cob = daily_price_chg.index[-4]
                    edate_3d_cob = daily_price_chg.index[-2] + dt.timedelta(hours=23, minutes=59)
                    sdate_3d_cob_1 = daily_price_chg.index[-5]
                    edate_3d_cob_1 = daily_price_chg.index[-3] + dt.timedelta(hours=23, minutes=59)
                    sdate_3d_cob_2 = daily_price_chg.index[-6]
                    edate_3d_cob_2 = daily_price_chg.index[-4] + dt.timedelta(hours=23, minutes=59)
                    sdate_3d_cob_3 = daily_price_chg.index[-7]
                    edate_3d_cob_3 = daily_price_chg.index[-5] + dt.timedelta(hours=23, minutes=59)
                    if edate_3d > intraday_price_chg.index[-1] or edate_3d not in intraday_price_chg.index:
                        three_d_chg = intraday_price_chg[(intraday_price_chg.index >= sdate_3d) & (intraday_price_chg.index <= edate_3d)]
                    else:
                        three_d_chg = intraday_price_chg.loc[sdate_3d:edate_3d]
                    if intraday_price_chg.index[-1] > edate_3d_cob - BDay(10):
                        sharpe_3d_cob = sr_3d_cob(intraday_price_chg, sdate_3d_cob, edate_3d_cob, instr, chg_type)
                        sharpe_3d_cob_1 = sr_3d_cob(intraday_price_chg, sdate_3d_cob_1, edate_3d_cob_1, instr, chg_type)
                        sharpe_3d_cob_2 = sr_3d_cob(intraday_price_chg, sdate_3d_cob_2, edate_3d_cob_2, instr, chg_type)
                        sharpe_3d_cob_3 = sr_3d_cob(intraday_price_chg, sdate_3d_cob_3, edate_3d_cob_3, instr, chg_type)
                        fig = go.Figure()
                        fig.add_trace(go.Candlestick(x=daily_price.index, open=daily_price['PX_OPEN'], high=daily_price['PX_HIGH'], low=daily_price['PX_LOW'], close=daily_price['PX_LAST']))
                        fig.update_layout(xaxis_rangeslider_visible=False, title=f"{key} OHLC")
                        table.figures_to_html([fig], filename=ut.convert_path_to_linux(f"{folder_path}gas\\candle_plots\\{key}.html"))
                    else:
                        sharpe_3d_cob = np.nan
                        sharpe_3d_cob_1 = np.nan
                        sharpe_3d_cob_2 = np.nan
                        sharpe_3d_cob_3 = np.nan
                    if len(three_d_chg) > 10:
                        three_d_chg = three_d_chg.drop(three_d_chg.index[0])
                        chg_3d = three_d_chg.mean()
                        stdev_3d = three_d_chg.std()
                        try:
                            sharpe_3d = chg_3d / stdev_3d * np.sqrt(len(three_d_chg))
                        except:
                            sharpe_3d = np.nan
                else:
                    sharpe_3d = np.nan
                    sharpe_3d_cob = np.nan
                    sharpe_3d_cob_1 = np.nan
                    sharpe_3d_cob_2 = np.nan
                    sharpe_3d_cob_3 = np.nan
                true_range = talib.tr(highp.values, lowp.values, closep.values)
                avg_true_range = talib.atr(highp.values, lowp.values, closep, 5)
                range_spike = true_range[-1] / avg_true_range[-2]
                range_spike1 = true_range[-2] / avg_true_range[-3]
                if closep[-1] > closep[-2] > closep[-3] > closep[-4] and closep[-4] < closep[-5]:
                    three_day = 'UP'
                elif closep[-1] < closep[-2] < closep[-3] < closep[-4] and closep[-4] > closep[-5]:
                    three_day = 'DOWN'
                else:
                    three_day = None
                if closep[-1] > highp[-1] - (highp[-1] - lowp[-1]) * .1:
                    current_price_range = 'HIGH'
                elif closep[-1] < lowp[-1] + (highp[-1] - lowp[-1]) * .1:
                    current_price_range = 'LOW'
                else:
                    current_price_range = None
                trend_index = 0
                if price_change > 0:
                    if range_spike > 1.2:
                        trend_index += 1
                    if volume_spike_5d > 1.2:
                        trend_index += 1
                    if three_day == 'UP':
                        trend_index += .5
                    if sharpe_3d > 1:
                        trend_index += sharpe_3d
                    elif sharpe_3d > .5 and sharpe_22d > 1:
                        trend_index += .5
                    if current_price_range == 'HIGH':
                        trend_index += .5
                    if post_close / avg_true_range[-2] > .2:
                        trend_index += .5
                elif price_change <= 0:
                    if range_spike > 1.2:
                        trend_index -= 1
                    if volume_spike_5d > 1.2:
                        trend_index -= 1
                    if three_day == 'DOWN':
                        trend_index -= .5
                    if sharpe_3d < -1:
                        trend_index += sharpe_3d
                    elif sharpe_3d < -.5 and sharpe_22d < -1:
                        trend_index -= .5
                    if current_price_range == 'LOW':
                        trend_index -= .5
                    if post_close / avg_true_range[-2] < -.2:
                        trend_index -= .5
                fdhs, fdls = talib.fisher_indicator(highp, lowp, closep)
                fdh = fdhs[-1]
                fdl = fdls[-1]
                fdh1 = fdhs[-2]
                fdl1 = fdls[-2]
                fdh2 = fdhs[-3]
                fdl2 = fdls[-3]
                fwhs, fwls = talib.fisher_indicator(highp_w, lowp_w, closep_w)
                fwhs = fwhs.reindex(daily_price.index)
                fwhs = fwhs.shift(1)
                fwhs.fillna(method='ffill', inplace=True)
                fwls = fwls.reindex(daily_price.index)
                fwls = fwls.shift(1)
                fwls.fillna(method='ffill', inplace=True)
                fwh = fwhs[-1]
                fwl = fwls[-1]
                fwh1 = fwhs[-2]
                fwl1 = fwls[-2]
                fwh2 = fwhs[-3]
                fwl2 = fwls[-3]
                fmhs, fmls = talib.fisher_indicator(highp_m, lowp_m, closep_m)
                fmhs = fmhs.reindex(daily_price.index)
                fmhs = fmhs.shift(1)
                fmhs.fillna(method='ffill', inplace=True)
                fmls = fmls.reindex(daily_price.index)
                fmls = fmls.shift(1)
                fmls.fillna(method='ffill', inplace=True)
                fmh = fmhs[-1]
                fml = fmls[-1]
                fmh1 = fmhs[-2]
                fml1 = fmls[-2]
                fmh2 = fmhs[-3]
                fml2 = fmls[-3]
                if fdl > fwh and fwh > fmh and fwl > fml and closep[-1] > fwl and (fdl1 < fwh1 or fwh1 < fmh1 or fwl1 < fml1):
                    fish = 'BUY'
                elif fdh < fwl and fwl < fml and fwh < fmh and closep[-1] < fwh and (fdh1 > fwl1 or fwl1 > fml1 or fwh1 > fmh1):
                    fish = 'SELL'
                else:
                    fish = None
                if fdl1 > fwh1 and fwh1 > fmh1 and fwl1 > fml1 and closep[-2] > fwl1 and (fdl2 < fwh2 or fwh2 < fmh2 or fwl2 < fml2):
                    fish1 = 'BUY'
                elif fdh1 < fwl1 and fwl1 < fml1 and fwh1 < fmh1 and closep[-2] < fwh1 and (fdh2 > fwl2 or fwl2 > fml2 or fwh2 > fmh2):
                    fish1 = 'SELL'
                else:
                    fish1 = None
                if fwl > fmh and closep[-1] > fmh and fwl1 < fmh1:
                    fishlt = 'BUY'
                elif fwh < fml and closep[-1] < fml and fwh1 > fml1:
                    fishlt = 'SELL'
                else:
                    fishlt = None
                if fwl1 > fmh1 and closep[-2] > fmh1 and fwl2 < fmh2:
                    fishlt1 = 'BUY'
                elif fwh1 < fml1 and closep[-2] < fml1 and fwh2 > fml2:
                    fishlt1 = 'SELL'
                else:
                    fishlt1 = None
                if closep[-1] > closep[-2] and closep[-1] > fwl and fdl > fwh and fwh < fml and fdl1 < fwh1:
                    counter = 'BUY'
                elif closep[-1] < closep[-2] and closep[-1] < fwh and fdh < fwl and fwl > fmh and fdh1 > fwl1:
                    counter = 'SELL'
                else:
                    counter = None
                if closep[-2] > closep[-3] and closep[-2] > fwl1 and fdl1 > fwh1 and fwh1 < fml1 and fdl2 < fwh2:
                    counter1 = 'BUY'
                elif closep[-2] < closep[-3] and closep[-2] < fwh1 and fdh1 < fwl1 and fwl1 > fmh1 and fdh2 > fwl2:
                    counter1 = 'SELL'
                else:
                    counter1 = None
                emva3 = talib.mva(closep.values, 3, 'e')
                emva8 = talib.mva(closep.values, 8, 'e')
                try:
                    emva14 = talib.mva(closep.values, 14, 'e')
                except:
                    emva14 = emva8
                try:
                    emva30 = talib.mva(closep.values, 30, 'e')
                except:
                    emva30 = emva8
                if closep[-1] > closep[-2] and emva3[-1] > emva8[-1] > emva14[-1] and closep[-1] > emva8[-1] and (emva3[-2] < emva8[-2] or emva8[-2] < emva14[-2]):
                    mom = 'BUY'
                elif closep[-1] < closep[-2] and emva3[-1] < emva8[-1] < emva14[-1] and closep[-1] < emva8[-1] and (emva3[-2] > emva8[-2] or emva8[-2] > emva14[-2]):
                    mom = 'SELL'
                else:
                    mom = None
                if closep[-2] > closep[-3] and emva3[-2] > emva8[-2] > emva14[-2] and closep[-2] > emva8[-2] and (emva3[-3] < emva8[-3] or emva8[-3] < emva14[-3]):
                    mom1 = 'BUY'
                elif closep[-2] < closep[-3] and emva3[-2] < emva8[-2] < emva14[-2] and closep[-2] < emva8[-2] and (emva3[-3] > emva8[-3] or emva8[-3] > emva14[-3]):
                    mom1 = 'SELL'
                else:
                    mom1 = None
                if emva8[-1] > emva14[-1] > emva30[-1] and closep[-1] > emva14[-1] and (emva8[-2] < emva14[-2] or emva14[-2] < emva30[-2]):
                    hybrid = 'BUY'
                elif emva8[-1] < emva14[-1] < emva30[-1] and closep[-1] < emva14[-1] and (emva8[-2] > emva14[-2] or emva14[-2] > emva30[-2]):
                    hybrid = 'SELL'
                else:
                    hybrid = None
                if emva8[-2] > emva14[-2] > emva30[-2] and closep[-2] > emva14[-2] and (emva8[-3] < emva14[-3] or emva14[-3] < emva30[-3]):
                    hybrid1 = 'BUY'
                elif emva8[-2] < emva14[-2] < emva30[-2] and closep[-2] < emva14[-2] and (emva8[-3] > emva14[-3] or emva14[-3] > emva30[-3]):
                    hybrid1 = 'SELL'
                else:
                    hybrid1 = None
                mom_signal = None
                hybrid_signal = None
                fish_signal = None
                fishlt_signal = None
                counter_signal = None
                mom_signal_ystd = None
                hybrid_signal_ystd = None
                fish_signal_ystd = None
                fishlt_signal_ystd = None
                counter_signal_ystd = None
                if trend_index > 2 and price_change > 0:
                    if mom == 'BUY' or mom1 == 'BUY':
                        mom_signal = 'BUY'
                    if mom1 == 'BUY':
                        mom_signal_ystd = 'BUY'
                    if hybrid == 'BUY' or hybrid1 == 'BUY':
                        hybrid_signal = 'BUY'
                    if hybrid1 == 'BUY':
                        hybrid_signal_ystd = 'BUY'
                    if fish == 'BUY' or fish1 == 'BUY':
                        fish_signal = 'BUY'
                    if fish1 == 'BUY':
                        fish_signal_ystd = 'BUY'
                    if fishlt == 'BUY' or fishlt1 == 'BUY':
                        fishlt_signal = 'BUY'
                    if fishlt1 == 'BUY':
                        fishlt_signal_ystd = 'BUY'
                    if counter == 'BUY' or counter1 == 'BUY':
                        counter_signal = 'BUY'
                    if counter1 == 'BUY':
                        counter_signal_ystd = 'BUY'
                elif trend_index < -2 and price_change < 0:
                    if mom == 'SELL' or mom1 == 'SELL':
                        mom_signal = 'SELL'
                    if mom1 == 'SELL':
                        mom_signal_ystd = 'SELL'
                    if hybrid == 'SELL' or hybrid1 == 'SELL':
                        hybrid_signal = 'SELL'
                    if hybrid1 == 'SELL':
                        hybrid_signal_ystd = 'SELL'
                    if fish == 'SELL' or fish1 == 'SELL':
                        fish_signal = 'SELL'
                    if fish1 == 'SELL':
                        fish_signal_ystd = 'SELL'
                    if fishlt == 'SELL' or fishlt1 == 'SELL':
                        fishlt_signal = 'SELL'
                    if fishlt1 == 'SELL':
                        fishlt_signal_ystd = 'SELL'
                    if counter == 'SELL' or counter1 == 'SELL':
                        counter_signal = 'SELL'
                    if counter1 == 'SELL':
                        counter_signal_ystd = 'SELL'
                if emva3[-2] > emva8[-2] > emva14[-2]:
                    mom2 = 'B'
                elif emva3[-2] < emva8[-2] < emva14[-2]:
                    mom2 = 'S'
                else:
                    mom2 = None
                if emva8[-2] > emva14[-2] > emva30[-2]:
                    hybrid2 = 'B'
                elif emva8[-2] < emva14[-2] < emva30[-2]:
                    hybrid2 = 'S'
                else:
                    hybrid2 = None
                if fdl1 > fwh1 > fmh1 and fwl1 > fml1:
                    fish2 = 'B'
                elif fdh1 < fwl1 < fml1 and fwh1 < fmh1:
                    fish2 = 'S'
                else:
                    fish2 = None
                if fwl1 > fmh1:
                    fishlt2 = 'B'
                elif fwh1 < fml1:
                    fishlt2 = 'S'
                else:
                    fishlt2 = None
                if fdl1 > fwh1 and fwh1 < fml:
                    counter2 = 'B'
                elif fdh1 < fwl1 and fwl1 > fmh1:
                    counter2 = 'S'
                else:
                    counter2 = None
                if len(imp_vol) == 0:
                    imp_vol = daily_price_chg.rolling(20).std() * 100 * np.sqrt(252)
                if len(imp_vol) > 0:
                    imp_vol.fillna(method='ffill', inplace=True)
                    if chg_type == 'SPRD':
                        kpi_on = (closep[-1] - closep[-2]) / (imp_vol[-2] / 100 / np.sqrt(252))
                        kpi_1d = (closep[-2] - closep[-3]) / (imp_vol[-3] / 100 / np.sqrt(252))
                        kpi = (closep - closep.shift(1)) / (imp_vol.shift(1) / 100 / np.sqrt(252))
                        kpi_avg = talib.exp_avg(kpi[-13:].values, a=1 - 2 / (13 + 1))
                    elif chg_type == 'FLAT':
                        kpi_on = (closep[-1] - closep[-2]) / (closep[-1] * imp_vol[-2] / 100 / np.sqrt(252))
                        kpi_1d = (closep[-2] - closep[-3]) / (closep[-2] * imp_vol[-3] / 100 / np.sqrt(252))
                        kpi = (closep - closep.shift(1)) / (closep * imp_vol.shift(1) / 100 / np.sqrt(252))
                        kpi_avg = talib.exp_avg(kpi[-13:].values, a=1 - 2 / (13 + 1))
                else:
                    kpi_on = np.nan
                    kpi_1d = np.nan
                    kpi_avg = np.nan
                if (mom_signal == 'BUY' or hybrid_signal == 'BUY' or fish_signal == 'BUY' or fishlt_signal == 'BUY' or
                    counter_signal == 'BUY' or trend_index > 3 or
                    mom_signal == 'SELL' or hybrid_signal == 'SELL' or fish_signal == 'SELL' or fishlt_signal == 'SELL' or
                    counter_signal == 'SELL' or trend_index < -3):
                    signal = 1
                else:
                    signal = 0
                output_list.append([instr, ticker, days_before_exp, last_price, kpi_on, kpi_1d,
                                    price_change, chg_type, fom_flag, fom_date, sto_ind, post_close, sharpe_3d, sharpe_22d,
                                    mom2, hybrid2, fish2, fishlt2, counter2,
                                    mom_signal, hybrid_signal, fish_signal, fishlt_signal, counter_signal,
                                    kpi_avg, trend_index, three_day, range_spike, volume_spike_5d, volume_spike_20d,
                                    current_price_range, signal, mom_signal_ystd, hybrid_signal_ystd, fish_signal_ystd,
                                    fishlt_signal_ystd, counter_signal_ystd, price_change_pct, chg_5d, fom_vwap, fom_delta])
                if range_spike1 * volume_spike_5d_1 > 2.25 and daily_price.index[-1] >= today():
                    follow_through = price_change_pct / price_change_pct1
                    alert_email_all.append([instr, ticker, follow_through, price_change_pct, range_spike, volume_spike_5d,
                                            price_change_pct1, range_spike1, volume_spike_5d_1, chg_type])
                if (sharpe_3d_cob > 0 and sharpe_22d > 0) or (sharpe_3d_cob < 0 and sharpe_22d < 0):
                    sr_sig = "TREND"
                elif (sharpe_3d_cob > 1 and sharpe_22d < -0.5) or (sharpe_3d_cob < -1 and sharpe_22d > 0.5):
                    sr_sig = "COUNTER"
                else:
                    sr_sig = "-"
                plot_link = f"<a href='{folder_path}gas\\candle_plots\\{key}.html'>{ticker}</a>"
                sharpe_ratio_list.append([instr, ticker, chg_type, sharpe_3d, sharpe_3d_cob, sharpe_22d, sr_sig,
                                          sharpe_3d_cob_1, sharpe_3d_cob_2, sharpe_3d_cob_3, plot_link])
            else:
                output_list.append([instr, ticker, np.nan, np.nan, np.nan, np.nan, np.nan, np.nan, np.nan, np.nan, np.nan, np.nan, np.nan, np.nan, np.nan, np.nan, np.nan, np.nan, np.nan, np.nan, np.nan, np.nan, np.nan, np.nan, np.nan, np.nan, np.nan, np.nan, np.nan, np.nan, np.nan, np.nan, np.nan, np.nan, np.nan, np.nan, np.nan, np.nan, np.nan, np.nan, np.nan])
    df = pd.DataFrame(output_list, columns=['Instr', 'Ticker', 'Expiry', 'Last Price', 'KPI', 'KPI t-1',
                                            'Change', 'Type', 'FOM', 'Ref Date', 'Stochs', 'Post Close', 'Shrp 3', 'Shrp 22',
                                            'M', 'H', 'F', 'LT', 'CT', 'Mom', 'Hybrid', 'Fish', 'FishLT', 'Counter',
                                            'KPI Index', 'Index', '3Day', 'RngSpike', 'VolSpike', 'Volspike20',
                                            'Percentile', 'Signal', 'Mom_ytsd', 'Hybrid_ytsd', 'Fish_ytsd',
                                            'FishLT_ytsd', 'Counter_ytsd', 'Change 1d', 'Change 5d', 'FOM Vwap', 'FOM PxDelta'])
    df.index.name = time.strftime('%Y-%m-%d %H:%M')
    df_all = pd.DataFrame(alert_email_all, columns=['Instr', 'Ticker', 'Follow Through', 'Change T', 'Range T', 'Volume T',
                                                   'Change T-1', 'Range T-1', 'Volume T-1', 'chg type'])
    sharpe_ratio_df = pd.DataFrame(sharpe_ratio_list,
                                   columns=['Instr', 'Ticker', 'Type', 'Shrp 3D Live', 'Shrp 3D COB', 'Shrp 22D', 'Signal',
                                            'Shrp 3D COB 1', 'Shrp 3D COB 2', 'Shrp 3D COB 3', '_plot_link'])
    return df, df_all, sharpe_ratio_df


def market_scan():
    ut.tic()
    ticker_dict = get_tickers()
    df, df_all, sharpe_ratio = get_basic_info(ticker_dict)
    df.to_csv(ut.convert_path_to_linux(f"{folder_path_gas}trend_signal_gas.csv"))
    df_all.to_csv(ut.convert_path_to_linux(f"{folder_path_gas}trend_alert_gas.csv"))
    sharpe_ratio.to_csv(ut.convert_path_to_linux(f"{folder_path_gas}sharpe_ratio_gas_power.csv"))
    ut.toc()


def table_format(excelpath, data, name="sheet1", title=None, save_subset_to_new_sheet=False, subset_range=None):
    writer = pd.ExcelWriter(excelpath, engine="xlsxwriter")
    df = pd.DataFrame()
    df.to_excel(writer, index=False, sheet_name=name)
    if save_subset_to_new_sheet:
        df.to_excel(writer, index=False, sheet_name="sheet2")

    workbook = writer.book
    worksheet = writer.sheets[name]

    head_format = workbook.add_format()
    head_format.set_bold()
    head_format.set_align("center")
    head_format.set_align("vcenter")
    head_format.set_bottom(1)

    head_format_border = workbook.add_format()
    head_format_border.set_bold()
    head_format_border.set_left(5)
    head_format_border.set_bottom(1)
    head_format_border.set_align("center")
    head_format_border.set_align("vcenter")

    head_format_bold_border = workbook.add_format()
    head_format_bold_border.set_bold()
    head_format_bold_border.set_right(1)
    head_format_bold_border.set_bottom(1)
    head_format_bold_border.set_align("center")
    head_format_bold_border.set_align("vcenter")

    head_format_dash_border = workbook.add_format()
    head_format_dash_border.set_bold()
    head_format_dash_border.set_right(7)
    head_format_dash_border.set_bottom(1)
    head_format_dash_border.set_align("center")
    head_format_dash_border.set_align("vcenter")


    str_format1 = workbook.add_format()
    str_format1.set_align("center")
    str_format1.set_align("vcenter")

    str_format1_bold_border = workbook.add_format()
    str_format1_bold_border.set_align("center")
    str_format1_bold_border.set_align("vcenter")
    str_format1_bold_border.set_right(1)

    str_format1_dash_border = workbook.add_format()
    str_format1_dash_border.set_align("center")
    str_format1_dash_border.set_align("vcenter")
    str_format1_dash_border.set_right(7)

    str_format_border = workbook.add_format()
    str_format_border.set_align("center")
    str_format_border.set_align("vcenter")
    str_format_border.set_left(5)

    str_format_bold = workbook.add_format()
    str_format_bold.set_align("center")
    str_format_bold.set_align("vcenter")
    str_format_bold.set_bold()


    str_format1_red_bold_boarder = workbook.add_format()
    str_format1_red_bold_boarder.set_align("center")
    str_format1_red_bold_boarder.set_align("center")
    str_format1_red_bold_boarder.set_bg_color('#FFC7CE')
    str_format1_red_bold_boarder.set_font_color('#9C0006')
    str_format1_red_bold_boarder.set_right(1)

    str_format1_red_dash_boarder = workbook.add_format()
    str_format1_red_dash_boarder.set_align("center")
    str_format1_red_dash_boarder.set_align("vcenter")
    str_format1_red_dash_boarder.set_bg_color('#FFC7CE')
    str_format1_red_dash_boarder.set_font_color('#9C0006')
    str_format1_red_dash_boarder.set_right(7)


    str_format1_green_bold_boarder = workbook.add_format()
    str_format1_green_bold_boarder.set_align("center")
    str_format1_green_bold_boarder.set_align("vcenter")
    str_format1_green_bold_boarder.set_bg_color('#C6EFCE')
    str_format1_green_bold_boarder.set_font_color('#006100')
    str_format1_green_bold_boarder.set_right(1)

    str_format1_green_dash_boarder = workbook.add_format()
    str_format1_green_dash_boarder.set_align("center")
    str_format1_green_dash_boarder.set_align("vcenter")
    str_format1_green_dash_boarder.set_bg_color('#C6EFCE')
    str_format1_green_dash_boarder.set_font_color('#006100')
    str_format1_green_dash_boarder.set_right(7)


    num_format = workbook.add_format()
    num_format.set_num_format("0.000")
    num_format.set_align("center")
    num_format.set_align("vcenter")


    num_format1 = workbook.add_format()
    num_format1.set_num_format("0.00;[Red]-0.00")
    num_format1.set_align("center")
    num_format1.set_align("vcenter")
    num_format1.set_right(1)

    num_format1_dash_boarder = workbook.add_format()
    num_format1_dash_boarder.set_num_format("0.00;[Red]-0.00")
    num_format1_dash_boarder.set_align("center")
    num_format1_dash_boarder.set_align("vcenter")
    num_format1_dash_boarder.set_right(7)


    num_format1_red = workbook.add_format()
    num_format1_red.set_num_format("0.00;[Red]-0.00")
    num_format1_red.set_align("center")
    num_format1_red.set_align("vcenter")
    num_format1_red.set_bg_color('#FFC7CE')
    num_format1_red.set_bold()


    num_format1_green = workbook.add_format()
    num_format1_green.set_num_format("0.00;[Red]-0.00")
    num_format1_green.set_align("center")
    num_format1_green.set_align("vcenter")
    num_format1_green.set_bg_color('#C6EFCE')
    num_format1_green.set_bold()


    num_format2 = workbook.add_format()
    num_format2.set_num_format("0%;[Red]-0%")
    num_format2.set_align("center")
    num_format2.set_align("vcenter")
    num_format2.set_right(1)

    num_format2_dash_boarder = workbook.add_format()
    num_format2_dash_boarder.set_num_format("0%;[Red]-0%")
    num_format2_dash_boarder.set_align("center")
    num_format2_dash_boarder.set_align("vcenter")
    num_format2_dash_boarder.set_right(7)

    num_format2_red = workbook.add_format()
    num_format2_red.set_num_format("0%;[Red]-0%")
    num_format2_red.set_align("center")
    num_format2_red.set_align("vcenter")
    num_format2_red.set_bg_color('#FFC7CE')
    num_format2_red.set_bold()
    num_format2_red.set_right(1)

    num_format2_red_dash_boarder = workbook.add_format()
    num_format2_red_dash_boarder.set_num_format("0%;[Red]-0%")
    num_format2_red_dash_boarder.set_align("center")
    num_format2_red_dash_boarder.set_align("vcenter")
    num_format2_red_dash_boarder.set_bg_color('#FFC7CE')
    num_format2_red_dash_boarder.set_bold()
    num_format2_red_dash_boarder.set_right(7)

    num_format2_green = workbook.add_format()
    num_format2_green.set_num_format("0%;[Red]-0%")
    num_format2_green.set_align("center")
    num_format2_green.set_align("vcenter")
    num_format2_green.set_bg_color('#C6EFCE')
    num_format2_green.set_bold()
    num_format2_green.set_right(1)

    num_format2_green_dash_boarder = workbook.add_format()
    num_format2_green_dash_boarder.set_num_format("0%;[Red]-0%")
    num_format2_green_dash_boarder.set_align("center")
    num_format2_green_dash_boarder.set_align("vcenter")
    num_format2_green_dash_boarder.set_bg_color('#C6EFCE')
    num_format2_green_dash_boarder.set_bold()
    num_format2_green_dash_boarder.set_right(7)

    num_format6 = workbook.add_format()
    num_format6.set_num_format("0.00%;[Red]-0.00%")
    num_format6.set_align("center")
    num_format6.set_align("vcenter")


    num_format3 = workbook.add_format()
    num_format3.set_num_format("0;[Red]-0")
    num_format3.set_align("center")
    num_format3.set_align("vcenter")
    num_format3.set_right(1)

    leftFormat = workbook.add_format({"left": 5})
    topFormat = workbook.add_format({"top": 5})

    worksheet.set_column("A:A", 10)
    worksheet.set_column("B:B", 16)
    worksheet.set_column("C:C", 6)
    worksheet.set_column("D:D", 0)
    worksheet.set_column("E:E", 6)
    worksheet.set_column("F:F", 6)
    worksheet.set_column("G:G", 0)
    worksheet.set_column("H:H", 0)
    worksheet.set_column("I:I", 4)
    worksheet.set_column("J:J", 0)
    worksheet.set_column("K:K", 6)
    worksheet.set_column("L:L", 8)
    worksheet.set_column("M:M", 6)
    worksheet.set_column("N:N", 6)
    worksheet.set_column("O:O", 2)
    worksheet.set_column("P:P", 2)
    worksheet.set_column("Q:Q", 2)
    worksheet.set_column("R:R", 2)
    worksheet.set_column("S:S", 2)
    worksheet.set_column("T:T", 4)
    worksheet.set_column("U:U", 4)
    worksheet.set_column("V:V", 4)
    worksheet.set_column("W:W", 4)
    worksheet.set_column("X:X", 4)
    worksheet.set_column("Y:Y", 6)
    worksheet.set_column("Z:Z", 6)
    worksheet.set_column("AA:AA", 6)
    worksheet.set_column("AB:AB", 8)
    worksheet.set_column("AC:AC", 8)
    worksheet.set_column("AD:AD", 8)

    row_num, col_num = data.shape
    scol = 1

    if title is not None:
        worksheet.write_string(0, 0, title, head_format)
        merge_format = workbook.add_format(
            {
                "bold": 1,
                "font_size": 12,
                "top": 5,
                "left": 5,
                "right": 5,
                "bottom": 1,
                "align": "center",
                "valign": "vcenter",
                "fg_color": "white",
            }
        )
        worksheet.merge_range("A1:AD1", title, merge_format)
        srow = 1
    else:
        srow = 0

    for row in range(srow, data.shape[0] + 1 + srow):
        worksheet.write(row, data.shape[1], "", leftFormat)

    for col in range(0, data.shape[1]):
        worksheet.write(data.shape[0] + 1 + srow, col, "", topFormat)

    for i in range(scol, col_num + 1):
        if i == 0:
            worksheet.write_string(srow, i, data.index.name, head_format_border)
        else:
            if i == scol:
                worksheet.write_string(srow, i - scol, data.columns[i - scol], head_format_border)
            else:
                if i in [3, 6, 12, 14, 19, 24, 27]:
                    worksheet.write_string(srow, i - scol, data.columns[i - scol], head_format_bold_border)
                else:
                    worksheet.write_string(srow, i - scol, data.columns[i - scol], head_format_dash_border)

    for i in range(1, row_num + 1):
        for j in range(scol, col_num + 1):
            if j in [
                1,
                2,
                8,
                9,
                10,
                11,
                15,
                16,
                17,
                18,
                19,
                20,
                21,
                22,
                23,
                24,
                30,
                31,
                32,
                33,
                34,
                35,
            ]:
                if data.iloc[i - 1, j - scol] in ["B", "BUY", "A+", "OS", "OS CROSS"]:
                    if j in [19, 24, 30]:
                        worksheet.write(
                            i + srow,
                            j - scol,
                            data.iloc[i - 1, j - scol],
                            str_format1_green_bold_boarder,
                        )
                    else:
                        worksheet.write(
                            i + srow,
                            j - scol,
                            data.iloc[i - 1, j - scol],
                            str_format1_green_dash_boarder,
                        )
                elif data.iloc[i - 1, j - scol] in ["S", "SELL", "A-", "OB", "OB CROSS"]:
                    if j in [19, 24, 30]:
                        worksheet.write(
                            i + srow,
                            j - scol,
                            data.iloc[i - 1, j - scol],
                            str_format1_red_bold_boarder,
                        )
                    else:
                        worksheet.write(
                            i + srow,
                            j - scol,
                            data.iloc[i - 1, j - scol],
                            str_format1_red_dash_boarder,
                        )
                else:
                    if j == 1:
                        worksheet.write(i + srow, j - scol, data.iloc[i - 1, j - scol], str_format_border)
                    elif j == 2:
                        if data.iloc[i - 1, 31:35].any():
                            worksheet.write(i + srow, j - scol, data.iloc[i - 1, j - scol], str_format1)
                        else:
                            worksheet.write(i + srow, j - scol, data.iloc[i - 1, j - scol], str_format_bold)
                    else:
                        try:
                            if j in [19, 24, 30]:
                                worksheet.write(
                                    i + srow,
                                    j - scol,
                                    data.iloc[i - 1, j - scol],
                                    str_format1_bold_border,
                                )
                            else:
                                worksheet.write(
                                    i + srow,
                                    j - scol,
                                    data.iloc[i - 1, j - scol],
                                    str_format1_dash_border,
                                )
                        except:
                            if j in [19, 24, 30]:
                                worksheet.write(i + srow, j - scol, "", str_format1_bold_border)
                            else:
                                worksheet.write(i + srow, j - scol, "", str_format1_dash_border)
            else:
                if j in [
                    3,
                ]:
                    if not np.isnan(data.iloc[i - 1, j - scol]):
                        worksheet.write(i + srow, j - scol, data.iloc[i - 1, j - scol], num_format3)
                    else:
                        worksheet.write(i + srow, j - scol, "", num_format3)
                elif j in [5, 6, 25, 27, 28] and not np.isnan(data.iloc[i - 1, j - scol]):
                    if data.iloc[i - 1, j - scol] > 1.5:
                        if j in [6, 27]:
                            worksheet.write(i + srow, j - scol, data.iloc[i - 1, j - scol], num_format2_green)
                        else:
                            worksheet.write(
                                i + srow,
                                j - scol,
                                data.iloc[i - 1, j - scol],
                                num_format2_green_dash_boarder,
                            )
                    elif data.iloc[i - 1, j - scol] < -1.5:
                        if j in [6, 27]:
                            worksheet.write(i + srow, j - scol, data.iloc[i - 1, j - scol], num_format2_red)
                        else:
                            worksheet.write(
                                i + srow,
                                j - scol,
                                data.iloc[i - 1, j - scol],
                                num_format2_red_dash_boarder,
                            )
                    else:
                        if j in [6, 27]:
                            worksheet.write(i + srow, j - scol, data.iloc[i - 1, j - scol], num_format2)
                        else:
                            worksheet.write(
                                i + srow,
                                j - scol,
                                data.iloc[i - 1, j - scol],
                                num_format2_dash_boarder,
                            )
                elif j in [
                    8,
                ] and not np.isnan(data.iloc[i - 1, j - scol]):
                    worksheet.write(i + srow, j - scol, data.iloc[i - 1, j - scol], num_format6)
                else:
                    if not np.isnan(data.iloc[i - 1, j - scol]):
                        if (
                            j
                            in [
                                26,
                            ]
                            and data.iloc[i - 1, j - scol] > 3
                        ):
                            worksheet.write(i + srow, j - scol, data.iloc[i - 1, j - scol], num_format1_green)
                        elif (
                            j
                            in [
                                26,
                            ]
                            and data.iloc[i - 1, j - scol] < -3
                        ):
                            worksheet.write(i + srow, j - scol, data.iloc[i - 1, j - scol], num_format1_red)
                        else:
                            if j in [12, 14]:
                                worksheet.write(i + srow, j - scol, data.iloc[i - 1, j - scol], num_format1)
                            else:
                                worksheet.write(
                                    i + srow,
                                    j - scol,
                                    data.iloc[i - 1, j - scol],
                                    num_format1_dash_boarder,
                                )
    writer.close()
    if save_subset_to_new_sheet:
        wb = openpyxl.load_workbook(excelpath, data_only=True)
        worksheet_in = wb["sheet1"]
        worksheet_out = wb["sheet2"]
        try:
            if subset_range is None:
                print(f"Error cannot save subset to sheet2, as subet is {subset_range}")
                return
            output_subset = worksheet_in[subset_range]
        except Exception as e:
            print(f"Error cannot save subset to sheet2, with {subset_range} due to {e}")
            return
        for i in output_subset:
            for cell in i:
                new_cell = worksheet_out.cell(row=cell.row, column=cell.column, value=cell.value)
                if cell.has_style:
                    new_cell.font = copy(cell.font)
                    new_cell.border = copy(cell.border)
                    new_cell.fill = copy(cell.fill)
                    new_cell.number_format = copy(cell.number_format)
                    new_cell.protection = copy(cell.protection)
                    new_cell.alignment = copy(cell.alignment)
        worksheet_out.merge_cells("A1:AD1")
        wb.save(excelpath)

    data[["KPI Index", "RngSpike", "VolSpike", "KPI", "KPI t-1"]] = data[["KPI Index", "RngSpike", "VolSpike", "KPI", "KPI t-1"]] * 100
    data["Expiry"] = data["Expiry"].fillna(-1).astype(int).replace(-1, "-")
    column_format = {}
    true_keys = ["B", "BUY", "A+", "OS", "OS CROSS"]
    false_keys = ["S", "SELL", "A-", "OB", "OB CROSS"]
    for col in data.columns:
        base_format = {"width": "55px", "text-align": "center"}

        if col in ["Ticker"]:
            base_format["width"] = "150px"
        elif col in ["KPI Index", "Index", "RngSpike"]:
            base_format["width"] = "55px"
        elif col in ["Mom", "Hybrid", "Fish", "FishLT", "Counter"]:
            base_format["width"] = "30px"
        elif col in ["M", "H", "F", "LT", "CT", "Instr"]:
            base_format["width"] = "5px"
        if col in [
            "H",
            "F",
            "LT",
            "Type",
            "Stocs",
            "Hybrid",
            "Fish",
            "FishLT",
            "Index",
            "FOM",
            "Stochs",
            "Volspike20",
        ]:
            base_format["right_border"] = {"style": "dotted", "size": 2}
            base_format["left_border"] = {"style": "dotted", "size": 2}
        if col in ["Shrp 22", "KPI t-1", "RngSpike", "Counter", "Expiry", "CT"]:
            base_format["right_border"] = True
        if col in ["Shrp 3"]:
            base_format["left_border"] = True
            base_format["right_border"] = {"style": "dotted", "size": 2}

        if col in ["KPI Index", "RngSpike", "VolSpike", "KPI", "KPI t-1"]:
            base_format["format"] = "{:.2f}%"
            base_format["width"] = "55px"

        if col in ["M", "H", "F", "LT", "CT", "FOM", "Mom", "Hybrid", "Fish", "FishLT", "Counter"]:
            base_format["highlight_on_key"] = {
                "columns": col,
                "true_key": true_keys,
                "false_key": false_keys,
            }
        if col in ["Index"]:
            base_format["highlight_on_range"] = {"columns": col, "min": -3, "max": 3}
        elif col in ["KPI", "KPI t-1", "RngSpike", "KPI Index", "VolSpike"]:
            base_format["highlight_on_range"] = {"columns": col, "min": -150, "max": 150}
        column_format[col] = base_format
    format_row = {i: {"bottom_border": {"style": "dotted"}} for i in range(len(data)) if i != len(data)}
    html_table = table.html_format(
        data,
        header_raw=f'<p style="text-align: center; font-family:Calibri; font-weight:bold; font-size:16px">{title}</p>',
        footer=None,
        hide_cols=[
            "Mom_ytsd",
            "Hybrid_ytsd",
            "Fish_ytsd",
            "FishLT_ytsd",
            "Counter_ytsd",
            "Change 1d",
            "Change 5d",
            *_photo_gap(1662, "remaining hidden column names after Change 5d"),
        ],
        format_column=column_format,
        format_row=format_row,
        one_bg_color=True,
        background_color="white",
        precision=2,
        show_date=False,
    )
    return html_table


def save_output(name, asset_class):
    print(f"Saving {name}")
    df = pd.read_csv(ut.convert_path_to_linux(folder_path_gas + "trend_signal_gas.csv"))
    df_signal = df.loc[(df['Signal'] == 1) & (df['Instr'].isin(asset_class)), :].copy()
    if len(df_signal) > 0:
        df_signal = df_signal.set_index(df_signal.columns[0])
        df_signal = df_signal.sort_values('Index', ascending=False)
        df_signal.drop(['3Day', 'Signal'], inplace=True, axis=1)
        output_path = ut.convert_path_to_linux(folder_path_gas + "sort_signal_{:s}.xlsx".format(name))
        output_path_html = ut.convert_path_to_linux(folder_path_gas + "sort_signal_{:s}.html".format(name))
        html_raw = table_format(excelpath=output_path, data=df_signal, name='sheet1', title=name,
                                **_photo_gap(1684, 'save_subset_to_new_sheet and remaining export arguments'))
        png_path = ut.convert_path_to_linux(folder_path_gas + 'sort_signal_{:s}.png'.format(name))
        c = 0
        while c < 100:
            try:
                if user != "pmlo25_svc" and sys.platform.startwith("win"):
                    excel2img.export_img(output_path, png_path, "", "sheet1!A1:AD{:d}".format(
                        len(df_signal) + _photo_gap(1691, 'row extent after len(df_signal)')))
                with open(output_path_html, "w") as html_file:
                    html_file.write(html_raw)
                break
            except:
                c += 1
                pass
        df_index = df.loc[(df['Index'] > 2) | (df['Index'] < -2), :].copy()
        df_index = df_index.set_index(df_index.columns[0])
        df_index.drop(['3Day', 'Signal'], inplace=True, axis=1)
        output_path = ut.convert_path_to_linux(folder_path_gas + "sort_index_{:s}.xlsx".format(name))
        output_path_html = ut.convert_path_to_linux(folder_path_gas + "sort_index_{:s}.html".format(name))
        html_raw = table_format(excelpath=output_path, data=df_index, name='sheet1', title=name,
                                **_photo_gap(1705, 'save_subset_to_new_sheet and remaining export arguments'))
        png_path = ut.convert_path_to_linux(folder_path_gas + 'sort_index_{:s}.png'.format(name))
        c = 0
        while c < 100:
            try:
                if user != "pmlo25_svc" and sys.platform.startwith("win"):
                    excel2img.export_img(output_path, png_path, "", "sheet1!A1:AD{:d}".format(
                        len(df_index) + _photo_gap(1711, 'row extent after len(df_index)')))
                with open(output_path_html, "w") as html_file:
                    html_file.write(html_raw)
                break
            except:
                c += 1
                pass
    else:
        try:
            os.remove(folder_path_gas + "sort_signal_{:s}.xlsx".format(name))
            os.remove(folder_path_gas + 'sort_signal_{:s}.png'.format(name))
            os.remove(folder_path_gas + 'sort_signal_{:s}.html'.format(name))
            os.remove(folder_path_gas + "sort_index_{:s}.xlsx".format(name))
            os.remove(folder_path_gas + 'sort_index_{:s}.png'.format(name))
            os.remove(folder_path_gas + 'sort_index_{:s}.html'.format(name))
        except:
            pass


def save_watchlist(name):
    df = pd.read_csv(ut.convert_path_to_linux(folder_path_gas + "trend_signal_alex.csv"))
    df_watchlist = pd.read_csv(ut.convert_path_to_linux(_photo_gap(1733, r'\\elementcapital.corp\ecns01\PM\Michel Kikano\...')))
    df_wl_name = df_watchlist.loc[df_watchlist['User'] == name, :]
    df_signal = df.loc[df['Ticker'].isin(df_wl_name['Ticker']), :].copy()
    if len(df_signal) > 0:
        _photo_gap(1738, 'unphotographed source lines1738–1739 before output_path')
        output_path = folder_path_gas + "sort_signal_{:s}.xlsx".format(name)
        output_path_html = folder_path_gas + "sort_signal_{:s}.html".format(name)
        table_format(excelpath=output_path, data=df_signal, name='sheet1', title=name,
                     **_photo_gap(1742, 'save_subset_to_new_sheet and remaining arguments'))
        png_path = folder_path_gas + 'sort_signal_{:s}.png'.format(name)
        c = 0
        while c < 100:
            try:
                excel2img.export_img(output_path, png_path, "", "sheet1!A1:AD{:d}".format(
                    len(df_signal) + _photo_gap(1748, 'row extent after len(df_signal)')))
                xlsx2html(output_path, output_path_html, sheet="sheet2",
                          default_cell_border=_photo_gap(1749, '1px solid... and any remaining arguments'))
                break
            except:
                c += 1
                pass
        df_index = df.loc[(df['Index'] > 2) | (df['Index'] < -2), :].copy()
        df_index = df_index.set_index(df_index.columns[0])
        df_index.drop(['3Day', 'Signal'], inplace=True, axis=1)
        output_path = folder_path_gas + "sort_index_{:s}.xlsx".format(name)
        output_path_html = folder_path_gas + "sort_index_{:s}.html".format(name)
        table_format(excelpath=output_path, data=df_index, name='sheet1', title=name,
                     **_photo_gap(1761, 'save_subset_to_new_sheet and remaining arguments'))
        png_path = folder_path_gas + 'sort_index_{:s}.png'.format(name)
        c = 0
        while c < 100:
            try:
                excel2img.export_img(output_path, png_path, "", "sheet1!A1:AD{:d}".format(
                    len(df_index) + _photo_gap(1767, 'row extent after len(df_index)')))
                xlsx2html(output_path, output_path_html, sheet="sheet2",
                          default_cell_border=_photo_gap(1768, '1px solid... and any remaining arguments'))
                break
            except:
                c += 1
                pass
    else:
        try:
            os.remove(folder_path_gas + "sort_signal_{:s}.xlsx".format(name))
            os.remove(folder_path_gas + 'sort_signal_{:s}.png'.format(name))
            os.remove(folder_path_gas + 'sort_signal_{:s}.html'.format(name))
            os.remove(folder_path_gas + "sort_index_{:s}.xlsx".format(name))
            os.remove(folder_path_gas + 'sort_index_{:s}.png'.format(name))
            os.remove(folder_path_gas + 'sort_index_{:s}.html'.format(name))
        except:
            pass


def send_trend_signal(asset_class_list, send_to):
    body = [table.html_text('Lastest update time is {:s}'.format(now_ldn().strftime('%Y-%m-%d %H:%M'))),
            *_photo_gap(1788, 'remaining initial email body entries after timestamp')]
    for i in asset_class_list:
        if os.path.exists('{:s}sort_signal_{:s}.html'.format(folder_path_gas, i)):
            with open(ut.convert_path_to_linux('{:s}sort_signal_{:s}.html'.format(folder_path_gas, i)), 'r') as f:
                html_lines = f.readlines()
            body.append(''.join(html_lines) + '</br>')
    if len(body) > 1:
        print('Sending Trend Signal Gas')
        send_email(send_to, subject='MarketScan: Trend Signal GAS', body=body)
    else:
        print('No trend signal for GAS')


def send_all_tables(send_to, asset_class=None):
    asset_class_dict = {
        'EU_GAS': ['EU_GAS'],
        'EU_POWER': ['EU_POWER'],
        'COAL': ['COAL'],
        'NG_STOCKS': ['NG_STOCKS'],
    }
    if asset_class is None:
        asset_class = list(asset_class_dict.keys())
    for key, val in asset_class_dict.items():
        save_output(key, val)
    send_trend_signal(asset_class, send_to=send_to)
    figs = []
    for i in list(asset_class_dict.keys()):
        ng_signal_path = ut.convert_path_to_linux(f"{folder_path_gas}sort_signal_{i}.html")
        if os.path.exists(ng_signal_path):
            with open(ng_signal_path, 'r') as f:
                html_lines = f.readlines()
            figs.append("<div style='font-family:Calibri;' >")
            figs.append('Lastest update time is {:s} </br>'.format(now_ldn().strftime('%Y-%m-%d %H:%M')))
            figs.append(f'{i.upper()} Signal</br>')
            figs.append(''.join(html_lines))
            figs.append('</br>')
    table.figures_to_html(figs, filename=ut.convert_path_to_linux(f"{folder_path}ng_signal.html"))


def send_sharpe_ratio_rank_gas(send_to):
    sr = pd.read_csv(ut.convert_path_to_linux(f"{folder_path}gas\\sharpe_ratio_gas_power.csv"))
    sr.set_index('Unnamed: 0', inplace=True)
    sr.dropna(axis=0, inplace=True)
    sr['abs'] = np.abs(sr['Shrp 3D COB'])
    sr = sr.sort_values('abs', ascending=False)
    sr = sr.loc[sr['abs'] > 1, :]
    sr['Ticker'] = [x for x in sr['_plot_link']]
    sr.drop(columns=['_plot_link'])

    def get_rank(data):
        data_ = np.abs(data)
        rank_ = data_.rank(ascending=False)
        return rank_

    def table_format(df, header):
        return table.html_format(
            df=df,
            header=header,
            hide_cols=['abs', 'new'],
            format_column={
                'Instr': {'width': '40px', 'text-align': 'center'},
                'Ticker': {'width': '120px', 'text-align': 'center'},
                'Shrp 3D Live': {'width': '80px', 'text-align': 'center'},
                'Shrp 3D COB': {'width': '80px', 'text-align': 'center'},
                'Shrp 22D': {'width': '80px', 'text-align': 'center'},
                'Signal': {'width': '80px', 'text-align': 'center'},
                'new': {'highlight': ['Ticker', 'new']},
            },
        )

    sr_us_ = sr.loc[sr['Instr'].isin(['US_GAS']), :]
    sr_us_flat_ = sr_us_.loc[sr['Type'].isin(['FLAT']), :]
    sr_us_flat_.reset_index(drop=True, inplace=True)
    sr_us_flat_['rank1'] = get_rank(sr_us_flat_['Shrp 3D COB 1'])
    sr_us_flat_['rank2'] = get_rank(sr_us_flat_['Shrp 3D COB 2'])
    sr_us_flat_['rank3'] = get_rank(sr_us_flat_['Shrp 3D COB 3'])
    sr_us_flat = sr_us_flat_.iloc[:10, :]
    sr_us_flat['new'] = 0
    for idx, row in sr_us_flat.iterrows():
        if (((row['rank1'] > 10) or (row['rank1'] <= 10 and np.sign(row['Shrp 3D COB']) != np.sign(row['Shrp 3D COB 1']))) and
            ((row['rank2'] > 10) or (row['rank2'] <= 10 and np.sign(row['Shrp 3D COB']) != np.sign(row['Shrp 3D COB 2']))) and
            ((row['rank3'] > 10) or (row['rank3'] <= 10 and np.sign(row['Shrp 3D COB']) != np.sign(row['Shrp 3D COB 3'])))):
            sr_us_flat.loc[idx, 'new'] = 1
    sr_us_flat = sr_us_flat[['Instr', 'Ticker', 'Shrp 3D Live', 'Shrp 3D COB', 'Shrp 22D', 'Signal', 'new']]
    figs = []
    figs.append("<div style='font-family:Calibri;' >")
    figs.append('<br>')
    figs.append('Green ticker is new signal.')
    figs.append('<br>')
    if len(sr_us_flat) > 0:
        sr_us_html = table_format(sr_us_flat, 'US GAS Flats 3d COB Sharpe Ratio Alert')
        figs.append(sr_us_html)
        figs.append('<br>')
    else:
        figs.append('No US GAS Flats 3d COB Sharpe Ratio Alert')
        figs.append('<br>')
    sr_us_spread_ = sr_us_.loc[sr['Type'].isin(['SPRD']), :]
    sr_us_spread_.reset_index(drop=True, inplace=True)
    sr_us_spread_['rank1'] = get_rank(sr_us_spread_['Shrp 3D COB 1'])
    sr_us_spread_['rank2'] = get_rank(sr_us_spread_['Shrp 3D COB 2'])
    sr_us_spread_['rank3'] = get_rank(sr_us_spread_['Shrp 3D COB 3'])
    sr_us_spread = sr_us_spread_.iloc[:10, :]
    sr_us_spread['new'] = 0
    for idx, row in sr_us_spread.iterrows():
        if (((row['rank1'] > 10) or (row['rank1'] <= 10 and np.sign(row['Shrp 3D COB']) != np.sign(row['Shrp 3D COB 1']))) and
            ((row['rank2'] > 10) or (row['rank2'] <= 10 and np.sign(row['Shrp 3D COB']) != np.sign(row['Shrp 3D COB 2']))) and
            ((row['rank3'] > 10) or (row['rank3'] <= 10 and np.sign(row['Shrp 3D COB']) != np.sign(row['Shrp 3D COB 3'])))):
            sr_us_flat.loc[idx, 'new'] = 1
    sr_us_spread = sr_us_spread[['Instr', 'Ticker', 'Shrp 3D Live', 'Shrp 3D COB', 'Shrp 22D', 'Signal', 'new']]
    if len(sr_us_spread) > 0:
        sr_us_html = table_format(sr_us_spread, 'US GAS Spreads 3d COB Sharpe Ratio')
        figs.append(sr_us_html)
        figs.append('<br>')
    else:
        figs.append('No US GAS Spreads 3d COB Sharpe Ratio')
        figs.append('<br>')

    sr_eu_ = sr.loc[sr['Instr'].isin(['EU_GAS']), :]
    sr_eu_flat_ = sr_eu_.loc[sr['Type'].isin(['FLAT']), :]
    sr_eu_flat_.reset_index(drop=True, inplace=True)
    sr_eu_flat_['rank1'] = get_rank(sr_eu_flat_['Shrp 3D COB 1'])
    sr_eu_flat_['rank2'] = get_rank(sr_eu_flat_['Shrp 3D COB 2'])
    sr_eu_flat_['rank3'] = get_rank(sr_eu_flat_['Shrp 3D COB 3'])
    sr_eu_flat = sr_eu_flat_.iloc[:10, :]
    sr_eu_flat['new'] = 0
    for idx, row in sr_eu_flat.iterrows():
        if (((row['rank1'] > 10) or (row['rank1'] <= 10 and np.sign(row['Shrp 3D COB']) != np.sign(row['Shrp 3D COB 1']))) and
            ((row['rank2'] > 10) or (row['rank2'] <= 10 and np.sign(row['Shrp 3D COB']) != np.sign(row['Shrp 3D COB 2']))) and
            ((row['rank3'] > 10) or (row['rank3'] <= 10 and np.sign(row['Shrp 3D COB']) != np.sign(row['Shrp 3D COB 3'])))):
            sr_eu_flat.loc[idx, 'new'] = 1
    sr_eu_flat = sr_eu_flat[['Instr', 'Ticker', 'Shrp 3D Live', 'Shrp 3D COB', 'Shrp 22D', 'Signal', 'new']]
    if len(sr_eu_flat) > 0:
        sr_eu_html = table_format(sr_eu_flat, 'EU GAS Flats 3d COB Sharpe Ratio')
        figs.append(sr_eu_html)
        figs.append('<br>')
    else:
        figs.append('No EU GAS Flats 3d COB Sharpe Ratio')
        figs.append('<br>')
    sr_eu_spreads_ = sr_eu_.loc[sr['Type'].isin(['SPRD']), :]
    sr_eu_spreads_.reset_index(drop=True, inplace=True)
    sr_eu_spreads_['rank1'] = get_rank(sr_eu_spreads_['Shrp 3D COB 1'])
    sr_eu_spreads_['rank2'] = get_rank(sr_eu_spreads_['Shrp 3D COB 2'])
    sr_eu_spreads_['rank3'] = get_rank(sr_eu_spreads_['Shrp 3D COB 3'])
    sr_eu_sprd = sr_eu_spreads_.iloc[:10, :]
    sr_eu_sprd['new'] = 0
    for idx, row in sr_eu_sprd.iterrows():
        if (((row['rank1'] > 10) or (row['rank1'] <= 10 and np.sign(row['Shrp 3D COB']) != np.sign(row['Shrp 3D COB 1']))) and
            ((row['rank2'] > 10) or (row['rank2'] <= 10 and np.sign(row['Shrp 3D COB']) != np.sign(row['Shrp 3D COB 2']))) and
            ((row['rank3'] > 10) or (row['rank3'] <= 10 and np.sign(row['Shrp 3D COB']) != np.sign(row['Shrp 3D COB 3'])))):
            sr_eu_sprd.loc[idx, 'new'] = 1
    sr_eu_sprd = sr_eu_sprd[['Instr', 'Ticker', 'Shrp 3D Live', 'Shrp 3D COB', 'Shrp 22D', 'Signal', 'new']]
    if len(sr_eu_sprd) > 0:
        sr_eu_html = table_format(sr_eu_sprd, 'EU GAS Spreads 3d COB Sharpe Ratio')
        figs.append(sr_eu_html)
        figs.append('<br>')
    else:
        figs.append('No EU GAS Spreads 3d COB Sharpe Ratio')
        figs.append('<br>')

    sr_other_ = sr.loc[sr['Instr'].isin(['EU_POWER', 'COAL', 'NG_STOCKS', 'OIL']), :]
    sr_other_flat_ = sr_other_.loc[sr['Type'].isin(['FLAT']), :]
    sr_other_flat_.reset_index(drop=True, inplace=True)
    sr_other_flat_['rank1'] = get_rank(sr_other_flat_['Shrp 3D COB 1'])
    sr_other_flat_['rank2'] = get_rank(sr_other_flat_['Shrp 3D COB 2'])
    sr_other_flat_['rank3'] = get_rank(sr_other_flat_['Shrp 3D COB 3'])
    sr_other_flat = sr_other_flat_.iloc[:10, :]
    sr_other_flat['new'] = 0
    for idx, row in sr_other_flat.iterrows():
        if (((row['rank1'] > 10) or (row['rank1'] <= 10 and np.sign(row['Shrp 3D COB']) != np.sign(row['Shrp 3D COB 1']))) and
            ((row['rank2'] > 10) or (row['rank2'] <= 10 and np.sign(row['Shrp 3D COB']) != np.sign(row['Shrp 3D COB 2']))) and
            ((row['rank3'] > 10) or (row['rank3'] <= 10 and np.sign(row['Shrp 3D COB']) != np.sign(row['Shrp 3D COB 3'])))):
            sr_other_flat.loc[idx, 'new'] = 1
    sr_other_flat = sr_other_flat[['Instr', 'Ticker', 'Shrp 3D Live', 'Shrp 3D COB', 'Shrp 22D', 'Signal', 'new']]
    if len(sr_other_flat) > 0:
        sr_other_html = table_format(sr_other_flat, 'Other Flats 3d COB Sharpe Ratio')
        figs.append(sr_other_html)
        figs.append('<br>')
    else:
        figs.append('No Other Flats 3d COB Sharpe Ratio')
        figs.append('<br>')
    sr_other_sprd_ = sr_other_.loc[sr['Type'].isin(['SPRD']), :]
    sr_other_sprd_.reset_index(drop=True, inplace=True)
    sr_other_sprd_['rank1'] = get_rank(sr_other_sprd_['Shrp 3D COB 1'])
    sr_other_sprd_['rank2'] = get_rank(sr_other_sprd_['Shrp 3D COB 2'])
    sr_other_sprd_['rank3'] = get_rank(sr_other_sprd_['Shrp 3D COB 3'])
    sr_other_sprd = sr_other_sprd_.iloc[:10, :]
    sr_other_sprd['new'] = 0
    for idx, row in sr_other_sprd.iterrows():
        if (((row['rank1'] > 10) or (row['rank1'] <= 10 and np.sign(row['Shrp 3D COB']) != np.sign(row['Shrp 3D COB 1']))) and
            ((row['rank2'] > 10) or (row['rank2'] <= 10 and np.sign(row['Shrp 3D COB']) != np.sign(row['Shrp 3D COB 2']))) and
            ((row['rank3'] > 10) or (row['rank3'] <= 10 and np.sign(row['Shrp 3D COB']) != np.sign(row['Shrp 3D COB 3'])))):
            sr_other_sprd.loc[idx, 'new'] = 1
    sr_other_sprd = sr_other_sprd[['Instr', 'Ticker', 'Shrp 3D Live', 'Shrp 3D COB', 'Shrp 22D', 'Signal', 'new']]
    if len(sr_other_sprd) > 0:
        sr_other_html = table_format(sr_other_sprd, 'Other Spreads 3d COB Sharpe Ratio')
        figs.append(sr_other_html)
        figs.append('<br>')
    else:
        figs.append('No Other Spreads 3d COB Sharpe Ratio')
        figs.append('<br>')
    table.figures_to_html(figs, ut.convert_path_to_linux(_photo_gap(2010, f"{html_path}\\gas\\cross_cmds\\sharpe_ratio_g...")))
    send_email(send_to=send_to, subject='Sharpe Ratio Rank Gas Power', body=figs,
               html_path=_photo_gap(2011, f"{html_path}..."))


def update():
    market_scan()
    range_vol_am.send_alert_email_gas(send_to=send_to)
    send_sharpe_ratio_rank_gas(send_to=send_to)


if __name__ == '__main__':
    update()
