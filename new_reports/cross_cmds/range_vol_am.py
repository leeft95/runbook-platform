import pandas as pd
import numpy as np
import datetime as dt
import sys
import os
import getpass
from dateutil.relativedelta import relativedelta, FR, SA
from dateutil.rrule import rrule, MONTHLY, WE
from pandas.tseries.offsets import BDay
import plotly.graph_objects as go
import plotly as py
from plotly.subplots import make_subplots
from functools import partial
import PyPDF2
import ecm.cmds.table as table
import ecm.cmds.chart as chart
if sys.platform.startswith("win"):
    import excel2img
import ecm.cmds.config as config
import ecm.cmds.bbg as bbg
import ecm.cmds.pyg as pyg
import ecm.cmds.sql as sql
import ecm.cmds.talib as talib
import ecm.cmds.ticker as tk
import ecm.cmds.utils as ut
import ecm.cmds.time_series as ts
import ecm.cmds.data as dv
import ecm.cmds.market_scan as ms
from ecm.data.api import RTHQueryClient
from ecm.cmds.config import root_path, output_path, html_path, url, oil_group, gas_group
from ecm.cmds._email import send_email
from ecm.cmds.cdr import today, getworkingdays, CDR, month_str2int, now_ldn
from ecm.cmds.utils import convert_path_to_linux
from pyg_mongo import *
from loguru import logger as log
from pandas.tseries.holiday import USFederalHolidayCalendar
from copy import copy
import openpyxl
import kaleido

from ecm.data.api import RTHQueryClient
from multiprocessing import Pool
from collections import ChainMap
import itertools
from ecm.atom.clients import ECMHttpError
from market_scan_helpers import MarketDataLoader
from market_scan_helpers import (
    fom_keys,
    name_dict,
    month_ticker_dict,
    month_ticker_dict1,
    diverg_list,
    option_dict,
    spots_dict,
    other_platts_dict
)
from market_scan_helpers.processing_utils import process_ticker_data
from options_volume import vol_alert
from options_volume import get_vol
import getpass

user = getpass.getuser()

bday_us = pd.offsets.CustomBusinessDay(calendar=USFederalHolidayCalendar())
send_to = list(set(oil_group + gas_group))
report_name = "Range-Vol and Divergence AM"
file_name = "range_vol_am"  # without .py
file_path = f"{root_path}\\autoreports\\reports\\cross_cmds\\{file_name}.py"
cc_csv_folder = f"{output_path}\\csvs\\cross_cmds"
cc_json_folder = f"{output_path}\\json\\cross_cmds"
cc_pdf_folder = f"{output_path}\\pdf\\cross_cmds"
folder_path = f"{cc_csv_folder}\\market_scan\\"

for file in [cc_csv_folder, cc_json_folder, cc_pdf_folder, folder_path, html_path]:
    if not os.path.exists(file):
        os.makedirs(file, exist_ok=True)


market_data_loader = MarketDataLoader(cache_dir=folder_path, use_cache=True)

risk_index_def_link = table.html_text(
    "Risk Index Definition",
    style="font-size: 16px; text-align:left; font-family:Calibri;",
    tag="div"
)


def get_tickers(name_dict=None):
    """Get ticker dictionary using MarketDataLoader"""
    return market_data_loader.get_tickers(name_dict, force_update=False)


def refresh_data(key, val, folder_path, use_bbg=True):
    """
    Refresh ticker data using MarketDataLoader

    Wrapper around MarketDataLoader.refresh_ticker_data() to maintain
    compatibility with existing code that expects {key: dict(...)} format
    """
    sdate = today() - dt.timedelta(days=35)
    edate = dt.datetime.now()

    ticker, instr, chg_type, enabled, has_vol = val

    if ticker is None:
        log.error(f"{key} has no ticker mapped!")
        return {key: dict(daily=pd.DataFrame(), weekly=pd.DataFrame(), monthly=pd.DataFrame(), intraday=pd.DataFrame())}
    if not enabled:
        log.warning(f"Skipping : {ticker}")
        return {key: dict(daily=pd.DataFrame(), weekly=pd.DataFrame(), monthly=pd.DataFrame(), intraday=pd.DataFrame())}

    result = market_data_loader.refresh_ticker_data(
        ticker=ticker,
        start_date=sdate,
        end_date=edate,
        folder_path=folder_path,
        use_bbg=use_bbg
    )

    return {key: dict(daily=result['daily'], weekly=result['weekly'], monthly=result['monthly'], intraday=result['intraday'])}


_sheet_order = [
    "RISK",
    "EQT",
    "FI",
    "CCY",
    "VOL",
    "PM",
    "BRT",
    "WTI",
    "DIST",
    "GASOLINE",
    "GAS",
    "HH",
    "TTF",
    "BM",
    "Bulks",
    "GR",
    "SFT",
]


def get_basic_info(ticker_dict, use_bbg=True):
    output_list = []
    edate = dt.datetime.now()
    current_fom = (today().replace(day=1) + 0 * bday_us).strftime("%Y-%m-%d")

    fom_path = ut.convert_path_to_linux("\\\\elementcapital.corp\\ecns01\\PM\\Michel Kikano\\Market Scan\\FOM.csv")
    fom_comments = pd.read_csv(fom_path)
    fom_comments.set_index("Unnamed: 0", inplace=True)
    fom_comments_ = fom_comments.copy()
    fom_comments_["days"] = 0

    missing_keys = [x for x in fom_comments.index if x not in fom_keys]
    if missing_keys:
        for k in missing_keys:
            log.warning(f"FOM Key {k} missing using default {current_fom}")
            fom_comments[k] = current_fom
    fom_dict = {}
    for idx, row in fom_comments_.iterrows():
        if isinstance(row["FOM Dates"], str):
            if len(row["FOM Dates"]) == 10:
                try:
                    fom_d0 = dt.datetime.strptime(row["FOM Dates"][:10], "%Y-%m-%d")
                except:
                    fom_d0 = dt.datetime.strptime(row["FOM Dates"][:10], "%m/%d/%Y")
                fom_comments_.loc[idx, "days"] = int(getworkingdays(fom_d0, today()))
                fom_comments_.loc[idx, "FOM Dates"] = str(dt.datetime.strftime(fom_d0, "%d/%b"))
                fom_d1 = today()
            else:
                try:
                    fom_d0 = dt.datetime.strptime(row["FOM Dates"][:10], "%Y-%m-%d")
                except:
                    fom_d0 = dt.datetime.strptime(row["FOM Dates"][:10], "%m/%d/%Y")
                try:
                    fom_d1 = dt.datetime.strptime(row["FOM Dates"][-10:], "%Y-%m-%d")
                except:
                    fom_d1 = dt.datetime.strptime(row["FOM Dates"][-10:], "%m/%d/%Y")
                fom_comments_.loc[idx, "days"] = int(getworkingdays(fom_d1, today()))
                fom_comments_.loc[idx, "FOM Dates"] = dt.datetime.strftime(fom_d0, "%d/%b") + " - " + dt.datetime.strftime(fom_d1, "%d/%b")
            if fom_d0 > today():
                fom_d0 = dt.datetime(today().year, today().month, 1)
        elif isinstance(row["FOM Dates"], dt.datetime):
            fom_comments_.loc[idx, "days"] = int(getworkingdays(row["FOM Dates"], today()))
            fom_comments_.loc[idx, "FOM Dates"] = str(dt.datetime.strftime(row["FOM Dates"], "%d/%b"))
        fom_dict[idx] = dict(sdate=fom_d0, edate=fom_d1)
    comments_path = ut.convert_path_to_linux(f"{folder_path}\\fom_comments.csv")
    try:
        fom_comments_.to_csv(comments_path)
    except Exception as e:
        log.error(f"Could not save {comments_path} due to \n{e}")

    alert_email_all = []
    alert_email_oil = []
    alert_email_ng = []
    alert_email_macro = []
    intraday_cum_volume = []
    alert_email_all_today = []
    alert_email_oil_today = []
    alert_email_ng_today = []
    alert_email_macro_today = []
    alert_bm = []
    alert_bm_pm = []
    sharpe_ratio_list = []

    weekly_high = []
    weekly_low = []
    week_dates = pd.bdate_range(today() - dt.timedelta(days=7), today())
    for i in range(len(week_dates) - 2, -1, -1):
        if week_dates[i].weekday() > week_dates[i + 1].weekday():
            week_sdate = week_dates[i + 1]
    log.info("Running Market Scan")
    log.info("Refreshing data")
    args = [(key[1], val, folder_path, use_bbg) for key, val in ticker_dict.items()]
    vwap_track_date_range = dict(sdate=fom_d0, edate=fom_d1)

    with Pool(10) as pool:
        latest_data_list = pool.starmap(refresh_data, args)
        latest_data = dict(ChainMap(*latest_data_list))
        args_process = [
            (
                key,
                val,
                latest_data.get(key[1], dict()),
                fom_comments,
                fom_dict,
                week_sdate,
                folder_path,
            )
            for key, val in ticker_dict.items()
        ]
        outputs = pool.starmap(process_ticker_data, args_process)

    for (
        output_list_,
        alert_email_all_,
        alert_email_oil_,
        alert_email_ng_,
        alert_email_macro_,
        alert_email_all_today_,
        alert_email_oil_today_,
        alert_email_ng_today_,
        alert_email_macro_today_,
        alert_bm_,
        alert_bm_pm_,
        sharpe_ratio_list_,
        weekly_high_,
        weekly_low_,
        intraday_cum_volume_,
    ) in outputs:
        output_list.extend(output_list_)
        alert_email_all.extend(alert_email_all_)
        alert_email_oil.extend(alert_email_oil_)
        alert_email_ng.extend(alert_email_ng_)
        alert_email_macro.extend(alert_email_macro_)
        alert_email_all_today.extend(alert_email_all_today_)
        alert_email_oil_today.extend(alert_email_oil_today_)
        alert_email_ng_today.extend(alert_email_ng_today_)
        alert_email_macro_today.extend(alert_email_macro_today_)
        alert_bm.extend(alert_bm_)
        alert_bm_pm.extend(alert_bm_pm_)
        sharpe_ratio_list.extend(sharpe_ratio_list_)
        weekly_low.extend(weekly_low_)
        weekly_high.extend(weekly_high_)
        intraday_cum_volume.extend(intraday_cum_volume_)
    df = pd.DataFrame(
        output_list,
        columns=[
            "Instr",
            "Ticker",
            "Generic",
            "Expiry",
            "Last Price",
            "KPI",
            "KPI t-1",
            "Change",
            "Type",
            "FOM",
            "Ref Date",
            "Stochs",
            "Post Close",
            "Shrp 3",
            "Shrp 22",
            "M",
            "H",
            "F",
            "LT",
            "CT",
            "Mom",
            "Hybrid",
            "Fish",
            "FishLT",
            "Counter",
            "KPI Index",
            "Index",
            "3Day",
            "RngSpike",
            "VolSpike",
            "Volspike20d",
            "Percentile",
            "Signal",
            "Mom_ytsd",
            "Hybrid_ytsd",
            "Fish_ytsd",
            "FishLT_ytsd",
            "Counter_ytsd",
            "Change 1d",
            "Change 5d",
            "IROLL_tag",
            "FOM PxDelta",
            "FOM Vwap",
            "FOM zscore",
        ],
    )
    df = pd.concat(
        [
            df[df["Instr"] != "RISK"].set_index("Generic").reindex(name_dict.keys()),
            df[(df["Instr"] == "RISK")].set_index("Generic").reindex(risk_sharpe_ticker_weights.keys()),
        ]
    ).reset_index(drop=True)
    df = df[~df.Instr.isna()]
    df.index.name = now_ldn().strftime("%Y-%m-%d %H:%M")

    df_all = pd.DataFrame(
        alert_email_all,
        columns=[
            "Instr",
            "Ticker",
            "Follow Through",
            "Change T",
            "Range T",
            "Volume T",
            "Change T-1",
            "Range T-1",
            "Volume T-1",
            "chg type",
            "Shrp 3D COB",
        ],
    )
    df_oil = pd.DataFrame(
        alert_email_oil,
        columns=[
            "Instr",
            "Ticker",
            "Follow Through",
            "Change T",
            "Range T",
            "Volume T",
            "Change T-1",
            "Range T-1",
            "Volume T-1",
            "chg type",
            "Shrp 3D COB",
        ],
    )
    df_ng = pd.DataFrame(
        alert_email_ng,
        columns=[
            "Instr",
            "Ticker",
            "Follow Through",
            "Change T",
            "Range T",
            "Volume T",
            "Change T-1",
            "Range T-1",
            "Volume T-1",
            "chg type",
            "Shrp 3D COB",
        ],
    )
    df_macro = pd.DataFrame(
        alert_email_macro,
        columns=[
            "Instr",
            "Ticker",
            "Follow Through",
            "Change T",
            "Range T",
            "Volume T",
            "Change T-1",
            "Range T-1",
            "Volume T-1",
            "chg type",
            "Shrp 3D COB",
        ],
    )
    df_intraday_vol = pd.DataFrame(intraday_cum_volume, columns=["Instr", "Ticker", "Volume T", "Change T", "chg type"])
    df_all_today = pd.DataFrame(
        alert_email_all_today,
        columns=[
            "Instr",
            "Ticker",
            "Follow Through",
            "Change T",
            "Range T",
            "Volume T",
            "Change T-1",
            "Range T-1",
            "Volume T-1",
            "chg type",
            "Shrp 3D COB",
        ],
    )
    df_oil_today = pd.DataFrame(
        alert_email_oil_today,
        columns=[
            "Instr",
            "Ticker",
            "Follow Through",
            "Change T",
            "Range T",
            "Volume T",
            "Change T-1",
            "Range T-1",
            "Volume T-1",
            "chg type",
            "Shrp 3D COB",
        ],
    )
    df_ng_today = pd.DataFrame(
        alert_email_ng_today,
        columns=[
            "Instr",
            "Ticker",
            "Follow Through",
            "Change T",
            "Range T",
            "Volume T",
            "Change T-1",
            "Range T-1",
            "Volume T-1",
            "chg type",
            "Shrp 3D COB",
        ],
    )
    df_macro_today = pd.DataFrame(
        alert_email_macro_today,
        columns=[
            "Instr",
            "Ticker",
            "Follow Through",
            "Change T",
            "Range T",
            "Volume T",
            "Change T-1",
            "Range T-1",
            "Volume T-1",
            "chg type",
            "Shrp 3D COB",
        ],
    )

    weekly_high_df = pd.DataFrame(
        weekly_high,
        columns=[
            "Instr",
            "Ticker",
            "Live",
            "Weekly High",
            "Change Today",
            "Range Today",
            "Volume Today",
        ],
    )
    weekly_low_df = pd.DataFrame(
        weekly_low,
        columns=[
            "Instr",
            "Ticker",
            "Live",
            "Weekly Low",
            "Change Today",
            "Range Today",
            "Volume Today",
        ],
    )

    alert_bm_df = pd.DataFrame(
        alert_bm,
        columns=[
            "Ticker",
            "Price change",
            "Price",
            "Price 20d MA",
            "20d z-score",
            "Date",
            "chg type",
        ],
    )
    alert_bm_pm_df = pd.DataFrame(
        alert_bm_pm,
        columns=[
            "Ticker",
            "Price change",
            "Price",
            "Price 20d MA",
            "20d z-score",
            "Date",
            "chg type",
        ],
    )

    sharpe_ratio_df = pd.DataFrame(
        sharpe_ratio_list,
        columns=[
            "Instr",
            "Ticker",
            "Type",
            "Shrp 3D Live",
            "Shrp 3D COB",
            "Shrp 22D",
            "Signal",
            "Shrp 3D COB 1",
            "Shrp 3D COB 2",
            "Shrp 3D COB 3",
            "KPI",
            "KPI t-1",
            "KPI Index",
            "Index",
            "_plot_link",
        ],
    )

    return (
        df,
        df_all,
        df_oil,
        df_ng,
        df_macro,
        weekly_high_df,
        weekly_low_df,
        df_intraday_vol,
        df_all_today,
        df_oil_today,
        df_ng_today,
        df_macro_today,
        alert_bm_df,
        alert_bm_pm_df,
        sharpe_ratio_df,
    )


def market_scan():
    log.info("Starting Market Scan")
    ut.tic()
    ticker_dict = get_tickers()
    (
        df,
        df_all,
        df_oil,
        df_ng,
        df_macro,
        weekly_high_df,
        weekly_low_df,
        df_intraday_vol,
        df_all_today,
        df_oil_today,
        df_ng_today,
        df_macro_today,
        df_bm_am,
        df_bm_pm,
        sharpe_ratio,
    ) = get_basic_info(ticker_dict, use_bbg=True)

    df.to_csv(ut.convert_path_to_linux(f"{folder_path}trend_signal.csv"))
    df_all.to_csv(ut.convert_path_to_linux(f"{folder_path}trend_alert_all.csv"))
    df_oil.to_csv(ut.convert_path_to_linux(f"{folder_path}trend_alert_oil.csv"))
    df_ng.to_csv(ut.convert_path_to_linux(f"{folder_path}trend_alert_ng.csv"))
    df_macro.to_csv(ut.convert_path_to_linux(f"{folder_path}trend_alert_macro.csv"))
    df_all_today.to_csv(ut.convert_path_to_linux(f"{folder_path}trend_alert_all_today.csv"))
    df_oil_today.to_csv(ut.convert_path_to_linux(f"{folder_path}trend_alert_oil_today.csv"))
    df_ng_today.to_csv(ut.convert_path_to_linux(f"{folder_path}trend_alert_ng_today.csv"))
    df_macro_today.to_csv(ut.convert_path_to_linux(f"{folder_path}trend_alert_macro_today.csv"))
    weekly_high_df.to_csv(ut.convert_path_to_linux(f"{folder_path}weekly_high.csv"))
    weekly_low_df.to_csv(ut.convert_path_to_linux(f"{folder_path}weekly_low.csv"))
    df_intraday_vol.to_csv(ut.convert_path_to_linux(f"{folder_path}intraday_cum_vol.csv"))
    df_bm_am.to_csv(ut.convert_path_to_linux(f"{folder_path}bm_alert_am.csv"))
    df_bm_pm.to_csv(ut.convert_path_to_linux(f"{folder_path}bm_alert_pm.csv"))
    sharpe_ratio.to_csv(ut.convert_path_to_linux(f"{folder_path}sharpe_ratio.csv"))
    diverge_ticker_dict = {x[1]: y for x, y in ticker_dict.items() if x[1] in list(itertools.chain.from_iterable(diverg_list))}
    df_diverge = get_diverge(diverge_ticker_dict)
    df_diverge.to_csv(ut.convert_path_to_linux(f"{folder_path}diverge_signal.csv"))
    df_fly, fly_charts = get_fly_switch(ticker_dict)
    df_fly.to_csv(ut.convert_path_to_linux(f"{folder_path}fly_switch_signal.csv"))
    if len(fly_charts) > 0:
        table.figures_to_html(fly_charts, convert_path_to_linux(f"{folder_path}charts\\fly_switch_signal.html"))
    ut.toc()
    log.info("Market Scan Complete")


def get_diverge(ticker_dict):
    log.info(f"Running Get Diverge")
    index_list = [f"{x[0]} / {x[1]}" for x in diverg_list]
    df = pd.DataFrame(np.nan, index=index_list, columns=["Ticker", "3d divergence", "8d divergence"])
    for idx, i in enumerate(diverg_list):
        ticker1 = ticker_dict[i[1]][0]
        price1 = ts.read_csv(f"{folder_path}price_daily\\{ticker1}.csv", index_name="date")
        if i[0] == "JKM/TTF":
            ticker0 = f"JKMTTF_{(today()+dt.timedelta(50)).strftime('%b%y')} Comdty"
            price0 = dv.ce_price(id=f"JKMTTFusd_{(today()+dt.timedelta(50)).strftime('%b.%y')}", nd=120)
            price0.columns = ["PX_LAST"]
            price0_chg = -price0["PX_LAST"].diff()
            price1 = price1.reindex(price0.index)
            price1.fillna(method="ffill", inplace=True)
        else:
            ticker0 = ticker_dict[i[0]][0]
            price0 = ts.read_csv(f"{folder_path}price_daily\\{ticker0}.csv", index_name="date")
            price0_chg = price0["PX_LAST"].diff()
        price1_chg = price1["PX_LAST"].diff()
        df.loc[index_list[idx], "Ticker"] = f"{ticker0.split(' ')[0]}/{ticker1.split(' ')[0]}"
        if price0_chg.index[-1] >= today():
            df.loc[index_list[idx], "Sprd change"] = price0_chg[-2]
            df.loc[index_list[idx], "Flat change"] = price1_chg[-2]
            if ("S:XBXB" in i[0] or "S:HOHO" in i[0]) and "S:COCO" in i[1]:
                df.loc[index_list[idx], "3d divergence"] = price0_chg.rolling(3).sum()[-2] * 0.42 / price1_chg.rolling(3).sum()[-2]
                df.loc[index_list[idx], "8d divergence"] = price0_chg.rolling(8).sum()[-2] * 0.42 / price1_chg.rolling(8).sum()[-2]
            else:
                if i[0] == "JKM/TTF":
                    if abs(price1_chg[-2]) / price1_chg[-23:-2].std() > 1.5:
                        valid = True
                    else:
                        valid = False
                else:
                    valid = True

                if valid:
                    df.loc[index_list[idx], "3d divergence"] = price0_chg.rolling(3).sum()[-2] / price1_chg.rolling(3).sum()[-2]
                    df.loc[index_list[idx], "8d divergence"] = price0_chg.rolling(8).sum()[-2] / price1_chg.rolling(8).sum()[-2]
        else:
            df.loc[index_list[idx], "Sprd change"] = price0_chg[-1]
            df.loc[index_list[idx], "Flat change"] = price1_chg[-1]
            if ("S:XBXB" in i[0] or "S:HOHO" in i[0]) and "S:COCO" in i[1]:
                df.loc[index_list[idx], "3d divergence"] = price0_chg.rolling(3).sum()[-1] * 0.42 / price1_chg.rolling(3).sum()[-1]
                df.loc[index_list[idx], "8d divergence"] = price0_chg.rolling(8).sum()[-1] * 0.42 / price1_chg.rolling(8).sum()[-1]
            else:
                if i[0] == "JKM/TTF":
                    if abs(price1_chg[-1]) / price1_chg[-22:].std() > 1.5:
                        valid = True
                    else:
                        valid = False
                else:
                    valid = True

                if valid:
                    df.loc[index_list[idx], "3d divergence"] = price0_chg.rolling(3).sum()[-1] / price1_chg.rolling(3).sum()[-1]
                    df.loc[index_list[idx], "8d divergence"] = price0_chg.rolling(8).sum()[-1] / price1_chg.rolling(8).sum()[-1]
    df = df[(df["3d divergence"] < 0) | (df["8d divergence"] < 0)]
    return df


def get_fly_switch(ticker_dict):
    log.info(f"Running Get Fly List")
    fly_list = [
        ("CO 1st Spread", "CO 2nd Spread", "CO 3rd Spread"),
        ("CL 1st Spread", "CL 2nd Spread", "CL 3rd Spread"),
        ("XB 1st Spread", "XB 2nd Spread", "XB 3rd Spread"),
        ("HO 1st Spread", "HO 2nd Spread", "HO 3rd Spread"),
        ("QS 1st Spread", "QS 2nd Spread", "QS 3rd Spread"),
        ("NG 1st Spread", "NG 2nd Spread", "NG 3rd Spread"),
        ("TZT 1st spread", "TZT 2nd spread", "TZT 3rd spread"),
    ]
    fly_ticker_dict = {x[1]: y for x, y in ticker_dict.items() if x[1] in list(itertools.chain.from_iterable(fly_list))}
    df = pd.DataFrame()
    figs = []
    for i in fly_list:
        log.info(f"Processing {i}")
        price0 = ts.read_csv(f"{folder_path}price_daily\\{fly_ticker_dict[i[0]][0]}.csv", index_name="date")
        price1 = ts.read_csv(f"{folder_path}price_daily\\{fly_ticker_dict[i[1]][0]}.csv", index_name="date")
        price2 = ts.read_csv(f"{folder_path}price_daily\\{fly_ticker_dict[i[2]][0]}.csv", index_name="date")
        fly1 = price0["PX_LAST"] - price1["PX_LAST"]
        fly2 = price1["PX_LAST"] - price2["PX_LAST"]
        fly1_ticker = f"{(fly_ticker_dict[i[0]][0]).split(' ')[0]}-{(fly_ticker_dict[i[1]][0]).split(' ')[0]}"
        fly2_ticker = f"{(fly_ticker_dict[i[1]][0]).split(' ')[0]}-{(fly_ticker_dict[i[2]][0]).split(' ')[0]}"
        if np.sign(fly1.iloc[-1]) != np.sign(fly1.iloc[-2]):
            df.loc[fly1_ticker, "Current Price"] = fly1.iloc[-1]
            df.loc[fly1_ticker, "Yesterday Price"] = fly1.iloc[-2]
            df.loc[fly1_ticker, "Daily Change"] = fly1.iloc[-1] - fly1.iloc[-2]
            figs.append(chart.line_chart(df=(fly1.iloc[-65:]).to_frame("Price"), title=f"{fly1_ticker}"))
        if np.sign(fly2.iloc[-1]) != np.sign(fly2.iloc[-2]):
            df.loc[fly2_ticker, "Current Price"] = fly2.iloc[-1]
            df.loc[fly2_ticker, "Yesterday Price"] = fly2.iloc[-2]
            df.loc[fly2_ticker, "Daily Change"] = fly2.iloc[-1] - fly2.iloc[-2]
            figs.append(chart.line_chart(df=(fly2.iloc[-65:]).to_frame("Price"), title=f"{fly2_ticker}"))
    return df, figs


def table_format(excelpath, data, name="sheet1", title=None, save_subset_to_new_sheet=False, subset_range=None):
    writer = pd.ExcelWriter(excelpath, engine="xlsxwriter", engine_kwargs={'options': {'nan_inf_to_errors': True}})
    df = pd.DataFrame()
    df.to_excel(writer, index=False, sheet_name=name)
    if save_subset_to_new_sheet:
        df.to_excel(writer, index=False, sheet_name="sheet2")

    workbook = writer.book
    worksheet = writer.sheets[name]

    head_format = workbook.add_format()
    head_format.set_bold()
    head_format.set_align("center")
    head_format.set_align("vcenter")
    head_format.set_bottom(1)

    head_format_border = workbook.add_format()
    head_format_border.set_bold()
    head_format_border.set_left(5)
    head_format_border.set_bottom(1)
    head_format_border.set_align("center")
    head_format_border.set_align("vcenter")

    head_format_bold_border = workbook.add_format()
    head_format_bold_border.set_bold()
    head_format_bold_border.set_right(1)
    head_format_bold_border.set_bottom(1)
    head_format_bold_border.set_align("center")
    head_format_bold_border.set_align("vcenter")

    head_format_dash_border = workbook.add_format()
    head_format_dash_border.set_bold()
    head_format_dash_border.set_right(7)
    head_format_dash_border.set_bottom(1)
    head_format_dash_border.set_align("center")
    head_format_dash_border.set_align("vcenter")


    str_format1 = workbook.add_format()
    str_format1.set_align("center")
    str_format1.set_align("vcenter")

    str_format1_bold_border = workbook.add_format()
    str_format1_bold_border.set_align("center")
    str_format1_bold_border.set_align("vcenter")
    str_format1_bold_border.set_right(1)

    str_format1_dash_border = workbook.add_format()
    str_format1_dash_border.set_align("center")
    str_format1_dash_border.set_align("vcenter")
    str_format1_dash_border.set_right(7)

    str_format_border = workbook.add_format()
    str_format_border.set_align("center")
    str_format_border.set_align("vcenter")
    str_format_border.set_left(5)

    str_format_bold = workbook.add_format()
    str_format_bold.set_align("center")
    str_format_bold.set_align("vcenter")
    str_format_bold.set_bold()


    str_format1_red_bold_boarder = workbook.add_format()
    str_format1_red_bold_boarder.set_align("center")
    str_format1_red_bold_boarder.set_align("vcenter")
    str_format1_red_bold_boarder.set_bg_color('#FFC7CE')
    str_format1_red_bold_boarder.set_font_color('#9C0006')
    str_format1_red_bold_boarder.set_right(1)

    str_format1_red_dash_boarder = workbook.add_format()
    str_format1_red_dash_boarder.set_align("center")
    str_format1_red_dash_boarder.set_align("vcenter")
    str_format1_red_dash_boarder.set_bg_color('#FFC7CE')
    str_format1_red_dash_boarder.set_font_color('#9C0006')
    str_format1_red_dash_boarder.set_right(7)


    str_format1_green_bold_boarder = workbook.add_format()
    str_format1_green_bold_boarder.set_align("center")
    str_format1_green_bold_boarder.set_align("vcenter")
    str_format1_green_bold_boarder.set_bg_color('#C6EFCE')
    str_format1_green_bold_boarder.set_font_color('#006100')
    str_format1_green_bold_boarder.set_right(1)

    str_format1_green_dash_boarder = workbook.add_format()
    str_format1_green_dash_boarder.set_align("center")
    str_format1_green_dash_boarder.set_align("vcenter")
    str_format1_green_dash_boarder.set_bg_color('#C6EFCE')
    str_format1_green_dash_boarder.set_font_color('#006100')
    str_format1_green_dash_boarder.set_right(7)


    num_format = workbook.add_format()
    num_format.set_num_format("0.000")
    num_format.set_align("center")
    num_format.set_align("vcenter")


    num_format1 = workbook.add_format()
    num_format1.set_num_format("0.00;[Red]-0.00")
    num_format1.set_align("center")
    num_format1.set_align("vcenter")
    num_format1.set_right(1)

    num_format1_dash_boarder = workbook.add_format()
    num_format1_dash_boarder.set_num_format("0.00;[Red]-0.00")
    num_format1_dash_boarder.set_align("center")
    num_format1_dash_boarder.set_align("vcenter")
    num_format1_dash_boarder.set_right(7)


    num_format1_red = workbook.add_format()
    num_format1_red.set_num_format("0.00;[Red]-0.00")
    num_format1_red.set_align("center")
    num_format1_red.set_align("vcenter")
    num_format1_red.set_bg_color('#FFC7CE')
    num_format1_red.set_bold()


    num_format1_green = workbook.add_format()
    num_format1_green.set_num_format("0.00;[Red]-0.00")
    num_format1_green.set_align("center")
    num_format1_green.set_align("vcenter")
    num_format1_green.set_bg_color('#C6EFCE')
    num_format1_green.set_bold()


    num_format2 = workbook.add_format()
    num_format2.set_num_format("0%;[Red]-0%")
    num_format2.set_align("center")
    num_format2.set_align("vcenter")
    num_format2.set_right(1)

    num_format2_dash_boarder = workbook.add_format()
    num_format2_dash_boarder.set_num_format("0%;[Red]-0%")
    num_format2_dash_boarder.set_align("center")
    num_format2_dash_boarder.set_align("vcenter")
    num_format2_dash_boarder.set_right(7)

    num_format2_red = workbook.add_format()
    num_format2_red.set_num_format("0%;[Red]-0%")
    num_format2_red.set_align("center")
    num_format2_red.set_align("vcenter")
    num_format2_red.set_bg_color('#FFC7CE')
    num_format2_red.set_bold()
    num_format2_red.set_right(1)

    num_format2_red_dash_boarder = workbook.add_format()
    num_format2_red_dash_boarder.set_num_format("0%;[Red]-0%")
    num_format2_red_dash_boarder.set_align("center")
    num_format2_red_dash_boarder.set_align("vcenter")
    num_format2_red_dash_boarder.set_bg_color('#FFC7CE')
    num_format2_red_dash_boarder.set_bold()
    num_format2_red_dash_boarder.set_right(7)

    num_format2_green = workbook.add_format()
    num_format2_green.set_num_format("0%;[Red]-0%")
    num_format2_green.set_align("center")
    num_format2_green.set_align("vcenter")
    num_format2_green.set_bg_color('#C6EFCE')
    num_format2_green.set_bold()
    num_format2_green.set_right(1)

    num_format2_green_dash_boarder = workbook.add_format()
    num_format2_green_dash_boarder.set_num_format("0%;[Red]-0%")
    num_format2_green_dash_boarder.set_align("center")
    num_format2_green_dash_boarder.set_align("vcenter")
    num_format2_green_dash_boarder.set_bg_color('#C6EFCE')
    num_format2_green_dash_boarder.set_bold()
    num_format2_green_dash_boarder.set_right(7)

    num_format6 = workbook.add_format()
    num_format6.set_num_format("0.00%;[Red]-0.00%")
    num_format6.set_align("center")
    num_format6.set_align("vcenter")


    num_format3 = workbook.add_format()
    num_format3.set_num_format("0;[Red]-0")
    num_format3.set_align("center")
    num_format3.set_align("vcenter")
    num_format3.set_right(1)

    leftFormat = workbook.add_format({"left": 5})
    topFormat = workbook.add_format({"top": 5})

    worksheet.set_column("A:A", 4)
    worksheet.set_column("B:B", 16)
    worksheet.set_column("C:C", 6)
    worksheet.set_column("D:D", 0)
    worksheet.set_column("E:E", 6)
    worksheet.set_column("F:F", 6)
    worksheet.set_column("G:G", 0)
    worksheet.set_column("H:H", 0)
    worksheet.set_column("I:I", 4)
    worksheet.set_column("J:J", 0)
    worksheet.set_column("K:K", 6)
    worksheet.set_column("L:L", 8)
    worksheet.set_column("M:M", 6)
    worksheet.set_column("N:N", 6)
    worksheet.set_column("O:O", 2)
    worksheet.set_column("P:P", 2)
    worksheet.set_column("Q:Q", 2)
    worksheet.set_column("R:R", 2)
    worksheet.set_column("S:S", 2)
    worksheet.set_column("T:T", 4)
    worksheet.set_column("U:U", 4)
    worksheet.set_column("V:V", 4)
    worksheet.set_column("W:W", 4)
    worksheet.set_column("X:X", 4)
    worksheet.set_column("Y:Y", 6)
    worksheet.set_column("Z:Z", 6)
    worksheet.set_column("AA:AA", 6)
    worksheet.set_column("AB:AB", 8)
    worksheet.set_column("AC:AC", 8)
    worksheet.set_column("AD:AD", 8)

    row_num, col_num = data.shape
    scol = 1

    if title is not None:
        worksheet.write_string(0, 0, title, head_format)
        merge_format = workbook.add_format(
            {
                "bold": 1,
                "font_size": 12,
                "top": 5,
                "left": 5,
                "right": 5,
                "bottom": 1,
                "align": "center",
                "valign": "vcenter",
                "fg_color": "white",
            }
        )
        worksheet.merge_range("A1:AD1", title, merge_format)
        srow = 1
    else:
        srow = 0

    for row in range(srow, data.shape[0] + 1 + srow):
        worksheet.write(row, data.shape[1], "", leftFormat)

    for col in range(0, data.shape[1]):
        worksheet.write(data.shape[0] + 1 + srow, col, "", topFormat)

    for i in range(scol, col_num + 1):
        if i == 0:
            worksheet.write_string(srow, i, data.index.name, head_format_border)
        else:
            if i == scol:
                worksheet.write_string(srow, i - scol, data.columns[i - scol], head_format_border)
            else:
                if i in [3, 6, 12, 14, 19, 24, 27]:
                    worksheet.write_string(srow, i - scol, data.columns[i - scol], head_format_bold_border)
                else:
                    worksheet.write_string(srow, i - scol, data.columns[i - scol], head_format_dash_border)

    for i in range(1, row_num + 1):
        for j in range(scol, col_num + 1):
            if j in [
                1,
                2,
                8,
                9,
                10,
                11,
                15,
                16,
                17,
                18,
                19,
                20,
                21,
                22,
                23,
                24,
                30,
                31,
                32,
                33,
                34,
                35,
                38,
            ]:
                if data.iloc[i - 1, j - scol] in ["B", "BUY", "A+", "OS", "OS CROSS"]:
                    if j in [19, 24, 30]:
                        worksheet.write(
                            i + srow,
                            j - scol,
                            data.iloc[i - 1, j - scol],
                            str_format1_green_bold_boarder,
                        )
                    else:
                        worksheet.write(
                            i + srow,
                            j - scol,
                            data.iloc[i - 1, j - scol],
                            str_format1_green_dash_boarder,
                        )
                elif data.iloc[i - 1, j - scol] in ["S", "SELL", "A-", "OB", "OB CROSS"]:
                    if j in [19, 24, 30]:
                        worksheet.write(
                            i + srow,
                            j - scol,
                            data.iloc[i - 1, j - scol],
                            str_format1_red_bold_boarder,
                        )
                    else:
                        worksheet.write(
                            i + srow,
                            j - scol,
                            data.iloc[i - 1, j - scol],
                            str_format1_red_dash_boarder,
                        )
                else:
                    if j == 1:
                        worksheet.write(i + srow, j - scol, data.iloc[i - 1, j - scol], str_format_border)
                    elif j == 2:
                        if data.iloc[i - 1, 31:38].any():
                            worksheet.write(i + srow, j - scol, data.iloc[i - 1, j - scol], str_format1)
                        else:
                            worksheet.write(i + srow, j - scol, data.iloc[i - 1, j - scol], str_format_bold)
                    else:
                        try:
                            if j in [19, 24, 30]:
                                worksheet.write(
                                    i + srow,
                                    j - scol,
                                    data.iloc[i - 1, j - scol],
                                    str_format1_bold_border,
                                )
                            else:
                                worksheet.write(
                                    i + srow,
                                    j - scol,
                                    data.iloc[i - 1, j - scol],
                                    str_format1_dash_border,
                                )
                        except:
                            if j in [19, 24, 30]:
                                worksheet.write(i + srow, j - scol, "", str_format1_bold_border)
                            else:
                                worksheet.write(i + srow, j - scol, "", str_format1_dash_border)
            else:
                if j in [
                    3,
                ]:
                    if not np.isnan(data.iloc[i - 1, j - scol]):
                        worksheet.write(i + srow, j - scol, data.iloc[i - 1, j - scol], num_format3)
                    else:
                        worksheet.write(i + srow, j - scol, "", num_format3)
                elif j in [5, 6, 25, 27, 28] and not np.isnan(data.iloc[i - 1, j - scol]):
                    if data.iloc[i - 1, j - scol] > 1.5:
                        if j in [6, 27]:
                            worksheet.write(i + srow, j - scol, data.iloc[i - 1, j - scol], num_format2_green)
                        else:
                            worksheet.write(
                                i + srow,
                                j - scol,
                                data.iloc[i - 1, j - scol],
                                num_format2_green_dash_boarder,
                            )
                    elif data.iloc[i - 1, j - scol] < -1.5:
                        if j in [6, 27]:
                            worksheet.write(i + srow, j - scol, data.iloc[i - 1, j - scol], num_format2_red)
                        else:
                            worksheet.write(
                                i + srow,
                                j - scol,
                                data.iloc[i - 1, j - scol],
                                num_format2_red_dash_boarder,
                            )
                    else:
                        if j in [6, 27]:
                            worksheet.write(i + srow, j - scol, data.iloc[i - 1, j - scol], num_format2)
                        else:
                            worksheet.write(
                                i + srow,
                                j - scol,
                                data.iloc[i - 1, j - scol],
                                num_format2_dash_boarder,
                            )
                elif j in [
                    8,
                ] and not np.isnan(data.iloc[i - 1, j - scol]):
                    worksheet.write(i + srow, j - scol, data.iloc[i - 1, j - scol], num_format6)
                else:
                    if not np.isnan(data.iloc[i - 1, j - scol]):
                        if (
                            j
                            in [
                                26,
                            ]
                            and data.iloc[i - 1, j - scol] > 3
                        ):
                            worksheet.write(i + srow, j - scol, data.iloc[i - 1, j - scol], num_format1_green)
                        elif (
                            j
                            in [
                                26,
                            ]
                            and data.iloc[i - 1, j - scol] < -3
                        ):
                            worksheet.write(i + srow, j - scol, data.iloc[i - 1, j - scol], num_format1_red)
                        else:
                            if j in [12, 14]:
                                worksheet.write(i + srow, j - scol, data.iloc[i - 1, j - scol], num_format1)
                            else:
                                worksheet.write(
                                    i + srow,
                                    j - scol,
                                    data.iloc[i - 1, j - scol],
                                    num_format1_dash_boarder,
                                )
    writer.close()
    if save_subset_to_new_sheet:
        wb = openpyxl.load_workbook(excelpath, data_only=True)
        worksheet_in = wb["sheet1"]
        worksheet_out = wb["sheet2"]
        try:
            if subset_range is None:
                log.error(f"Error cannot save subset to sheet2, as subet is {subset_range}")
                return
            output_subset = worksheet_in[subset_range]
        except Exception as e:
            log.error(f"Error cannot save subset to sheet2, with {subset_range} due to {e}")
            return
        for i in output_subset:
            for cell in i:
                new_cell = worksheet_out.cell(row=cell.row, column=cell.column, value=cell.value)
                if cell.has_style:
                    new_cell.font = copy(cell.font)
                    new_cell.border = copy(cell.border)
                    new_cell.fill = copy(cell.fill)
                    new_cell.number_format = copy(cell.number_format)
                    new_cell.protection = copy(cell.protection)
                    new_cell.alignment = copy(cell.alignment)
        worksheet_out.merge_cells("A1:AD1")
        try:
            wb.save(excelpath)
        except:
            log.error(f"Could not save {excelpath}")

    data[["KPI Index", "RngSpike", "VolSpike", "KPI", "KPI t-1"]] = data[["KPI Index", "RngSpike", "VolSpike", "KPI", "KPI t-1"]] * 100
    data["Expiry"] = data["Expiry"].fillna(-1).astype(int).replace(-1, "")
    column_format = {}
    true_keys = ["B", "BUY", "A+", "OS", "OS CROSS"]
    false_keys = ["S", "SELL", "A-", "OB", "OB CROSS"]
    for col in data.columns:
        base_format = {"width": "55px", "text-align": "center"}

        if col in ["Ticker"]:
            base_format["width"] = "150px"
        elif col in ["KPI Index", "Index", "RngSpike"]:
            base_format["width"] = "55px"
        elif col in ["Mom", "Hybrid", "Fish", "FishLT", "Counter"]:
            base_format["width"] = "30px"
        elif col in ["M", "H", "F", "LT", "CT", "Instr"]:
            base_format["width"] = "5px"
        if col in [
            "H",
            "F",
            "LT",
            "Type",
            "Stocs",
            "Hybrid",
            "Fish",
            "FishLT",
            "Index",
            "FOM",
            "Stochs",
            "Volspike20",
        ]:
            base_format["right_border"] = {"style": "dotted", "size": 2}
            base_format["left_border"] = {"style": "dotted", "size": 2}
        if col in ["Shrp 22", "KPI t-1", "RngSpike", "Counter", "Expiry", "CT"]:
            base_format["right_border"] = True
        if col in ["Shrp 3"]:
            base_format["left_border"] = True
            base_format["right_border"] = {"style": "dotted", "size": 2}

        if col in ["KPI Index", "RngSpike", "VolSpike", "KPI", "KPI t-1"]:
            base_format["format"] = "{:.2f}%"

        if col in ["M", "H", "F", "LT", "CT", "FOM", "Mom", "Hybrid", "Fish", "FishLT", "Counter"]:
            base_format["highlight_on_key"] = {
                "columns": col,
                "true_key": true_keys,
                "false_key": false_keys,
            }
        if col in ["Index"]:
            base_format["highlight_on_range"] = {"columns": col, "min": -3, "max": 3}
        elif col in ["KPI", "KPI t-1", "RngSpike", "KPI Index", "VolSpike"]:
            base_format["highlight_on_range"] = {"columns": col, "min": -150, "max": 150}
        column_format[col] = base_format
    format_row = {i: {"bottom_border": {"style": "dotted"}} for i in range(len(data)) if i != len(data)}
    html_table = table.html_format(
        data,
        header_raw=f'<p style="text-align: center; font-family:Calibri; font-weight:bold; font-size:16px">{title}</p>',
        footer=None,
        hide_cols=[
            "Mom_ytsd",
            "Hybrid_ytsd",
            "Fish_ytsd",
            "FishLT_ytsd",
            "Counter_ytsd",
            "Change 1d",
            "Change 5d",
            "IROLL_tag",
            "Ref Date",
        ],
        format_column=column_format,
        format_row=format_row,
        one_bg_color=True,
        background_color="white",
        precision=2,
        show_date=False,
    )
    return html_table


def save_output(name, asset_class):
    log.info(f"Saving Signal for {name}")
    df = pd.read_csv(ut.convert_path_to_linux(folder_path + "trend_signal.csv"))
    df_signal = df.loc[(df["Signal"] == 1) & (df["Instr"].isin(asset_class)), :].copy()

    if len(df_signal) > 0:
        df_signal = df_signal.set_index(df_signal.columns[0])
        df_signal = df_signal.sort_values("Index", ascending=False)
        df_signal.drop(["3Day", "Signal"], inplace=True, axis=1)
        output_path = ut.convert_path_to_linux(folder_path + "sort_signal_{:s}.xlsx".format(name))
        output_path_html = ut.convert_path_to_linux(folder_path + "sort_signal_{:s}.html".format(name))
        html_raw = table_format(
            excelpath=output_path,
            data=df_signal,
            name="sheet1",
            title=name,
            save_subset_to_new_sheet=True,
            subset_range="A:AD",
        )
        png_path = folder_path + "sort_signal_{:s}.png".format(name)

        c = 0
        while c < 10:
            try:
                if user != "pmlo25_svc" and sys.platform.startswith("win"):
                    excel2img.export_img(output_path, png_path, "", "sheet1!A1:AD{:d}".format(len(df_signal) + 2))
                with open(output_path_html, "w") as html_file:
                    html_file.write(html_raw)
                break
            except Exception as e:
                c += 1
                pass

        df_index = df.loc[(df["Index"] > 2) | (df["Index"] < -2), :].copy()
        df_index = df_index.set_index(df_index.columns[0])
        df_index.drop(["3Day", "Signal"], inplace=True, axis=1)
        output_path = ut.convert_path_to_linux(folder_path + "sort_index_{:s}.xlsx".format(name))
        output_path_html = ut.convert_path_to_linux(folder_path + "sort_index_{:s}.html".format(name))
        html_raw = table_format(
            excelpath=output_path,
            data=df_index,
            name="sheet1",
            title=name,
            save_subset_to_new_sheet=True,
            subset_range="A:AD",
        )
        png_path = folder_path + "sort_index_{:s}.png".format(name)

        c = 0
        while c < 10:
            try:
                if user != "pmlo25_svc" and sys.platform.startswith("win"):
                    excel2img.export_img(output_path, png_path, "", "sheet1!A1:AD{:d}".format(len(df_index) + 2))
                with open(output_path_html, "w") as html_file:
                    html_file.write(html_raw)
                break
            except:
                c += 1
                pass
    else:
        try:
            os.remove(ut.convert_path_to_linux(folder_path + "sort_signal_{:s}.xlsx".format(name)))
            os.remove(ut.convert_path_to_linux(folder_path + "sort_signal_{:s}.png".format(name)))
            os.remove(ut.convert_path_to_linux(folder_path + "sort_signal_{:s}.html".format(name)))
            os.remove(ut.convert_path_to_linux(folder_path + "sort_index_{:s}.xlsx".format(name)))
            os.remove(ut.convert_path_to_linux(folder_path + "sort_index_{:s}.png".format(name)))
            os.remove(ut.convert_path_to_linux(folder_path + "sort_index_{:s}.html".format(name)))

        except:
            pass


def send_trend_signal(asset_class_list, send_to):
    index_files = []
    index_html_str = []
    for asset_class in asset_class_list:
        file_name = ut.convert_path_to_linux(folder_path + "sort_index_{:s}.html".format(asset_class))
        if os.path.exists(file_name):
            with open(file_name, "r") as f:
                html_lines = f.readlines()
            index_html_str.append("".join(html_lines) + "</br>")
            index_files.append(file_name)

    body = [
        table.html_text("Lastest update time is {:s}".format(now_ldn().strftime("%Y-%m-%d %H:%M"))),
    ]
    for i in asset_class_list:
        if os.path.exists(ut.convert_path_to_linux("{:s}sort_signal_{:s}.html".format(folder_path, i))):
            with open(ut.convert_path_to_linux("{:s}sort_signal_{:s}.html".format(folder_path, i)), "r") as f:
                html_lines = f.readlines()
            body.append("".join(html_lines) + "</br>")

    send_email(send_to, subject="MarketScan: Trend Signal", body=body)


def send_all_tables(send_to, asset_class=None):
    asset_class_dict = {
        "Macro": ["EQT", "FI", "CCY", "VOL"],
        "Oil": ["WTI", "BRT", "DIST", "GASOLINE", "NSea"],
        "GAS": ["HH", "TTF"],
        "Metal": ["PM", "BM", "Bulks"],
        "AGS": ["GR", "SFT"],
    }

    if asset_class is None:
        asset_class = list(asset_class_dict.keys())

    for key, val in asset_class_dict.items():
        save_output(key, val)
    print(f"Sending Signal email")
    send_trend_signal(asset_class, send_to=send_to)
    figs = []
    ng_signal_path = ut.convert_path_to_linux(folder_path + "sort_signal_Oil.html")
    if os.path.exists(ng_signal_path):
        with open(ng_signal_path, "r") as f:
            html_lines = f.readlines()
        figs.append("<div style='font-family:Calibri;' >")
        figs.append("Lastest update time is {:s} </br>".format(now_ldn().strftime("%Y-%m-%d %H:%M")))
        figs.append("Oil Signal</br>")
        figs.append("".join(html_lines))
        table.figures_to_html(
            figs,
            filename=convert_path_to_linux(folder_path + "oil_signal.html"),
            add_generation_timestamp=False,
        )
    figs = []
    ng_signal_path = ut.convert_path_to_linux(folder_path + "sort_signal_Metal.html")
    if os.path.exists(ng_signal_path):
        with open(ng_signal_path, "r") as f:
            html_lines = f.readlines()
        figs.append("<div style='font-family:Calibri;' >")
        figs.append("Lastest update time is {:s} </br>".format(now_ldn().strftime("%Y-%m-%d %H:%M")))
        figs.append("Metal Signal</br>")
        figs.append("".join(html_lines))
        table.figures_to_html(
            figs,
            filename=convert_path_to_linux(folder_path + "metal_signal.html"),
            add_generation_timestamp=False,
        )


def send_alert_email_am(send_to):
    df_all = pd.read_csv(ut.convert_path_to_linux(f"{folder_path}trend_alert_all.csv"))
    if len(df_all) > 0:
        df_all.set_index("Unnamed: 0", inplace=True)
        html_alert = get_alert_email_html(df_all)

        figs = []
        for idx, row in df_all.iterrows():
            price = ts.read_csv(f"{folder_path}price_daily\\{row['Ticker']}.csv", index_name="date")
            required_cols = ["PX_HIGH", "PX_LOW", "PX_LAST", "PX_OPEN", "VOLUME"]
            if price.empty or not all(col in price.columns for col in required_cols):
                log.warning(f"Skipping chart for {row['Ticker']}: missing OHLCV columns")
                continue
            highp = price["PX_HIGH"].values
            lowp = price["PX_LOW"].values
            closep = price["PX_LAST"].values
            avg_tr = talib.atr(highp, lowp, closep, 3)
            fig = make_subplots(
                rows=3,
                cols=1,
                row_heights=[0.6, 0.2, 0.2],
                shared_xaxes=True,
                vertical_spacing=0.02,
            )
            fig.add_trace(
                go.Candlestick(
                    x=price.index[-25:],
                    open=price["PX_OPEN"].iloc[-25:],
                    high=price["PX_HIGH"].iloc[-25:],
                    low=price["PX_LOW"].iloc[-25:],
                    close=price["PX_LAST"].iloc[-25:],
                    name="Price",
                )
            )
            fig["layout"]["yaxis1"]["title"] = "Price"
            fig.add_trace(
                go.Bar(
                    name="Volume",
                    x=price.index[-25:],
                    y=price["VOLUME"].iloc[-25:],
                    showlegend=True,
                ),
                row=2,
                col=1,
            )
            fig["layout"]["yaxis2"]["title"] = "Volume"
            fig.add_trace(
                go.Scatter(
                    x=price.index[-25:],
                    y=avg_tr[-25:],
                    showlegend=True,
                    name="3d ATR",
                    line=dict(width=1),
                ),
                row=3,
                col=1,
            )
            fig["layout"]["yaxis3"]["title"] = "ATR"
            fig.update_xaxes(
                rangebreaks=[
                    dict(bounds=["sat", "mon"]),  # hide weekends
                    dict(values=["2015-12-25", "2016-01-01"]),  # hide Christmas and New Year's
                ]
            )
            fig.update_layout(
                title={"text": f"Range-Vol - {row['Ticker']}", "x": 0.5, "xanchor": "center"},
                barmode="stack",
                width=900,
                height=600,
                xaxis_rangeslider_visible=False,
            )
            figs.append(fig)
        table.figures_to_html(figs, filename=convert_path_to_linux(f"{folder_path}charts\\range_vol.html"))
        if html_alert is None:
            html_alert = "No Range-Vol alert today"
    else:
        html_alert = "No Range-Vol alert today"

    df_macro = pd.read_csv(ut.convert_path_to_linux(f"{folder_path}trend_alert_macro.csv"))
    if len(df_macro) > 0:
        df_macro.set_index("Unnamed: 0", inplace=True)
        html_macro_alert = get_alert_email_html(df_macro)

        figs = []
        for idx, row in df_macro.iterrows():
            price = ts.read_csv(f"{folder_path}price_daily\\{row['Ticker']}.csv", index_name="date")
            required_cols = ["PX_HIGH", "PX_LOW", "PX_LAST", "PX_OPEN", "VOLUME"]
            if price.empty or not all(col in price.columns for col in required_cols):
                log.warning(f"Skipping chart for {row['Ticker']}: missing OHLCV columns")
                continue
            highp = price["PX_HIGH"].values
            lowp = price["PX_LOW"].values
            closep = price["PX_LAST"].values
            avg_tr = talib.atr(highp, lowp, closep, 3)
            fig = make_subplots(
                rows=3,
                cols=1,
                row_heights=[0.6, 0.2, 0.2],
                shared_xaxes=True,
                vertical_spacing=0.02,
            )
            fig.add_trace(
                go.Candlestick(
                    x=price.index[-25:],
                    open=price["PX_OPEN"].iloc[-25:],
                    high=price["PX_HIGH"].iloc[-25:],
                    low=price["PX_LOW"].iloc[-25:],
                    close=price["PX_LAST"].iloc[-25:],
                    name="Price",
                )
            )
            fig["layout"]["yaxis1"]["title"] = "Price"
            fig.add_trace(
                go.Bar(
                    name="Volume",
                    x=price.index[-25:],
                    y=price["VOLUME"].iloc[-25:],
                    showlegend=True,
                ),
                row=2,
                col=1,
            )
            fig["layout"]["yaxis2"]["title"] = "Volume"
            fig.add_trace(
                go.Scatter(
                    x=price.index[-25:],
                    y=avg_tr[-25:],
                    showlegend=True,
                    name="3d ATR",
                    line=dict(width=1),
                ),
                row=3,
                col=1,
            )
            fig["layout"]["yaxis3"]["title"] = "ATR"
            fig.update_xaxes(
                rangebreaks=[
                    dict(bounds=["sat", "mon"]),  # hide weekends
                    dict(values=["2015-12-25", "2016-01-01"]),  # hide Christmas and New Year's
                ]
            )
            fig.update_layout(
                title={"text": f"Range-Vol - {row['Ticker']}", "x": 0.5, "xanchor": "center"},
                barmode="stack",
                width=900,
                height=600,
                xaxis_rangeslider_visible=False,
            )
            figs.append(fig)
        table.figures_to_html(figs, filename=convert_path_to_linux(f"{folder_path}charts\\range_vol_macro.html"))
        if html_alert is None:
            html_alert = "No Macro Range-Vol alert today"
    else:
        html_macro_alert = "No Macro Range-Vol alert today"

    df_all = pd.read_csv(ut.convert_path_to_linux(f"{folder_path}intraday_cum_vol.csv"))
    df_all = df_all.loc[~df_all["Instr"].isin(["GR", "SFT", "Bulks"]), :]
    df_all = df_all.loc[df_all["chg type"] == "FLAT", :]
    if len(df_all) > 0:
        df_all.set_index("Unnamed: 0", inplace=True)
        df_all.reset_index(drop=True, inplace=True)
        df_all["Ticker"] = [x.rpartition(" ")[0] for x in df_all["Ticker"]]
        pct_format_rows = df_all.loc[df_all["chg type"] == "FLAT", :].index
        dec_format_rows = df_all.loc[df_all["chg type"] == "SPRD", :].index
        html_intravol = table.html_format(
            df=df_all,
            precision=2,
            hide_cols=["chg type"],
            show_date=True,
            format_column={
                "Instr": {"width": "20px", "text-align": "center", "right_border": True},
                "Ticker": {"width": "160px", "text-align": "center", "right_border": True},
                "Volume T": {
                    "width": "80px",
                    "text-align": "center",
                    "format": "{:.0%}",
                    "right_border": True,
                },
                "Change T": {"width": "80px", "text-align": "center"},
            },
            format_row={
                tuple(pct_format_rows): {"format": "{:.2%}", "columns": ["Change T"]},
                tuple(dec_format_rows): {"format": "{:.2f}", "columns": ["Change T"]},
            },
        )
    else:
        html_intravol = "No intraday cumulative volume spike"

    htmls, _ = vol_alert()
    html_option = " ".join(htmls[2:4])
    html_imp_vs_real = " ".join(htmls[6:8])
    html_iv_chg = htmls[9]

    df_diverge = pd.read_csv(ut.convert_path_to_linux(f"{folder_path}\\diverge_signal.csv"))
    df_diverge.set_index("Unnamed: 0", inplace=True)
    if len(df_diverge) > 0:
        html_diverge = send_diverge_emails(df_diverge)

        figs_diverg = []
        for idx, row in df_diverge.iterrows():
            ticker0 = row["Ticker"].split("/")[0] + " Comdty"
            ticker1 = row["Ticker"].split("/")[1] + " Comdty"
            i0 = idx.split(" / ")[0]
            i1 = idx.split(" / ")[1]
            price0 = ts.read_csv(f"{folder_path}price_daily\\{ticker0}.csv", index_name="date")
            if price0.empty:
                print(f"Error: No data in {folder_path}price_daily\\{ticker0}.csv skipping.")
                continue
            price1 = ts.read_csv(f"{folder_path}price_daily\\{ticker1}.csv", index_name="date")
            price0_chg = price0["PX_LAST"].diff()
            price1_chg = price1["PX_LAST"].diff()
            if ("S:XBXB" in i0 or "S:HOHO" in i0) and "S:COCO" in i1:
                diverg_ratio_3d = price0_chg.rolling(3).sum() * 0.42 / price1_chg.rolling(3).sum()
                diverg_ratio_8d = price0_chg.rolling(8).sum() * 0.42 / price1_chg.rolling(8).sum()
            else:
                diverg_ratio_3d = price0_chg.rolling(3).sum() / price1_chg.rolling(3).sum()
                diverg_ratio_8d = price0_chg.rolling(8).sum() / price1_chg.rolling(8).sum()
            price1 = price1.reindex(price0.index)
            diverg_ratio_3d = diverg_ratio_3d.reindex(price0.index)
            diverg_ratio_3d = diverg_ratio_3d.to_frame("Ratio_3d")
            diverg_ratio_8d = diverg_ratio_8d.reindex(price0.index)
            diverg_ratio_8d = diverg_ratio_8d.to_frame("Ratio_8d")
            fig = make_subplots(
                rows=3,
                cols=1,
                row_heights=[0.4, 0.4, 0.2],
                shared_xaxes=True,
                vertical_spacing=0.02,
            )
            fig.add_trace(
                go.Candlestick(
                    x=price0.index[-25:],
                    open=price0["PX_OPEN"].iloc[-25:],
                    high=price0["PX_HIGH"].iloc[-25:],
                    low=price0["PX_LOW"].iloc[-25:],
                    close=price0["PX_LAST"].iloc[-25:],
                    name=ticker0,
                )
            )
            fig.add_trace(
                go.Candlestick(
                    x=price1.index[-25:],
                    open=price1["PX_OPEN"].iloc[-25:],
                    high=price1["PX_HIGH"].iloc[-25:],
                    low=price1["PX_LOW"].iloc[-25:],
                    close=price1["PX_LAST"].iloc[-25:],
                    name=ticker1,
                ),
                row=2,
                col=1,
            )
            fig.add_trace(
                go.Scatter(
                    x=diverg_ratio_3d.index[-25:],
                    y=diverg_ratio_3d["Ratio_3d"].iloc[-25:],
                    showlegend=True,
                    name="Ratio_3d",
                    line=dict(width=1),
                ),
                row=3,
                col=1,
            )
            fig.add_trace(
                go.Scatter(
                    x=diverg_ratio_8d.index[-25:],
                    y=diverg_ratio_8d["Ratio_8d"].iloc[-25:],
                    showlegend=True,
                    name="Ratio_8d",
                    line=dict(width=1),
                ),
                row=3,
                col=1,
            )
            fig.update_xaxes(
                rangebreaks=[
                    dict(bounds=["sat", "mon"]),  # hide weekends
                    dict(values=["2015-12-25", "2016-01-01"]),  # hide Christmas and New Year's
                ]
            )
            fig.update_layout(
                title={"text": f"Divergence - {row['Ticker']}", "x": 0.5, "xanchor": "center"},
                barmode="stack",
                width=900,
                height=600,
                xaxis1=dict(rangeslider=dict(visible=False)),
                xaxis2=dict(rangeslider=dict(visible=False)),
            )
            figs_diverg.append(fig)
        table.figures_to_html(figs_diverg, filename=convert_path_to_linux(f"{folder_path}charts\\divergence.html"))

    else:
        html_diverge = "No divergence alert today"
    table.to_html([html_diverge], convert_path_to_linux(f"{folder_path}divergence_only.html"), add_home=False, add_generation_timestamp=False)

    fly_switch_signal = pd.read_csv(ut.convert_path_to_linux(f"{folder_path}\\fly_switch_signal.csv"))
    fly_switch_signal.set_index("Unnamed: 0", inplace=True)
    if len(fly_switch_signal) > 0:
        fly_switch_signal.index.name = "Name"
        fly_switch_signal.reset_index(inplace=True)
        fly_swith_html = table.html_format(
            fly_switch_signal,
            precision=2,
            format_column={
                "Name": {"width": "160px", "text-align": "center", "right_border": True},
                tuple(fly_switch_signal.columns[1:]): {"width": "100px", "text-align": "center"},
            },
        )
    else:
        fly_swith_html = "No fly switch alert today"
    table.to_html([fly_swith_html], convert_path_to_linux(f"{folder_path}fly_switch_only.html"), add_home=False, add_generation_timestamp=False)

    alert_spot = spot_price_alert()
    if len(alert_spot) > 0:
        html_spot = table.html_format(
            alert_spot,
            precision=2,
            format_column={
                "Spot": {"width": "160px", "text-align": "center", "right_border": True},
                "Market": {"width": "160px", "text-align": "center", "right_border": True},
                "Price change": {"width": "100px", "text-align": "center"},
                "Linked sprd chg": {"width": "100px", "text-align": "center"},
                "Price": {"width": "100px", "text-align": "center"},
                "Price 20d MA": {"width": "100px", "text-align": "center"},
                "20d z-score": {"width": "100px", "text-align": "center"},
                "Latest Date": {"width": "100px", "text-align": "center"},
            },
        )
        spot_alert_outpath = convert_path_to_linux(
            f"{folder_path}\\physical_oil\\spot_alert_{today().strftime('%Y%m%d')}.txt",
        )
        with open(spot_alert_outpath, mode="w") as file:
            file.write(html_spot)
    else:
        html_spot = "No physical spot price alert today"
    table.to_html([html_spot], convert_path_to_linux(f"{folder_path}spot_only.html"), add_home=False, add_generation_timestamp=False)

    alert_phys, product_alert_phys = physical_price_alert()
    if len(alert_phys) > 0:
        html_phys = table.html_format(
            alert_phys,
            precision=2,
            format_column={
                "Crude": {"width": "160px", "text-align": "center", "right_border": True},
                "Market": {"width": "160px", "text-align": "center", "right_border": True},
                "Price change": {"width": "100px", "text-align": "center"},
                "Linked sprd chg": {"width": "100px", "text-align": "center"},
                "Price": {"width": "100px", "text-align": "center"},
                "Price 20d MA": {"width": "100px", "text-align": "center"},
                "20d z-score": {"width": "100px", "text-align": "center"},
                "Latest Date": {"width": "100px", "text-align": "center"},
            },
        )
    else:
        html_phys = "No physical crude alert today"
    table.to_html([html_phys], convert_path_to_linux(f"{folder_path}physical_only.html"), add_home=False, add_generation_timestamp=False)
    if len(product_alert_phys) > 0:
        html_product_phys = table.html_format(
            product_alert_phys,
            precision=2,
            format_column={
                "Product": {"width": "160px", "text-align": "center", "right_border": True},
                "Market": {"width": "160px", "text-align": "center", "right_border": True},
                "Price change": {"width": "100px", "text-align": "center"},
                "Linked sprd chg": {"width": "100px", "text-align": "center"},
                "Price": {"width": "100px", "text-align": "center"},
                "Price 20d MA": {"width": "100px", "text-align": "center"},
                "20d z-score": {"width": "100px", "text-align": "center"},
                "Latest Date": {"width": "100px", "text-align": "center"},
            },
        )
    else:
        html_product_phys = "No physical product alert today"
    table.to_html([html_product_phys], convert_path_to_linux(f"{folder_path}physical_product_only.html"), add_home=False, add_generation_timestamp=False)

    alert_crack = crack_price_alert()
    if len(alert_crack) > 0:
        html_crack = table.html_format(
            alert_crack,
            precision=2,
            format_column={
                "Crack": {"width": "160px", "text-align": "center", "right_border": True},
                "Market": {"width": "160px", "text-align": "center", "right_border": True},
                "Price change": {"width": "100px", "text-align": "center"},
                "Linked sprd chg": {"width": "100px", "text-align": "center"},
                "Price": {"width": "100px", "text-align": "center"},
                "Price 20d MA": {"width": "100px", "text-align": "center"},
                "20d z-score": {"width": "100px", "text-align": "center"},
                "Latest Date": {"width": "100px", "text-align": "center"},
            },
        )
    else:
        html_crack = "No crack alert today"
    table.to_html([html_crack], convert_path_to_linux(f"{folder_path}crack_only.html"), add_home=False, add_generation_timestamp=False)

    alert_swap, product_alert_swap = swap_price_alert()
    if len(alert_swap) > 0:
        html_swap = table.html_format(
            alert_swap,
            precision=2,
            format_column={
                "Crude": {"width": "160px", "text-align": "center", "right_border": True},
                "Market": {"width": "160px", "text-align": "center", "right_border": True},
                "Price change": {"width": "100px", "text-align": "center"},
                "Linked sprd chg": {"width": "100px", "text-align": "center"},
                "Price": {"width": "100px", "text-align": "center"},
                "Price 20d MA": {"width": "100px", "text-align": "center"},
                "20d z-score": {"width": "100px", "text-align": "center"},
                "Latest Date": {"width": "100px", "text-align": "center"},
            },
        )
    else:
        html_swap = "No crude swap alert today"
    table.to_html([html_swap], convert_path_to_linux(f"{folder_path}swap_only.html"), add_home=False, add_generation_timestamp=False)
    if len(product_alert_swap) > 0:
        html_product_swap = table.html_format(
            product_alert_swap,
            precision=2,
            format_column={
                "Product": {"width": "160px", "text-align": "center", "right_border": True},
                "Market": {"width": "160px", "text-align": "center", "right_border": True},
                "Price change": {"width": "100px", "text-align": "center"},
                "Linked sprd chg": {"width": "100px", "text-align": "center"},
                "Price": {"width": "100px", "text-align": "center"},
                "Price 20d MA": {"width": "100px", "text-align": "center"},
                "20d z-score": {"width": "100px", "text-align": "center"},
                "Latest Date": {"width": "100px", "text-align": "center"},
            },
        )
    else:
        html_product_swap = "No product swap alert today"
    table.to_html([html_product_swap], convert_path_to_linux(f"{folder_path}product_swap_only.html"), add_home=False, add_generation_timestamp=False)

    eu_gas = pd.read_csv(ut.convert_path_to_linux(f"{output_path}\\csvs\\ttf\\ttf_alert.csv"))
    eu_gas.set_index("Unnamed: 0", inplace=True)
    if len(eu_gas) > 0:
        html_eugas = table.html_format(
            eu_gas,
            precision=2,
            format_column={
                "Hub": {"width": "160px", "text-align": "center"},
                "Tenor": {"width": "100px", "text-align": "center"},
                "Price change": {"width": "100px", "text-align": "center"},
                "Price": {"width": "100px", "text-align": "center"},
                "Price 20d MA": {"width": "100px", "text-align": "center"},
                "3m z-score": {"width": "100px", "text-align": "center"},
            },
        )
    else:
        html_eugas = "No EU gas spread alert today"
    table.to_html([html_eugas], convert_path_to_linux(f"{folder_path}ttf_only.html"), add_home=False, add_generation_timestamp=False)

    us_gas = pd.read_csv(ut.convert_path_to_linux(f"{output_path}\\csvs\\gas\\basis\\usgas_alert.csv"))
    us_gas.set_index("Unnamed: 0", inplace=True)
    if len(us_gas) > 0:
        html_usgas = table.html_format(
            us_gas,
            precision=2,
            format_column={
                "Hub": {"width": "160px", "text-align": "center"},
                "Tenor": {"width": "100px", "text-align": "center"},
                "Price change": {"width": "100px", "text-align": "center"},
                "Price": {"width": "100px", "text-align": "center"},
                "Price 20d MA": {"width": "100px", "text-align": "center"},
                "3m z-score": {"width": "100px", "text-align": "center"},
                "Latest Date": {"width": "100px", "text-align": "center"},
            },
        )
    else:
        html_usgas = "No US gas spread alert today"
    table.to_html([html_usgas], convert_path_to_linux(f"{folder_path}usgas_only.html"), add_home=False, add_generation_timestamp=False)

    df_bm = pd.read_csv(ut.convert_path_to_linux(f"{folder_path}bm_alert_am.csv"))
    if len(df_bm) > 0:
        df_bm.set_index("Unnamed: 0", inplace=True)
        df_bm.reset_index(drop=True, inplace=True)
        df_bm["Ticker"] = [x.rpartition(" ")[0] for x in df_bm["Ticker"]]
        pct_format_rows = df_bm.loc[df_bm["chg type"] == "FLAT", :].index
        dec_format_rows = df_bm.loc[df_bm["chg type"] == "SPRD", :].index
        html_bm = table.html_format(
            df_bm,
            precision=2,
            hide_cols=["chg type"],
            format_column={
                "Ticker": {"width": "160px", "text-align": "center"},
                "Price change": {"width": "100px", "text-align": "center"},
                "Price": {"width": "100px", "text-align": "center"},
                "Price 20d MA": {"width": "100px", "text-align": "center"},
                "20d z-score": {"width": "100px", "text-align": "center"},
                "Date": {"width": "100px", "text-align": "center"},
            },
            format_row={
                tuple(pct_format_rows): {"format": "{:.2%}", "columns": ["Price change"]},
                tuple(dec_format_rows): {"format": "{:.2f}", "columns": ["Price change"]},
            },
        )
    else:
        html_bm = "No base metal alert"
    table.to_html([html_bm], convert_path_to_linux(f"{folder_path}bm_only.html"), add_home=False, add_generation_timestamp=False)
    risk_index_html = get_risk_index_html()

    html_table = [
        "<div style='font-family:Calibri;' >",
        f"Update time {now_ldn().strftime('%Y-%m-%d %H:%M')}. <br>",
        "<b>Risk Index</b><br>",
        risk_index_html,
        "<b>Commods Range-Vol alert: </b><br>",
        '<a href="https://.elementcapital.corp/display/LO25/Range-Vol+and+Divergence">Description</a><br>',
        html_alert,
        '<a href="{}charts\\range_vol.html">RangeVol Charts</a><br>'.format(folder_path),
        "<br><b>Macro Range-Vol alert: </b><br>",
        html_macro_alert,
        '<a href="{}charts\\range_vol_macro.html">Macro RangeVol Charts</a><br>'.format(folder_path),
        "<br><b>Intraday cumulative volume alert: </b><br>",
        html_intravol,
        "<br><b>Option volume alert: </b><br>",
        html_option,
        '<a href="{}charts\\range_vol_option.html">Option Charts</a><br>'.format(folder_path),
        "<br><b>Implied vs realized volatility alert: </b><br>",
        html_imp_vs_real,
        "<br><b>Implied volatility change alert: </b><br>",
        html_iv_chg,
        "<br><b>Spread and flat price divergence alert: </b><br>",
        html_diverge,
        '<a href="{}charts\\divergence.html">Divergence Charts</a><br>'.format(folder_path),
        "<br><b>Fly switch alert: </b><br>",
        fly_swith_html,
        '<a href="{}charts\\fly_switch_signal.html">Fly Switch Charts</a><br>'.format(folder_path),
        "<br><b>Physical Oil price alert: </b><br>",
        html_phys,
        '<a href="{}charts\\physical_alert.html">Physical Crude Charts</a><br>'.format(folder_path),
        html_product_phys,
        '<a href="{}charts\\physical_product_alert.html">Physical Product Charts</a><br>'.format(folder_path),
        "<br><b>Oil swap alert: </b><br>",
        html_swap,
        '<a href="{}charts\\swap_alert.html">Crude Swap Charts</a><br>'.format(folder_path),
        html_product_swap,
        '<a href="{}charts\\product_swap_alert.html">Product Swap Charts</a><br>'.format(folder_path),
        html_crack,
        '<a href="{}charts\\crack_alert.html">Crack Charts</a><br>'.format(folder_path),
        "<br><b>EU gas spread alert: </b><br>",
        html_eugas,
        '<a href="{}\\ttf\\links\\ttf_alert_charts.html">EU Gas Charts</a><br>'.format(html_path),
        "<br><b>US gas spread alert: </b><br>",
        html_usgas,
        '<a href="{}\\gas\\links\\usgas_alert_charts.html">US Gas Charts</a><br>'.format(html_path),
        "<br><b>Base metal alert: </b><br>",
        html_bm,
    ]

    table.figures_to_html(
        html_table,
        filename=convert_path_to_linux(f"{html_path}\\cross_cmds\\market_scan_alert.html"),
        task_name=report_name,
    )
    generate_range_vol_page_by_asset_class(asset_class="Oil")
    generate_range_vol_page_by_asset_class(asset_class="GAS")
    generate_range_vol_page_by_asset_class(asset_class="Metal")
    send_email(
        send_to=send_to,
        subject="Range-Vol and Divergence AM",
        body=[html_table],
        html_path=ut.convert_path_to_linux(f"{html_path}\\cross_cmds\\market_scan_alert.html"),
    )


def send_alert_email_pm(send_to):
    df_all = pd.read_csv(ut.convert_path_to_linux(f"{folder_path}trend_alert_all_today.csv"))
    if len(df_all) > 0:
        df_all.set_index("Unnamed: 0", inplace=True)
        html_alert = get_alert_email_html(df_all)

        figs = []
        for idx, row in df_all.iterrows():
            if not np.isnan(row["Volume T"]):
                price = ts.read_csv(f"{folder_path}price_daily\\{row['Ticker']}.csv", index_name="date")
                required_cols = ["PX_HIGH", "PX_LOW", "PX_LAST", "PX_OPEN", "VOLUME"]
                if price.empty or not all(col in price.columns for col in required_cols):
                    log.warning(f"Skipping chart for {row['Ticker']}: missing OHLCV columns")
                    continue
                highp = price["PX_HIGH"].values
                lowp = price["PX_LOW"].values
                closep = price["PX_LAST"].values
                avg_tr = talib.atr(highp, lowp, closep, 3)
                fig = make_subplots(
                    rows=3,
                    cols=1,
                    row_heights=[0.6, 0.2, 0.2],
                    shared_xaxes=True,
                    vertical_spacing=0.02,
                )
                fig.add_trace(
                    go.Candlestick(
                        x=price.index[-25:],
                        open=price["PX_OPEN"].iloc[-25:],
                        high=price["PX_HIGH"].iloc[-25:],
                        low=price["PX_LOW"].iloc[-25:],
                        close=price["PX_LAST"].iloc[-25:],
                        name="Price",
                    )
                )
                fig["layout"]["yaxis1"]["title"] = "Price"
                fig.add_trace(
                    go.Bar(
                        name="Volume",
                        x=price.index[-25:],
                        y=price["VOLUME"].iloc[-25:],
                        showlegend=True,
                    ),
                    row=2,
                    col=1,
                )
                fig["layout"]["yaxis2"]["title"] = "Volume"
                fig.add_trace(
                    go.Scatter(
                        x=price.index[-25:],
                        y=avg_tr[-25:],
                        showlegend=True,
                        name="3d ATR",
                        line=dict(width=1),
                    ),
                    row=3,
                    col=1,
                )
                fig["layout"]["yaxis3"]["title"] = "ATR"
                fig.update_xaxes(
                    rangebreaks=[
                        dict(bounds=["sat", "mon"]),  # hide weekends
                        dict(values=["2015-12-25", "2016-01-01"]),  # hide Christmas and New Year's
                    ]
                )
                fig.update_layout(
                    title={"text": f"Range-Vol - {row['Ticker']}", "x": 0.5, "xanchor": "center"},
                    barmode="stack",
                    width=900,
                    height=600,
                    xaxis_rangeslider_visible=False,
                )
                figs.append(fig)
        table.figures_to_html(figs, filename=convert_path_to_linux(f"{folder_path}charts\\range_vol.html"))
        if html_alert is None:
            html_alert = "No Range-Vol alert today<br>"
    else:
        html_alert = "No Range-Vol alert today<br>"

    df_macro = pd.read_csv(ut.convert_path_to_linux(f"{folder_path}trend_alert_macro_today.csv"))
    if len(df_macro) > 0:
        df_macro.set_index("Unnamed: 0", inplace=True)
        html_macro_alert = get_alert_email_html(df_macro)

        figs = []
        for idx, row in df_macro.iterrows():
            if not np.isnan(row["Volume T"]):
                price = ts.read_csv(f"{folder_path}price_daily\\{row['Ticker']}.csv", index_name="date")
                required_cols = ["PX_HIGH", "PX_LOW", "PX_LAST", "PX_OPEN", "VOLUME"]
                if price.empty or not all(col in price.columns for col in required_cols):
                    log.warning(f"Skipping chart for {row['Ticker']}: missing OHLCV columns")
                    continue
                highp = price["PX_HIGH"].values
                lowp = price["PX_LOW"].values
                closep = price["PX_LAST"].values
                avg_tr = talib.atr(highp, lowp, closep, 3)
                fig = make_subplots(
                    rows=3,
                    cols=1,
                    row_heights=[0.6, 0.2, 0.2],
                    shared_xaxes=True,
                    vertical_spacing=0.02,
                )
                fig.add_trace(
                    go.Candlestick(
                        x=price.index[-25:],
                        open=price["PX_OPEN"].iloc[-25:],
                        high=price["PX_HIGH"].iloc[-25:],
                        low=price["PX_LOW"].iloc[-25:],
                        close=price["PX_LAST"].iloc[-25:],
                        name="Price",
                    )
                )
                fig["layout"]["yaxis1"]["title"] = "Price"
                fig.add_trace(
                    go.Bar(
                        name="Volume",
                        x=price.index[-25:],
                        y=price["VOLUME"].iloc[-25:],
                        showlegend=True,
                    ),
                    row=2,
                    col=1,
                )
                fig["layout"]["yaxis2"]["title"] = "Volume"
                fig.add_trace(
                    go.Scatter(
                        x=price.index[-25:],
                        y=avg_tr[-25:],
                        showlegend=True,
                        name="3d ATR",
                        line=dict(width=1),
                    ),
                    row=3,
                    col=1,
                )
                fig["layout"]["yaxis3"]["title"] = "ATR"
                fig.update_xaxes(
                    rangebreaks=[
                        dict(bounds=["sat", "mon"]),  # hide weekends
                        dict(values=["2015-12-25", "2016-01-01"]),  # hide Christmas and New Year's
                    ]
                )
                fig.update_layout(
                    title={"text": f"Range-Vol - {row['Ticker']}", "x": 0.5, "xanchor": "center"},
                    barmode="stack",
                    width=900,
                    height=600,
                    xaxis_rangeslider_visible=False,
                )
                figs.append(fig)
        table.figures_to_html(figs, filename=convert_path_to_linux(f"{folder_path}charts\\range_vol_macro.html"))
        if html_macro_alert is None:
            html_macro_alert = "No Macro Range-Vol alert today<br>"
    else:
        html_macro_alert = "No Macro Range-Vol alert today<br>"

    df_all = pd.read_csv(ut.convert_path_to_linux(f"{folder_path}intraday_cum_vol.csv"))
    df_all = df_all.loc[df_all["chg type"] == "FLAT", :]
    if len(df_all) > 0:
        df_all.set_index("Unnamed: 0", inplace=True)
        df_all.reset_index(drop=True, inplace=True)
        df_all["Ticker"] = [x.split(" ")[0] for x in df_all["Ticker"]]
        pct_format_rows = df_all.loc[df_all["chg type"] == "FLAT", :].index
        dec_format_rows = df_all.loc[df_all["chg type"] == "SPRD", :].index
        html_intravol = table.html_format(
            df=df_all,
            precision=2,
            hide_cols=["chg type"],
            format_column={
                "Instr": {"width": "20px", "text-align": "center", "right_border": True},
                "Ticker": {"width": "160px", "text-align": "center", "right_border": True},
                "Volume T": {
                    "width": "80px",
                    "text-align": "center",
                    "format": "{:.0%}",
                    "right_border": True,
                },
                "Change T": {"width": "80px", "text-align": "center"},
            },
            format_row={
                tuple(pct_format_rows): {"format": "{:.2%}", "columns": ["Change T"]},
                tuple(dec_format_rows): {"format": "{:.2f}", "columns": ["Change T"]},
            },
        )
    else:
        html_intravol = "No intraday cumulative volume spike<br>"

    try:
        htmls, _ = vol_alert()
        html_option = htmls[2]
    except Exception as e:
        print(f"Error in vol_alert: {e}")
        htmls = ["No option volume alert<br>", "No option IV alert<br>", "No option skew alert<br>"]
        html_option = htmls[1]

    df_bm = pd.read_csv(ut.convert_path_to_linux(f"{folder_path}bm_alert_pm.csv"))
    if len(df_bm) > 0:
        df_bm.set_index("Unnamed: 0", inplace=True)
        df_bm.reset_index(drop=True, inplace=True)
        df_bm["Ticker"] = [x.split(" ")[0] for x in df_bm["Ticker"]]
        pct_format_rows = df_bm.loc[df_bm["chg type"] == "FLAT", :].index
        dec_format_rows = df_bm.loc[df_bm["chg type"] == "SPRD", :].index
        html_bm = table.html_format(
            df_bm,
            precision=2,
            hide_cols=["chg type"],
            format_column={
                "Ticker": {"width": "160px", "text-align": "center"},
                "Price change": {"width": "100px", "text-align": "center"},
                "Price": {"width": "100px", "text-align": "center"},
                "Price 20d MA": {"width": "100px", "text-align": "center"},
                "20d z-score": {"width": "100px", "text-align": "center"},
                "Date": {"width": "100px", "text-align": "center"},
            },
            format_row={
                tuple(pct_format_rows): {"format": "{:.2%}", "columns": ["Price change"]},
                tuple(dec_format_rows): {"format": "{:.2f}", "columns": ["Price change"]},
            },
        )
    else:
        html_bm = "No base metal alert<br>"

    risk_index_html = get_risk_index_html()
    html_table = [
        "<div style='font-family:Calibri;' >",
        f"Update time {now_ldn().strftime('%Y-%m-%d %H:%M')}. <br>",
        "<b>Risk Index</b><br>",
        risk_index_html,
        "<b>Range-Vol alert: </b><br>",
        '<a href="https://elementcapital.atlassian.net/wiki/spaces/LO25/pages/1774685381/Range-Vol+and+Divergence">Description</a><br>',
        html_alert,
        '<a href="{}charts\\range_vol.html">RangeVol Charts</a><br>'.format(folder_path),
        "<br><b>Macro Range-Vol alert: </b><br>",
        html_macro_alert,
        '<a href="{}charts\\range_vol_macro.html">Macro RangeVol Charts</a><br>'.format(folder_path),
        "<br><b>Intraday cumulative volume alert: </b><br>",
        html_intravol,
        "<br><b>Option volume alert: </b><br>",
        html_option,
        '<a href="{}charts\\range_vol_option.html">Option Charts</a><br>'.format(folder_path),
        "<br><b>Base metal alert: </b><br>",
        html_bm,
    ]

    table.figures_to_html(
        html_table,
        filename=convert_path_to_linux(f"{html_path}\\cross_cmds\\market_scan_pm_alert.html"),
        task_name="Range-Vol and Divergence PM",
    )
    send_email(
        send_to=send_to,
        subject="Range-Vol and Divergence PM",
        body=[
            html_table,
            "<br><br>This is automated email sent at {:s}. <br><br>".format(now_ldn().strftime("%Y-%m-%d %H:%M")),
        ],
        html_path=f"{html_path}\\cross_cmds\\market_scan_pm_alert.html",
    )


def contract_oi_chart(active="NGA Comdty"):
    fut_chain = bbg.bbulkref(active, "FUT_CHAIN", raw=False)
    sdate = today() - BDay(30)
    edate = today()
    if active in fut_chain.columns:
        fut_chain = fut_chain[active]
    if isinstance(fut_chain, pd.DataFrame):
        fut_chain = list(fut_chain["FUT_CHAIN"].values[:36])
    else:
        fut_chain = list(fut_chain.values[:36])
    oi = bbg.bdh(fut_chain, ["OPEN_INT"], sdate, edate)
    if isinstance(oi.columns, pd.MultiIndex):
        oi = oi[fut_chain]
    oi = oi[fut_chain]
    base_ticker = tk.decompose_ticker(oi.columns[0])["base_ticker"]
    oi.columns = [x.split(" ")[0][len(base_ticker):] for x in oi.columns]
    oi_mchg = oi.iloc[-1, :] - oi.iloc[-22, :]
    oi_wchg = oi.iloc[-1, :] - oi.iloc[-5, :]
    oi_dchg = oi.iloc[-1, :] - oi.iloc[-2, :]
    oi_chg = pd.concat([oi_dchg, oi_wchg, oi_mchg], axis=1)
    oi_chg.columns = ["daily", "weekly", "monthly"]
    return bar_chart(oi_chg, title=f"{base_ticker} Open Interest Change", barmode="group")


def send_alert_email_gas(send_to):
    df_all = pd.read_csv(ut.convert_path_to_linux(f"{folder_path}gas\\trend_alert_gas.csv"))
    if len(df_all) > 0:
        df_all.set_index("Unnamed: 0", inplace=True)
        html_alert = get_alert_email_html(df_all)

        figs = []
        for idx, row in df_all.iterrows():
            price = ts.read_csv(f"{folder_path}price_daily\\{row['Ticker']}.csv", index_name="date")
            required_cols = ["PX_HIGH", "PX_LOW", "PX_LAST", "PX_OPEN", "VOLUME"]
            if price.empty or not all(col in price.columns for col in required_cols):
                log.warning(f"Skipping chart for {row['Ticker']}: missing OHLCV columns")
                continue
            highp = price["PX_HIGH"].values
            lowp = price["PX_LOW"].values
            closep = price["PX_LAST"].values
            avg_tr = talib.atr(highp, lowp, closep, 3)
            fig = make_subplots(
                rows=3,
                cols=1,
                row_heights=[0.6, 0.2, 0.2],
                shared_xaxes=True,
                vertical_spacing=0.02,
            )
            fig.add_trace(
                go.Candlestick(
                    x=price.index[-25:],
                    open=price["PX_OPEN"].iloc[-25:],
                    high=price["PX_HIGH"].iloc[-25:],
                    low=price["PX_LOW"].iloc[-25:],
                    close=price["PX_LAST"].iloc[-25:],
                    name="Price",
                )
            )
            fig["layout"]["yaxis1"]["title"] = "Price"
            fig.add_trace(
                go.Bar(
                    name="Volume",
                    x=price.index[-25:],
                    y=price["VOLUME"].iloc[-25:],
                    showlegend=True,
                ),
                row=2,
                col=1,
            )
            fig["layout"]["yaxis2"]["title"] = "Volume"
            fig.add_trace(
                go.Scatter(
                    x=price.index[-25:],
                    y=avg_tr[-25:],
                    showlegend=True,
                    name="3d ATR",
                    line=dict(width=1),
                ),
                row=3,
                col=1,
            )
            fig["layout"]["yaxis3"]["title"] = "ATR"
            fig.update_xaxes(
                rangebreaks=[
                    dict(bounds=["sat", "mon"]),  # hide weekends
                    dict(values=["2015-12-25", "2016-01-01"]),  # hide Christmas and New Year's
                ]
            )
            fig.update_layout(
                title={"text": f"Range-Vol - {row['Ticker']}", "x": 0.5, "xanchor": "center"},
                barmode="stack",
                width=900,
                height=600,
                xaxis_rangeslider_visible=False,
            )
            figs.append(fig)
        table.figures_to_html(figs, filename=convert_path_to_linux(f"{folder_path}charts\\range_vol_gas.html"), task_name="Range-Vol and Divergence - Gas")
        if html_alert is None:
            html_alert = "No Range-Vol alert today<br>"
    else:
        html_alert = "No Range-Vol alert today<br>"

    df_all = pd.read_csv(ut.convert_path_to_linux(f"{folder_path}intraday_cum_vol.csv"))
    df_all = df_all.loc[df_all["Instr"].isin(["TTF", "NG", "HH"]), :]
    if len(df_all) > 0:
        df_all.set_index("Unnamed: 0", inplace=True)
        df_all.reset_index(drop=True, inplace=True)
        pct_format_rows = df_all.loc[df_all["chg type"] == "FLAT", :].index
        dec_format_rows = df_all.loc[df_all["chg type"] == "SPRD", :].index
        html_intravol = table.html_format(
            df=df_all,
            precision=2,
            hide_cols=["chg type"],
            format_column={
                "Instr": {"width": "20px", "text-align": "center", "right_border": True},
                "Ticker": {"width": "160px", "text-align": "center", "right_border": True},
                "Volume T": {
                    "width": "80px",
                    "text-align": "center",
                    "format": "{:.0%}",
                    "right_border": True,
                },
                "Change T": {"width": "80px", "text-align": "center"},
            },
            format_row={
                tuple(pct_format_rows): {"format": "{:.2%}", "columns": ["Change T"]},
                tuple(dec_format_rows): {"format": "{:.2f}", "columns": ["Change T"]},
            },
        )
    else:
        html_intravol = "No intraday cumulative volume spike<br>"

    try:
        htmls, _ = vol_alert(subset=["NG", "TTF", "EUA"])
        html_option = htmls[2]
    except Exception as e:
        print(f"Error in vol_alert: {e}")
        htmls = ["No option volume alert<br>", "No option IV alert<br>", "No option skew alert<br>"]
        html_option = htmls[1]

    imp_vs_real_vol = pd.read_csv(ut.convert_path_to_linux(f"{folder_path}imp_vs_real_vol_gas.csv"))
    imp_vs_real_vol.set_index("Unnamed: 0", inplace=True)
    if len(imp_vs_real_vol) > 0:
        html_imp_vs_real = table.html_format(
            df=imp_vs_real_vol,
            precision=2,
            hide_cols=["thr", "_thr"],
            format_column={
                "Name": {"width": "60px", "text-align": "left", "right_border": True},
                "3m IMP Chg": {"width": "100px", "text-align": "center", "right_border": False},
                "BE/SDMVATR": {"width": "100px", "text-align": "center", "right_border": True},
                "1m vs 8d": {"width": "100px", "text-align": "center"},
                "1m vs 8d zscore": {
                    "width": "100px",
                    "text-align": "center",
                    "right_border": True,
                    "highlight": ["1m vs 8d zscore", "thr", "_thr"],
                },
                "3m vs 13d": {"width": "100px", "text-align": "center"},
                "3m vs 13d zscore": {
                    "width": "100px",
                    "text-align": "center",
                    "right_border": True,
                    "highlight": ["3m vs 13d zscore", "thr", "_thr"],
                },
                "6m vs 34d": {"width": "100px", "text-align": "center"},
                "6m vs 34d zscore": {
                    "width": "100px",
                    "text-align": "center",
                    "right_border": True,
                    "highlight": ["6m vs 34d zscore", "thr", "_thr"],
                },
                "12m vs 50d": {"width": "100px", "text-align": "center"},
                "12m vs 50d zscore": {
                    "width": "100px",
                    "text-align": "center",
                    "right_border": True,
                    "highlight": ["12m vs 50d zscore", "thr", "_thr"],
                },
                "13d realized": {"width": "100px", "text-align": "center"},
                "13d realized zscore": {"width": "100px", "text-align": "center"},
            },
        )
    else:
        html_imp_vs_real = "No implied vs realized vol alert<br>"

    iv_chg = pd.read_csv(ut.convert_path_to_linux(f"{folder_path}imp_vs_real_vol_gas_ivchg.csv"))
    iv_chg.set_index("Unnamed: 0", inplace=True)
    if len(iv_chg) > 0:
        html_iv_chg = table.html_format(
            df=iv_chg,
            precision=2,
            hide_cols=["thr", "_thr"],
            format_column={
                "Name": {"width": "60px", "text-align": "left", "right_border": True},
                "Price Chg": {
                    "width": "100px",
                    "text-align": "center",
                    "right_border": False,
                    "highlight": ["Price * IV (1m)", "thr", "_thr"],
                },
                "Price * IV (1m)": {
                    "width": "100px",
                    "text-align": "center",
                    "right_border": False,
                    "highlight": ["Price * IV (1m)", "thr", "_thr"],
                },
                "1m IV Chg": {"width": "100px", "text-align": "center"},
                "1m IV Chg zscore": {
                    "width": "100px",
                    "text-align": "center",
                    "right_border": True,
                    "highlight": ["1m IV Chg zscore", "thr", "_thr"],
                },
                "3m IV Chg": {"width": "100px", "text-align": "center"},
                "3m IV Chg zscore": {
                    "width": "100px",
                    "text-align": "center",
                    "right_border": True,
                    "highlight": ["3m IV Chg zscore", "thr", "_thr"],
                },
                "6m IV Chg": {"width": "100px", "text-align": "center"},
                "6m IV Chg zscore": {
                    "width": "100px",
                    "text-align": "center",
                    "right_border": True,
                    "highlight": ["6m IV Chg zscore", "thr", "_thr"],
                },
                "12m IV Chg": {"width": "100px", "text-align": "center"},
                "12m IV Chg zscore": {
                    "width": "100px",
                    "text-align": "center",
                    "right_border": True,
                    "highlight": ["12m IV Chg zscore", "thr", "_thr"],
                },
            },
        )
    else:
        html_iv_chg = "No implied vol change alert<br>"

    df_diverge = pd.read_csv(ut.convert_path_to_linux(f"{folder_path}\\diverge_signal.csv"))
    df_diverge.set_index("Unnamed: 0", inplace=True)
    df_diverge = df_diverge.loc[
        df_diverge.index.isin(
            [
                "NG 1st Spread / NGA Comdty",
                "S:NGNG 2-6 Comdty / NGA Comdty",
                "TZT 1st spread / TZTA Comdty",
            ]
        ),
        :,
    ]
    if len(df_diverge) > 0:
        html_diverge = send_diverge_emails(df_diverge)

        figs_diverg = []
        for idx, row in df_diverge.iterrows():
            ticker0 = row["Ticker"].split("/")[0] + " Comdty"
            ticker1 = row["Ticker"].split("/")[1] + " Comdty"
            i0 = idx.split(" / ")[0]
            i1 = idx.split(" / ")[1]
            price0 = ts.read_csv(f"{folder_path}price_daily\\{ticker0}.csv", index_name="date")
            price1 = ts.read_csv(f"{folder_path}price_daily\\{ticker1}.csv", index_name="date")
            price0_chg = price0["PX_LAST"].diff()
            price1_chg = price1["PX_LAST"].diff()
            if ("S:XBXB" in i0 or "S:HOHO" in i0) and "S:COCO" in i1:
                diverg_ratio_3d = price0_chg.rolling(3).sum() * 0.42 / price1_chg.rolling(3).sum()
                diverg_ratio_8d = price0_chg.rolling(8).sum() * 0.42 / price1_chg.rolling(8).sum()
            else:
                diverg_ratio_3d = price0_chg.rolling(3).sum() / price1_chg.rolling(3).sum()
                diverg_ratio_8d = price0_chg.rolling(8).sum() / price1_chg.rolling(8).sum()
            price1 = price1.reindex(price0.index)
            diverg_ratio_3d = diverg_ratio_3d.reindex(price0.index)
            diverg_ratio_3d = diverg_ratio_3d.to_frame("Ratio_3d")
            diverg_ratio_8d = diverg_ratio_8d.reindex(price0.index)
            diverg_ratio_8d = diverg_ratio_8d.to_frame("Ratio_8d")
            fig = make_subplots(
                rows=3,
                cols=1,
                row_heights=[0.4, 0.4, 0.2],
                shared_xaxes=True,
                vertical_spacing=0.02,
            )
            fig.add_trace(
                go.Candlestick(
                    x=price0.index[-25:],
                    open=price0["PX_OPEN"].iloc[-25:],
                    high=price0["PX_HIGH"].iloc[-25:],
                    low=price0["PX_LOW"].iloc[-25:],
                    close=price0["PX_LAST"].iloc[-25:],
                    name=ticker0,
                )
            )
            fig.add_trace(
                go.Candlestick(
                    x=price1.index[-25:],
                    open=price1["PX_OPEN"].iloc[-25:],
                    high=price1["PX_HIGH"].iloc[-25:],
                    low=price1["PX_LOW"].iloc[-25:],
                    close=price1["PX_LAST"].iloc[-25:],
                    name=ticker1,
                ),
                row=2,
                col=1,
            )
            fig.add_trace(
                go.Scatter(
                    x=diverg_ratio_3d.index[-25:],
                    y=diverg_ratio_3d["Ratio_3d"].iloc[-25:],
                    showlegend=True,
                    name="Ratio_1d",
                    line=dict(width=1),
                ),
                row=3,
                col=1,
            )
            fig.add_trace(
                go.Scatter(
                    x=diverg_ratio_8d.index[-25:],
                    y=diverg_ratio_8d["Ratio_8d"].iloc[-25:],
                    showlegend=True,
                    name="Ratio_3d",
                    line=dict(width=1),
                ),
                row=3,
                col=1,
            )
            fig.update_xaxes(
                rangebreaks=[
                    dict(bounds=["sat", "mon"]),  # hide weekends
                    dict(values=["2015-12-25", "2016-01-01"]),  # hide Christmas and New Year's
                ]
            )
            fig.update_layout(
                title={"text": f"Divergence - {row['Ticker']}", "x": 0.5, "xanchor": "center"},
                barmode="stack",
                width=900,
                height=600,
                xaxis1=dict(rangeslider=dict(visible=False)),
                xaxis2=dict(rangeslider=dict(visible=False)),
            )
            figs_diverg.append(fig)
        table.figures_to_html(figs_diverg, filename=convert_path_to_linux(f"{folder_path}charts\\divergence_gas.html"))

    else:
        html_diverge = "No divergence alert today<br>"

    eu_gas = pd.read_csv(ut.convert_path_to_linux(f"{output_path}\\csvs\\ttf\\ttf_alert.csv"))
    eu_gas.set_index("Unnamed: 0", inplace=True)
    if len(eu_gas) > 0:
        html_eugas = table.html_format(
            df=eu_gas,
            precision=2,
            format_column={
                "Hub": {"width": "160px", "text-align": "center"},
                "Tenor": {"width": "100px", "text-align": "center"},
                "Price change": {"width": "100px", "text-align": "center"},
                "Price": {"width": "100px", "text-align": "center"},
                "Price 20d MA": {"width": "100px", "text-align": "center"},
                "3m z-score": {"width": "100px", "text-align": "center"},
            },
        )
    else:
        html_eugas = "No EU gas spread alert today"

    spark = pd.read_csv(ut.convert_path_to_linux(f"{output_path}\\csvs\\ttf\\spark_alert.csv"))
    spark.set_index("Unnamed: 0", inplace=True)
    if len(spark) > 0:
        html_spark = table.html_format(
            df=spark,
            precision=2,
            format_column={
                "Hub": {"width": "160px", "text-align": "center"},
                "Tenor": {"width": "100px", "text-align": "center"},
                "Price change": {"width": "100px", "text-align": "center"},
                "Price": {"width": "100px", "text-align": "center"},
                "Price 20d MA": {"width": "100px", "text-align": "center"},
                "3m z-score": {"width": "100px", "text-align": "center"},
            },
        )
    else:
        html_spark = "No spark alert today<br>"

    us_gas = pd.read_csv(ut.convert_path_to_linux(f"{output_path}\\csvs\\gas\\basis\\usgas_alert.csv"))
    us_gas.set_index("Unnamed: 0", inplace=True)
    if len(us_gas) > 0:
        html_usgas = table.html_format(
            df=us_gas,
            precision=2,
            format_column={
                "Hub": {"width": "160px", "text-align": "center"},
                "Tenor": {"width": "100px", "text-align": "center"},
                "Price change": {"width": "100px", "text-align": "center"},
                "Price": {"width": "100px", "text-align": "center"},
                "Price 20d MA": {"width": "100px", "text-align": "center"},
                "3m z-score": {"width": "100px", "text-align": "center"},
                "Latest Date": {"width": "100px", "text-align": "center"},
            },
        )
    else:
        html_usgas = "No US gas spread alert today<br>"

    tzt_oi = contract_oi_chart(active="TZTA Comdty")
    ng_oi = contract_oi_chart(active="NGA Comdty")
    risk_index_html = get_risk_index_html()
    html_table = [
        "<div style='font-family:Calibri;' >",
        f"Update time {now_ldn().strftime('%Y-%m-%d %H:%M')}. <br>",
        "<b>Risk Index</b><br>",
        risk_index_html,
        '<a href="https://elementcapital.atlassian.net/wiki/spaces/LO25/pages/1774685381/Range-Vol+and+Divergence">Description</a><br>',
        "<b>Range-Vol alert: </b><br>",
        html_alert,
        '<a href="{}charts\\range_vol_gas.html">RangeVol Charts</a><br>'.format(folder_path),
        "<br><b>Intraday cumulative volume alert: </b><br>",
        html_intravol,
        "<br><b>Option volume alert: </b><br>",
        html_option,
        '<a href="{}charts\\range_vol_option_gas.html">Option Charts</a><br>'.format(folder_path),
        "<br><b>Implied vs realized volatility alert: </b><br>",
        html_imp_vs_real,
        "<br><b>Implied volatility change alert: </b><br>",
        html_iv_chg,
        "<br><b>Spread and flat price divergence alert: </b><br>",
        html_diverge,
        '<a href="{}charts\\divergence_gas.html">Divergence Charts</a><br>'.format(folder_path),
        "<br><b>EU gas spread alert: </b><br>",
        html_eugas,
        '<a href="{}\\ttf\\links\\ttf_alert_charts.html">EU Gas Charts</a><br>'.format(html_path),
        "<br><b>Spark alert: </b><br>",
        html_spark,
        '<a href="{}\\ttf\\links\\spark_alert_charts.html">Spark Charts</a><br>'.format(html_path),
        "<br><b>US gas spread alert: </b><br>",
        html_usgas,
        '<a href="{}\\gas\\links\\usgas_alert_charts.html">US Gas Charts</a><br>'.format(html_path),
        "<br><b>TZT and NG OI change: </b><br>",
        tzt_oi,
        ng_oi,
    ]

    table.figures_to_html(
        html_table,
        filename=convert_path_to_linux(f"{html_path}\\cross_cmds\\market_scan_alert_gas.html"),
        task_name=report_name,
    )
    send_email(
        send_to=send_to,
        subject="Range-Vol and Divergence - Gas",
        body=[
            html_table,
            "<br><br>This is automated email sent at {:s}. <br><br>".format(now_ldn().strftime("%Y-%m-%d %H:%M")),
        ],
        attachments=[
            ut.convert_path_to_linux(f"{html_path}\\ttf\\eu_gas_basis.html"),
            ut.convert_path_to_linux(f"{html_path}\\gas\\us_gas_basis.html"),
        ],
        html_path=f"{html_path}\\cross_cmds\\market_scan_alert_gas.html",
    )


def get_alert_email_html(df):
    df.fillna(value=0, inplace=True)
    df_ = df.dropna(axis=0)
    df_ = df_.loc[~df_["Instr"].isin(["GR", "SFT", "Bulks"]), :]
    df_["Ticker"] = [x.split(" ")[0] for x in df_["Ticker"]]
    df_.reset_index(drop=True, inplace=True)
    if not df.empty:
        pct_format_rows = df_.loc[df_["chg type"] == "FLAT", :].index
        dec_format_rows = df_.loc[df_["chg type"] == "SPRD", :].index
        html_df = table.html_format(
            df_,
            precision=2,
            hide_cols=["chg type"],
            format_column={
                "Instr": {"width": "20px", "text-align": "center", "right_border": True},
                "Ticker": {"width": "160px", "text-align": "center", "right_border": True},
                "Follow Through": {
                    "width": "60px",
                    "text-align": "center",
                    "highlight": ["Follow Through"],
                    "right_border": True,
                },
                "Change T": {"width": "60px", "text-align": "center"},
                "Range T": {"width": "60px", "text-align": "center", "format": "{:.0%}"},
                "Volume T": {
                    "width": "60px",
                    "text-align": "center",
                    "format": "{:.0%}",
                    "right_border": True,
                },
                "Change T-1": {"width": "60px", "text-align": "center"},
                "Range T-1": {"width": "60px", "text-align": "center", "format": "{:.0%}"},
                "Volume T-1": {"width": "60px", "text-align": "center", "format": "{:.0%}"},
            },
            format_row={
                tuple(pct_format_rows): {"format": "{:.2%}", "columns": ["Change T", "Change T-1"]},
                tuple(dec_format_rows): {"format": "{:.2f}", "columns": ["Change T", "Change T-1"]},
            },
        )
        return html_df
    else:
        return None


def send_diverge_emails(df):
    df_ = df.dropna()
    df_.index.name = "Spread/Flat"
    df_.reset_index(inplace=True)
    for idx, row in df_.iterrows():
        name_list = row["Spread/Flat"].split(" ")
        for idx1, i in enumerate(name_list):
            if i == "Comdty":
                name_list.pop(idx1)
        df_.loc[idx, "Spread/Flat"] = " ".join(name_list)
    html_df = table.html_format(
        df_,
        precision=2,
        format_column={
            "Spread/Flat": {"width": "240px", "text-align": "center", "right_border": True},
            "Ticker": {"width": "160px", "text-align": "center", "right_border": True},
            "3d divergence": {"width": "100px", "text-align": "center"},
            "8d divergence": {"width": "100px", "text-align": "center"},
            "Sprd change": {"width": "100px", "text-align": "center"},
            "Flat change": {"width": "100px", "text-align": "center"},
        },
    )
    return html_df


def bar_chart(data, data1=None, title=None, **kwargs):
    file_path = kwargs.get("file_path", None)
    y1_axis_title = kwargs.get("y1_axis_title", None)
    y2_axis_title = kwargs.get("y2_axis_title", None)
    y3_axis_title = kwargs.get("y3_axis_title", None)
    y4_axis_title = kwargs.get("y4_axis_title", None)
    y5_axis_title = kwargs.get("y5_axis_title", None)
    y6_axis_title = kwargs.get("y6_axis_title", None)
    x_axis_title = kwargs.get("x_axis_title", None)
    bar_columns = kwargs.get("bar_columns", None)
    line_columns = kwargs.get("line_columns", None)
    data2 = kwargs.get("data2", None)
    data3 = kwargs.get("data3", None)
    second_y = kwargs.get("second_y", False)

    if data1 is None:
        if second_y:
            fig = make_subplots(specs=[[{"secondary_y": True}]])
        else:
            fig = make_subplots()
    elif data1 is not None and data2 is None:
        if second_y:
            fig = make_subplots(
                rows=2,
                cols=1,
                row_heights=[0.7, 0.3],
                shared_xaxes=True,
                vertical_spacing=0.02,
                specs=[[{"secondary_y": True}], [{"secondary_y": False}]],
            )
        else:
            fig = make_subplots(rows=2, cols=1, row_heights=[0.7, 0.3], shared_xaxes=True, vertical_spacing=0.02)
    else:
        if second_y:
            fig = make_subplots(
                rows=3,
                cols=1,
                row_heights=[0.6, 0.2, 0.2],
                shared_xaxes=True,
                vertical_spacing=0.02,
                specs=[[{"secondary_y": True}], [{"secondary_y": True}], [{"secondary_y": True}]],
            )
        else:
            fig = make_subplots(
                rows=3,
                cols=1,
                row_heights=[0.6, 0.2, 0.2],
                shared_xaxes=True,
                vertical_spacing=0.02,
            )

    if bar_columns is not None:
        for col in bar_columns:
            fig.add_trace(go.Bar(name=col, x=data.index, y=data[col].values, showlegend=True))
    else:
        for col in data.columns:
            fig.add_trace(go.Bar(name=col, x=data.index, y=data[col].values, showlegend=True))

    if line_columns is not None:
        if isinstance(line_columns, list):
            for col in line_columns:
                fig.add_trace(go.Scatter(x=data.index, y=data[col], showlegend=True, name=col, line=dict(width=1)))
        elif isinstance(line_columns, dict):
            for col, val in line_columns.items():
                fig.add_trace(
                    go.Scatter(
                        x=data.index,
                        y=data[col],
                        showlegend=True,
                        name=col,
                        line=dict(width=1 + val),
                    ),
                    secondary_y=val,
                )
    fig["layout"]["yaxis1"]["title"] = y1_axis_title
    if y2_axis_title is not None:
        fig["layout"]["yaxis2"]["title"] = y2_axis_title

    if data1 is not None:
        fig.add_trace(
            go.Bar(x=data.index, y=data1.iloc[:, 0], showlegend=True, name=data1.columns[0]),
            row=2,
            col=1,
        )
        fig["layout"]["yaxis3"]["title"] = y3_axis_title
    if data2 is not None:
        fig.add_trace(
            go.Bar(name=data2.columns[0], x=data.index, y=data2.iloc[:, 0].values, showlegend=True),
            row=3,
            col=1,
        )
        fig["layout"]["yaxis5"]["title"] = y5_axis_title
    if data3 is not None:
        fig.add_trace(
            go.Scatter(
                x=data.index,
                y=data3.iloc[:, 0],
                showlegend=True,
                name=data3.columns[0],
                line=dict(width=2),
            ),
            row=3,
            col=1,
            secondary_y=True,
        )
        fig["layout"]["yaxis6"]["title"] = y6_axis_title

    fig.update_xaxes(
        rangebreaks=[
            dict(bounds=["sat", "mon"]),  # hide weekends
            dict(values=["2015-12-25", "2016-01-01"]),  # hide Christmas and New Year's
        ]
    )

    fig.update_layout(title={"text": title, "x": 0.5, "xanchor": "center"}, barmode="stack", width=900, height=600)

    if file_path is not None:
        py.offline.plot(fig, auto_open=False, filename=file_path)
    return fig


def spot_price_alert():
    log.info("Running Spot Price Alert")
    db = partial(mongo_table, db="data", table="platts", url=url, pk=["name", "ticker", "platts_ticker"])
    ticker_dict = get_tickers()
    linked_sprd_dict = {}
    sp_ticker_dict = {x[1]: y for x, y in ticker_dict.items() if x[1] in list(itertools.chain.from_iterable(spots_dict.values()))}
    for key, val in spots_dict.items():
        if val[3] is not None and val[3] not in list(linked_sprd_dict.keys()):
            linked_sprd_dict[val[3]] = sp_ticker_dict[val[3]]
    ticker_list = [v[0] for k, v in linked_sprd_dict.items()]
    linked_price = bbg.bdh(
        ticker_list,
        ["PX_LAST"],
        sdate=today() - dt.timedelta(days=7),
        edate=today() - dt.timedelta(days=1),
    )
    if isinstance(linked_price.columns, pd.MultiIndex):
        linked_price = linked_price[ticker_list]
    linked_price_chg = linked_price.diff().iloc[-1]
    linked_price_dict = {}
    for key, val in spots_dict.items():
        if val[3] is not None:
            linked_price_dict[key] = linked_price_chg[linked_sprd_dict[val[3]][0]]
        else:
            linked_price_dict[key] = np.nan

    alert_dict = {}
    figs = []

    for key, val in spots_dict.items():
        print(key)
        if key in ["TD22", "TD25", "TC2", "TC5"]:
            if key == "TD22":
                print("break")
            cur_m = (today() + relativedelta(day=1)).strftime("%#m/%#d/%Y") if sys.platform.startswith("win") else (today() + relativedelta(day=1)).strftime("%-m/%-d/%Y")
            price = sql.read_sql(f"Select * from ICE_OIL where CONTRACT='{val[0]}' and STRIP='{cur_m}' order by TRADE_DATE")
            price.drop_duplicates(inplace=True)
            price.set_index("TRADE_DATE", inplace=True)
            price.index = pd.to_datetime(price.index)
            price = price["SETTLEMENT_PRICE"]
            price_chg = price.diff()
        else:
            if val[4]:
                try:
                    price = pyg.get_data(db, ticker=val[0])
                except:
                    price = bbg.bdh(val[0], ["PX_LAST"], sdate=today() - BDay(80), edate=today() - BDay(1))
                    if isinstance(price.columns, pd.MultiIndex):
                        price = price[val[0]]
                if price.empty:
                    continue
                price = price.iloc[:, 0]
                price_chg = price.diff()
            else:
                try:
                    price0 = pyg.get_data(db, ticker=val[0])
                except:
                    price0 = bbg.bdh(val[0], ["PX_LAST"], sdate=today() - BDay(80), edate=today() - BDay(1))
                    if isinstance(price0.columns, pd.MultiIndex):
                        price0 = price0[val[0]]
                price0 = price0.iloc[:, 0]
                try:
                    price1 = pyg.get_data(db, ticker=val[1])
                except:
                    price1 = bbg.bdh(val[1], ["PX_LAST"], sdate=today() - BDay(80), edate=today() - BDay(1))
                    if isinstance(price1.columns, pd.MultiIndex):
                        price1 = price1[val[0]]
                price1 = price1.iloc[:, 0]
                price = price0 - price1
                price_chg = price.diff()
        price_mean = price.rolling(20).mean()
        price_std = price_chg.rolling(20).std()
        zscore = price_chg[-1] / price_std[-1]
        if zscore > 1.5 or zscore < -1.5:
            alert_dict[key] = [
                val[2],
                price_chg[-1],
                linked_price_dict[key],
                price[-1],
                price_mean[-1],
                zscore,
                dt.datetime.strftime(price_chg.index[-1], "%Y-%m-%d"),
            ]
            figs.append(chart.line_chart(df=(price.iloc[-65:]).to_frame("Price"), title=f"{key}"))

    if len(alert_dict) > 0:
        alert_df = pd.DataFrame.from_dict(alert_dict, orient="index")
        alert_df.columns = [
            "Market",
            "Price change",
            "Linked sprd chg",
            "Price",
            "Price 20d MA",
            "20d z-score",
            "Latest Date",
        ]
        alert_df.index.name = "Spot"
        alert_df.reset_index(inplace=True)
        table.figures_to_html(figs, filename=convert_path_to_linux(f"{folder_path}charts\\spot_alert.html"))
        return alert_df
    else:
        return pd.DataFrame()


def crack_price_alert():
    log.info("Running Crack Price Alert")
    crack_dict = {
        "EBOB Crack": ["AEB", "AEO", "EU"],
        "GO Crack": ["ULD", "S:QSQS Comdty", "EU"],
        "Naphtha Crack": ["NBB", "NEC", "EU"],
        "FO 3.5 Crack": ["BOA", "BAR", "EU"],
        "XBCL": ["S:XBCL Comdty", "S:XBXB Comdty", "US"],
        "HOCL": ["S:HOCL Comdty", "S:HOHO Comdty", "US"],
    }

    alert_dict = {}
    figs = []

    for key, val in crack_dict.items():
        print(key)
        if key in ["XBCL", "HOCL"]:
            live_cl = bbg.live_contract("CLA Comdty")
            live_cl_next = bbg.live_contract("CLA Comdty", seq=1)
            crack_ticker = f"{val[0].split(' ')[0]} {live_cl['m']}{str(int(live_cl['y']))[-2:]}-{live_cl['m']}{str(int(live_cl['y']))[-2:]} Comdty"
            spread_ticker = f"{val[1].split(' ')[0]} {live_cl['m']}{str(int(live_cl['y']))[-2:]}-{live_cl_next['m']}{str(int(live_cl_next['y']))[-2:]} Comdty"
            price_crack = bbg.bdh(crack_ticker, ["PX_LAST"], sdate=today() - BDay(80), edate=today() - BDay(1))
            price_spread = bbg.bdh(spread_ticker, ["PX_LAST"], sdate=today() - BDay(80), edate=today() - BDay(1))
        else:
            if sys.platform.startswith("win"):
                strip = (today() + relativedelta(day=31) + relativedelta(days=1)).strftime("%#m/%#d/%Y")
                strip_next = (today() + relativedelta(day=31) + relativedelta(days=1) + relativedelta(months=1)).strftime("%#m/%#d/%Y")
            else:
                strip = (today() + relativedelta(day=31) + relativedelta(days=1)).strftime("%-m/%-d/%Y")
                strip_next = (today() + relativedelta(day=31) + relativedelta(days=1) + relativedelta(months=1)).strftime("%-m/%-d/%Y")
            price_crack = sql.read_sql(f"select * from ICE_OIL_{val[0]} where STRIP='{strip}' order by TRADE_DATE")
            if len(val[1]) > 3:
                near_m = today() + relativedelta(day=31) + relativedelta(days=1)
                far_m = today() + relativedelta(day=31) + relativedelta(days=1) + relativedelta(months=1)
                spread_ticker = f"{val[1].split(' ')[0]} {tk.month_int2str[near_m.month]}{str(near_m.year)[-2:]}-{tk.month_int2str[far_m.month]}{str(far_m.year)[-2:]} Comdty"
                price_spread = bbg.bdh(spread_ticker, ["PX_LAST"], sdate=today() - BDay(80), edate=today() - BDay(1))
            else:
                price0 = sql.read_sql(f"select * from ICE_OIL_{val[1]} where STRIP='{strip}' order by TRADE_DATE")
                price1 = sql.read_sql(f"select * from ICE_OIL_{val[1]} where STRIP='{strip_next}' order by TRADE_DATE")

            def format_ice_price(price):
                price1 = (price.set_index("TRADE_DATE"))["SETTLEMENT_PRICE"].to_frame("PX_LAST")
                price1 = price1[~price1.index.duplicated(keep='last')]
                price1.index = pd.to_datetime(price1.index)
                return price1

            price_crack = format_ice_price(price_crack)
            if len(val[1]) > 3:
                pass
            else:
                price0 = format_ice_price(price0)
                price1 = format_ice_price(price1)
                price_spread = (price0 - price1.reindex(price0.index)).dropna()

        price_crack = price_crack.iloc[:, 0]
        price_spread = price_spread.iloc[:, 0]
        price_crack_chg = price_crack.diff()
        price_spread_chg = price_spread.diff()

        price_mean = price_crack.rolling(20).mean()
        price_std = price_crack_chg.rolling(20).std()
        zscore = price_crack_chg[-1] / price_std[-1]
        if zscore > 1.5 or zscore < -1.5:
            alert_dict[key] = [
                val[2],
                price_crack_chg[-1],
                price_spread_chg[-1],
                price_crack[-1],
                price_mean[-1],
                zscore,
                dt.datetime.strftime(price_crack_chg.index[-1], "%Y-%m-%d"),
            ]
            figs.append(chart.line_chart(df=(price_crack.iloc[-65:]).to_frame("Price"), title=f"{key}"))

    if len(alert_dict) > 0:
        alert_df = pd.DataFrame.from_dict(alert_dict, orient="index")
        alert_df.columns = [
            "Market",
            "Price change",
            "Linked sprd chg",
            "Price",
            "Price 20d MA",
            "20d z-score",
            "Latest Date",
        ]
        alert_df.index.name = "Crack"
        alert_df.reset_index(inplace=True)
        table.figures_to_html(figs, filename=convert_path_to_linux(f"{folder_path}charts\\crack_alert.html"))
        return alert_df
    else:
        return pd.DataFrame()


def physical_price_alert():
    log.info("Running Physical Price Alert")

    from ecm.cmds.config import output_path
    total_crude = ts.read_csv(f"{output_path}\\csvs\\oil\\global_physical_crude_detail.csv", index_name="date")
    eu_crude = ts.read_csv(f"{output_path}\\csvs\\oil\\physical_crude_detail_eu.csv", index_name="date")
    us_crude = ts.read_csv(f"{output_path}\\csvs\\oil\\physical_crude_detail_us.csv", index_name="date")
    asia_crude = ts.read_csv(f"{output_path}\\csvs\\oil\\physical_crude_detail_asia.csv", index_name="date")
    total_gasoil = ts.read_csv(f"{output_path}\\csvs\\oil\\global_physical_gasoil_detail.csv", index_name="date")
    eu_gasoil = ts.read_csv(f"{output_path}\\csvs\\oil\\eu_physical_gasoil_detail.csv", index_name="date")
    us_gasoil = ts.read_csv(f"{output_path}\\csvs\\oil\\us_physical_gasoil_detail.csv", index_name="date")
    asia_gasoil = ts.read_csv(f"{output_path}\\csvs\\oil\\asia_physical_gasoil_detail.csv", index_name="Unnamed: 0")
    total_gasoline = ts.read_csv(f"{output_path}\\csvs\\oil\\global_physical_gasoline_detail.csv", index_name="date")
    eu_gasoline = ts.read_csv(f"{output_path}\\csvs\\oil\\eu_physical_gasoline_detail.csv", index_name="date")
    us_gasoline = ts.read_csv(f"{output_path}\\csvs\\oil\\us_physical_gasoline_detail.csv", index_name="date")
    asia_gasoline = ts.read_csv(f"{output_path}\\csvs\\oil\\asia_physical_gasoline_detail.csv", index_name="date")
    swaps = pd.concat(
        [
            total_crude.iloc[:, 0],
            eu_crude.drop(eu_crude.columns[1], axis=1),
            us_crude.drop(us_crude.columns[1], axis=1),
            asia_crude.drop(asia_crude.columns[1], axis=1),
            total_gasoil.iloc[:, 0],
            eu_gasoil.drop(eu_gasoil.columns[1], axis=1),
            us_gasoil.drop(us_gasoil.columns[1], axis=1),
            asia_gasoil.drop(asia_gasoil.columns[1], axis=1),
            total_gasoline.iloc[:, 0],
            eu_gasoline.drop(eu_gasoline.columns[1], axis=1),
            us_gasoline.drop(us_gasoline.columns[1], axis=1),
            asia_gasoline.drop(asia_gasoline.columns[1], axis=1),
        ],
        axis=1,
    )
    swaps.fillna(method="ffill", inplace=True)
    swaps_dict = {}
    swaps_dict["Crd-P-Global"] = ["Global Crude", 0]
    for i in eu_crude.columns:
        swaps_dict[i] = ["EU Crude", 1]
    for i in us_crude.columns:
        swaps_dict[i] = ["US Crude", 2]
    for i in asia_crude.columns:
        swaps_dict[i] = ["Asia Crude", 3]
    swaps_dict["GO-P-Global"] = ["Global Gasoil", 4]
    for i in eu_gasoil.columns:
        swaps_dict[i] = ["EU Gasoil", 5]
    for i in us_gasoil.columns:
        swaps_dict[i] = ["US Gasoil", 6]
    for i in asia_gasoil.columns:
        swaps_dict[i] = ["Asia Gasoil", 7]
    swaps_dict["MOGAS-P-Global"] = ["Global Gasoline", 8]
    for i in eu_gasoline.columns:
        swaps_dict[i] = ["EU Gasoline", 9]
    for i in us_gasoline.columns:
        swaps_dict[i] = ["US Gasoline", 10]
    for i in asia_gasoline.columns:
        swaps_dict[i] = ["Asia Gasoline", 11]
    figs = []
    product_figs = []
    alert_dict = {}
    product_alert_dict = {}

    swaps_chg = swaps.diff()
    price12 = pd.DataFrame()
    price12[total_crude.columns[1]] = total_crude.iloc[:, 1]
    price12[eu_crude.columns[1]] = eu_crude.iloc[:, 1]
    price12[us_crude.columns[1]] = us_crude.iloc[:, 1]
    price12[asia_crude.columns[1]] = asia_crude.iloc[:, 1]
    price12[total_gasoil.columns[1]] = total_gasoil.iloc[:, 1]
    price12[eu_gasoil.columns[1]] = eu_gasoil.iloc[:, 1]
    price12[us_gasoil.columns[1]] = us_gasoil.iloc[:, 1]
    price12[asia_gasoil.columns[1]] = asia_gasoil.iloc[:, 1]
    price12[total_gasoline.columns[1]] = total_gasoline.iloc[:, 1]
    price12[eu_gasoline.columns[1]] = eu_gasoline.iloc[:, 1]
    price12[us_gasoline.columns[1]] = us_gasoline.iloc[:, 1]
    price12[asia_gasoline.columns[1]] = asia_gasoline.iloc[:, 1]
    price12 = price12.reindex(swaps.index, method="ffill")
    price12.fillna(method="ffill", inplace=True)
    price12_chg = price12.diff()

    for i in swaps.columns:
        price = swaps[i]
        price_chg = swaps_chg[i]
        price_mean = price.rolling(20).mean()
        price_std = price_chg.rolling(20).std()
        zscore = price_chg[-1] / price_std[-1]
        if zscore > 1.5 or zscore < -1.5:
            if swaps_dict[i][1] in [0, 1, 2, 3]:
                alert_dict[i] = [
                    swaps_dict[i][0],
                    price_chg[-1],
                    price12_chg.iloc[-1, swaps_dict[i][1]],
                    price[-1],
                    price_mean[-1],
                    zscore,
                    dt.datetime.strftime(price_chg.index[-1], "%Y-%m-%d"),
                ]
                figs.append(chart.line_chart(df=(price.iloc[-65:]).to_frame("Price"), title=f"{i}"))
            else:
                product_alert_dict[i] = [
                    swaps_dict[i][0],
                    price_chg[-1],
                    price12_chg.iloc[-1, swaps_dict[i][1]],
                    price[-1],
                    price_mean[-1],
                    zscore,
                    dt.datetime.strftime(price_chg.index[-1], "%Y-%m-%d"),
                ]
                product_figs.append(chart.line_chart(df=(price.iloc[-65:]).to_frame("Price"), title=f"{i}"))
    if len(alert_dict) > 0:
        alert_df = pd.DataFrame.from_dict(alert_dict, orient="index")
        alert_df.columns = [
            "Market",
            "Price change",
            "Linked sprd chg",
            "Price",
            "Price 20d MA",
            "20d z-score",
            "Latest Date",
        ]
        alert_df.index.name = "Crude"
        alert_df.reset_index(inplace=True)
        table.figures_to_html(figs, filename=convert_path_to_linux(f"{folder_path}charts\\physical_alert.html"))
    else:
        alert_df = pd.DataFrame()
    if len(product_alert_dict) > 0:
        product_alert_df = pd.DataFrame.from_dict(product_alert_dict, orient="index")
        product_alert_df.columns = [
            "Market",
            "Price change",
            "Linked sprd chg",
            "Price",
            "Price 20d MA",
            "20d z-score",
            "Latest Date",
        ]
        product_alert_df.index.name = "Product"
        product_alert_df.reset_index(inplace=True)
        table.figures_to_html(product_figs, filename=convert_path_to_linux(f"{folder_path}charts\\physical_product_alert.html"))
    else:
        product_alert_df = pd.DataFrame()
    return alert_df, product_alert_df


def swap_price_alert():
    log.info("Running Swap Price Alert")

    from ecm.cmds.config import output_path
    total_swap = ts.read_csv(f"{output_path}\\csvs\\oil\\global_swap_detail.csv", index_name="TRADE_DATE")
    eu_swap = ts.read_csv(f"{output_path}\\csvs\\oil\\eu_swap_detail.csv", index_name="TRADE_DATE")
    us_swap = ts.read_csv(f"{output_path}\\csvs\\oil\\us_swap_detail.csv", index_name="TRADE_DATE")
    asia_swap = ts.read_csv(f"{output_path}\\csvs\\oil\\asia_swap_detail.csv", index_name="TRADE_DATE")
    price26 = ts.read_csv(f"{output_path}\\csvs\\oil\\global_crude_26.csv", index_name="TRADE_DATE")
    total_gasoil_swap = ts.read_csv(f"{output_path}\\csvs\\oil\\global_gasoil_swap_detail.csv", index_name="TRADE_DATE")
    eu_gasoil_swap = ts.read_csv(f"{output_path}\\csvs\\oil\\eu_gasoil_swap_detail.csv", index_name="TRADE_DATE")
    asia_gasoil_swap = ts.read_csv(f"{output_path}\\csvs\\oil\\asia_gasoil_swap_detail.csv", index_name="TRADE_DATE")
    total_gasoline_swap = ts.read_csv(f"{output_path}\\csvs\\oil\\global_gasoline_swap_detail.csv", index_name="TRADE_DATE")
    eu_gasoline_swap = ts.read_csv(f"{output_path}\\csvs\\oil\\eu_gasoline_swap_detail.csv", index_name="TRADE_DATE")
    asia_gasoline_swap = ts.read_csv(f"{output_path}\\csvs\\oil\\asia_gasoline_swap_detail.csv", index_name="TRADE_DATE")
    swaps = pd.concat(
        [
            total_swap.iloc[:, 0],
            eu_swap,
            us_swap,
            asia_swap,
            total_gasoil_swap.iloc[:, 0],
            eu_gasoil_swap.drop(eu_gasoil_swap.columns[1], axis=1),
            asia_gasoil_swap,
            total_gasoline_swap.drop(total_gasoline_swap.columns[1], axis=1),
            eu_gasoline_swap.iloc[:, 2:],
            asia_gasoline_swap.iloc[:, 2:],
        ],
        axis=1,
    )
    swaps.fillna(method="ffill", inplace=True)
    swaps_dict = {}
    swaps_dict["Crd-S-Global"] = ["Global Crude", 0]
    for i in eu_swap.columns:
        swaps_dict[i] = ["EU Crude", 1]
    for i in us_swap.columns:
        swaps_dict[i] = ["US Crude", 2]
    for i in asia_swap.columns:
        swaps_dict[i] = ["Asia Crude", 3]
    swaps_dict["GO-S-Global"] = ["Global Gasoil", 4]
    for i in eu_gasoil_swap.columns:
        swaps_dict[i] = ["EU Gasoil", 5]
    for i in asia_gasoil_swap.columns:
        swaps_dict[i] = ["Asia Gasoil", 6]
    swaps_dict["MOGAS-S-Global"] = ["Global Gasoline", 8]
    for i in eu_gasoline_swap.columns:
        swaps_dict[i] = ["EU Gasoline", 7]
    for i in asia_gasoline_swap.columns:
        swaps_dict[i] = ["Asia Gasoline", 8]
    figs = []
    product_figs = []
    alert_dict = {}
    product_alert_dict = {}

    swaps_chg = swaps.diff()
    price26[total_gasoil_swap.columns[1]] = total_gasoil_swap.iloc[:, 1]
    price26[eu_gasoil_swap.columns[1]] = eu_gasoil_swap.iloc[:, 1]
    price26[asia_gasoil_swap.columns[1]] = asia_gasoil_swap.iloc[:, 1]
    price26[total_gasoline_swap.columns[1]] = total_gasoline_swap.iloc[:, 1]
    price26[eu_gasoline_swap.columns[1]] = eu_gasoline_swap.iloc[:, 1]
    price26[asia_gasoline_swap.columns[1]] = asia_gasoline_swap.iloc[:, 1]
    price26 = price26.reindex(swaps.index, method="ffill")
    price26.fillna(method="ffill", inplace=True)
    price26_chg = price26.diff()

    for i in swaps.columns:
        price = swaps[i]
        price_chg = swaps_chg[i]
        if isinstance(price, pd.DataFrame):
            price = price.iloc[:, 0]
            price_chg = price_chg.iloc[:, 0]
        price_mean = price.rolling(20).mean()
        price_std = price_chg.rolling(20).std()
        zscore = price_chg[-1] / price_std[-1]
        if zscore > 1.5 or zscore < -1.5:
            if swaps_dict[i][1] in [0, 1, 2, 3]:
                alert_dict[i] = [
                    swaps_dict[i][0],
                    price_chg[-1],
                    price26_chg.iloc[-1, swaps_dict[i][1]],
                    price[-1],
                    price_mean[-1],
                    zscore,
                    dt.datetime.strftime(price_chg.index[-1], "%Y-%m-%d"),
                ]
                figs.append(chart.line_chart(df=(price.iloc[-65:]).to_frame("Price"), title=f"{i}"))
            else:
                product_alert_dict[i] = [
                    swaps_dict[i][0],
                    price_chg[-1],
                    price26_chg.iloc[-1, swaps_dict[i][1]],
                    price[-1],
                    price_mean[-1],
                    zscore,
                    dt.datetime.strftime(price_chg.index[-1], "%Y-%m-%d"),
                ]
                product_figs.append(chart.line_chart(df=(price.iloc[-65:]).to_frame("Price"), title=f"{i}"))
    if len(alert_dict) > 0:
        alert_df = pd.DataFrame.from_dict(alert_dict, orient="index")
        alert_df.columns = [
            "Market",
            "Price change",
            "Linked sprd chg",
            "Price",
            "Price 20d MA",
            "20d z-score",
            "Latest Date",
        ]
        alert_df.index.name = "Crude"
        alert_df.reset_index(inplace=True)
        table.figures_to_html(figs, filename=convert_path_to_linux(f"{folder_path}charts\\swap_alert.html"))
    else:
        alert_df = pd.DataFrame()
    if len(product_alert_dict) > 0:
        product_alert_df = pd.DataFrame.from_dict(product_alert_dict, orient="index")
        product_alert_df.columns = [
            "Market",
            "Price change",
            "Linked sprd chg",
            "Price",
            "Price 20d MA",
            "20d z-score",
            "Latest Date",
        ]
        product_alert_df.index.name = "Product"
        product_alert_df.reset_index(inplace=True)
        table.figures_to_html(product_figs, filename=convert_path_to_linux(f"{folder_path}charts\\product_swap_alert.html"))
    else:
        product_alert_df = pd.DataFrame()
    return alert_df, product_alert_df


risk_sharpe_ticker_weights = {
    "SHSZ300 Index": 0.03,
    "KWEB US Equity": 0.02,
    "ESA Index": 0.1,
    "GXA Index": 0.1,
    "MSZZCYDE Index": 0.05,
    "DXY Curncy": -0.2,
    "COA Comdty": 0.13,
    "LMCADS03 Comdty": 0.12,
    "USGG12M Index": 0.03,
    "USGG2YR Index": 0.03,
    "USGG5YR Index": 0.03,
    "USGG10YR Index": 0.03,
    "USGG30YR Index": 0.03,
    "GECU10YR Index": 0.1,
}
risk_sharpe_eqts = ["SHSZ300 Index", "KWEB US Equity", "ESA Index", "GXA Index", "MSZZCYDE Index"]
risk_sharpe_ccy = ["DXY Curncy"]
risk_sharpe_cmds = ["COA Comdty", "LMCADS03 Comdty"]
risk_sharpe_trsy = [
    "USGG12M Index",
    "USGG2YR Index",
    "USGG5YR Index",
    "USGG10YR Index",
    "USGG30YR Index",
    "GECU10YR Index",
]


def get_risk_index_html():
    def get_rank(data):
        data_ = np.abs(data)
        rank_ = data_.rank(ascending=False)
        return rank_

    def table_format(df, header):
        return table.html_format(
            df=df,
            header=header,
            hide_cols=["abs", "new"],
            format_column={
                "Instr": {"width": "40px", "text-align": "center"},
                "Ticker": {"width": "120px", "text-align": "center"},
                "Shrp 3D Live": {"width": "80px", "text-align": "center"},
                "Shrp 3D COB": {"width": "80px", "text-align": "center"},
                "Shrp 22D": {"width": "80px", "text-align": "center"},
                "Signal": {"width": "80px", "text-align": "center"},
                "new": {"highlight": ["Ticker", "new"]},
            },
        )

    def format_table_html_component(df, header):
        df[["KPI", "KPI t-1", "KPI Index"]] = df[["KPI", "KPI t-1", "KPI Index"]] * 100
        return table.html_format(
            df=df,
            header=header,
            hide_cols=["abs", "new", "Shrp 3D COB 1", "Shrp 3D COB 2", "Shrp 3D COB 3"],
            format_column={
                "Shrp 3D Live": {"width": "80px", "text-align": "center"},
                "Shrp 3D COB": {"width": "80px", "text-align": "center"},
                "Shrp 22D": {"width": "80px", "text-align": "center"},
                "KPI": {"width": "80px", "text-align": "center", "format": "{:,.0f}%"},
                "KPI Index": {"width": "80px", "text-align": "center", "format": "{:,.0f}%"},
                "KPI t-1": {"width": "80px", "text-align": "center", "format": "{:,.0f}%"},
                "Index": {"width": "80px", "text-align": "center", "format": "{:,.2f}"},
                "new": {"highlight": ["Ticker", "new"]},
            },
        )

    c_path = f"{folder_path}\\market_scan_risk_componenets.html"

    sr = pd.read_csv(convert_path_to_linux(f"{folder_path}sharpe_ratio.csv"))
    sr.set_index("Unnamed: 0", inplace=True)
    sr.dropna(axis=0, inplace=True)
    sr["abs"] = np.abs(sr["Shrp 3D COB"])
    sr = sr.sort_values("abs", ascending=False)
    sr["Ticker"] = [x for x in sr["_plot_link"]]
    sr.drop(columns=["_plot_link"])
    sr_risk_sharpe = sr.loc[sr["Instr"].isin(["RISK"]), :]
    num_cols = sr_risk_sharpe.select_dtypes(include=["number"]).columns.tolist()
    sr_risk_sharpe["Names"] = sr_risk_sharpe.Ticker.str.split(".html").apply(lambda x: x[0].split(r"\\")[-1])

    sr_weights_eq_weights_sum = sum([risk_sharpe_ticker_weights[x] for x in risk_sharpe_eqts])
    sr_weighted_eqt_dict = {x: y / sr_weights_eq_weights_sum for x, y in risk_sharpe_ticker_weights.items() if x in risk_sharpe_eqts}
    sr_risk_sharpe_eqt = sr_risk_sharpe.loc[sr_risk_sharpe["Names"].isin(risk_sharpe_eqts), :].set_index("Names")
    sr_risk_sharpe_eqt_num = sr_risk_sharpe_eqt.loc[risk_sharpe_eqts][num_cols].multiply(pd.Series(sr_weighted_eqt_dict), axis=0)
    sr_risk_sharpe_eqt_out = sr_risk_sharpe_eqt.loc[risk_sharpe_eqts][[x for x in sr_risk_sharpe_eqt.columns]]
    sr_risk_sharpe_eqt = sr_risk_sharpe_eqt_num.sum().to_frame().T
    sr_risk_sharpe_eqt = sr_risk_sharpe_eqt.drop(columns=["abs"])
    sr_risk_sharpe_eqt[["KPI", "KPI t-1", "KPI Index"]] = sr_risk_sharpe_eqt[["KPI", "KPI t-1", "KPI Index"]] * 100
    sr_risk_sharpe_eqt.index = ["RISK_EQT"]
    sr_risk_sharpe_eqt.index.name = ""
    sr_risk_sharpe_eqt = sr_risk_sharpe_eqt.rename(columns={"Shrp 3D COB": "Shrp 3", "Shrp 22D": "Shrp 22", "Shrp 3D Live": "Shrp 3 Live"})
    sr_risk_sharpe_eqt = sr_risk_sharpe_eqt[["KPI", "KPI t-1", "Shrp 3 Live", "Shrp 3", "Shrp 22", "KPI Index", "Index"]]
    sr_risk_sharpe_eqt = sr_risk_sharpe_eqt.reset_index()

    sr_weights_ccy_weights_sum = sum([risk_sharpe_ticker_weights[x] for x in risk_sharpe_ccy])
    sr_weighted_ccy_dict = {x: y / sr_weights_ccy_weights_sum for x, y in risk_sharpe_ticker_weights.items() if x in risk_sharpe_ccy}
    sr_risk_sharpe_ccy = sr_risk_sharpe.loc[sr_risk_sharpe["Names"].isin(risk_sharpe_ccy), :].set_index("Names")
    sr_risk_sharpe_ccy_num = sr_risk_sharpe_ccy.loc[risk_sharpe_ccy][num_cols].multiply(pd.Series(sr_weighted_ccy_dict), axis=0)
    sr_risk_sharpe_ccy_out = sr_risk_sharpe_ccy.loc[risk_sharpe_ccy][[x for x in sr_risk_sharpe_ccy.columns]]
    sr_risk_sharpe_ccy = sr_risk_sharpe_ccy_num.sum().to_frame().T
    sr_risk_sharpe_ccy = sr_risk_sharpe_ccy.drop(columns=["abs"])
    sr_risk_sharpe_ccy[["KPI", "KPI t-1", "KPI Index"]] = sr_risk_sharpe_ccy[["KPI", "KPI t-1", "KPI Index"]] * 100
    sr_risk_sharpe_ccy.index = ["RISK_USD"]
    sr_risk_sharpe_ccy.index.name = ""
    sr_risk_sharpe_ccy = sr_risk_sharpe_ccy.rename(columns={"Shrp 3D COB": "Shrp 3", "Shrp 22D": "Shrp 22", "Shrp 3D Live": "Shrp 3 Live"})
    sr_risk_sharpe_ccy = sr_risk_sharpe_ccy[["KPI", "KPI t-1", "Shrp 3 Live", "Shrp 3", "Shrp 22", "KPI Index", "Index"]]
    sr_risk_sharpe_ccy = sr_risk_sharpe_ccy.reset_index()

    sr_weights_cmds_weights_sum = sum([risk_sharpe_ticker_weights[x] for x in risk_sharpe_cmds])
    sr_weighted_cmds_dict = {x: y / sr_weights_cmds_weights_sum for x, y in risk_sharpe_ticker_weights.items() if x in risk_sharpe_cmds}
    sr_risk_sharpe_cmds = sr_risk_sharpe.loc[sr_risk_sharpe["Names"].isin(risk_sharpe_cmds), :].set_index("Names")
    sr_risk_sharpe_cmds_num = sr_risk_sharpe_cmds.loc[risk_sharpe_cmds][num_cols].multiply(pd.Series(sr_weighted_cmds_dict), axis=0)
    sr_risk_sharpe_cmds_out = sr_risk_sharpe_cmds.loc[risk_sharpe_cmds][[x for x in sr_risk_sharpe_cmds.columns]]
    sr_risk_sharpe_cmds = sr_risk_sharpe_cmds_num.sum().to_frame().T
    sr_risk_sharpe_cmds = sr_risk_sharpe_cmds.drop(columns=["abs"])
    sr_risk_sharpe_cmds[["KPI", "KPI t-1", "KPI Index"]] = sr_risk_sharpe_cmds[["KPI", "KPI t-1", "KPI Index"]] * 100
    sr_risk_sharpe_cmds.index = ["RISK_CMDS"]
    sr_risk_sharpe_cmds.index.name = ""
    sr_risk_sharpe_cmds = sr_risk_sharpe_cmds.rename(columns={"Shrp 3D COB": "Shrp 3", "Shrp 22D": "Shrp 22", "Shrp 3D Live": "Shrp 3 Live"})
    sr_risk_sharpe_cmds = sr_risk_sharpe_cmds[["KPI", "KPI t-1", "Shrp 3 Live", "Shrp 3", "Shrp 22", "KPI Index", "Index"]]
    sr_risk_sharpe_cmds = sr_risk_sharpe_cmds.reset_index()

    sr_weights_trsy_weights_sum = sum([risk_sharpe_ticker_weights[x] for x in risk_sharpe_trsy])
    sr_weighted_trsy_dict = {x: y / sr_weights_trsy_weights_sum for x, y in risk_sharpe_ticker_weights.items() if x in risk_sharpe_trsy}
    sr_risk_sharpe_trsy = sr_risk_sharpe.loc[sr_risk_sharpe["Names"].isin(risk_sharpe_trsy), :].set_index("Names")
    sr_risk_sharpe_trsy_num = sr_risk_sharpe_trsy.loc[risk_sharpe_trsy][num_cols].multiply(pd.Series(sr_weighted_trsy_dict), axis=0)
    sr_risk_sharpe_trsy_out = sr_risk_sharpe_trsy.loc[risk_sharpe_trsy][[x for x in sr_risk_sharpe_trsy.columns]]
    sr_risk_sharpe_trsy = sr_risk_sharpe_trsy_num.sum().to_frame().T
    sr_risk_sharpe_trsy = sr_risk_sharpe_trsy.drop(columns=["abs"])
    sr_risk_sharpe_trsy[["KPI", "KPI t-1", "KPI Index"]] = sr_risk_sharpe_trsy[["KPI", "KPI t-1", "KPI Index"]] * 100
    sr_risk_sharpe_trsy.index = ["RISK_RATES"]
    sr_risk_sharpe_trsy.index.name = ""
    sr_risk_sharpe_trsy = sr_risk_sharpe_trsy.rename(columns={"Shrp 3D COB": "Shrp 3", "Shrp 22D": "Shrp 22", "Shrp 3D Live": "Shrp 3 Live"})
    sr_risk_sharpe_trsy = sr_risk_sharpe_trsy[["KPI", "KPI t-1", "Shrp 3 Live", "Shrp 3", "Shrp 22", "KPI Index", "Index"]]
    sr_risk_sharpe_trsy = sr_risk_sharpe_trsy.reset_index()

    weighted_sr_risk = sr_risk_sharpe.set_index("Names").loc[risk_sharpe_ticker_weights.keys()][num_cols].multiply(pd.Series(risk_sharpe_ticker_weights), axis=0).sum().to_frame().T
    weighted_sr_risk = weighted_sr_risk.drop(columns=["abs"])
    weighted_sr_risk[["KPI", "KPI t-1", "KPI Index"]] = weighted_sr_risk[["KPI", "KPI t-1", "KPI Index"]] * 100
    weighted_sr_risk.index = [f"<a href='{c_path}'>RISK</a>"]
    weighted_sr_risk.index.name = ""
    weighted_sr_risk = weighted_sr_risk.rename(columns={"Shrp 3D COB": "Shrp 3", "Shrp 22D": "Shrp 22", "Shrp 3D Live": "Shrp 3 Live"})
    weighted_sr_risk = weighted_sr_risk[["KPI", "KPI t-1", "Shrp 3 Live", "Shrp 3", "Shrp 22", "KPI Index", "Index"]]
    weighted_sr_risk = weighted_sr_risk.reset_index()
    weighted_sr_risk = pd.concat(
        [
            weighted_sr_risk,
            sr_risk_sharpe_eqt,
            sr_risk_sharpe_ccy,
            sr_risk_sharpe_cmds,
            sr_risk_sharpe_trsy,
        ],
        axis=0,
    ).reset_index(drop=True)
    sr_risk_sharpe_eqt_out = sr_risk_sharpe_eqt_out.reset_index().drop(columns=["abs", "_plot_link", "Names"])
    sr_risk_sharpe_eqt_out = sr_risk_sharpe_eqt_out.set_index("Ticker")
    sr_risk_sharpe_eqt_out.index.name = ""
    sr_risk_sharpe_ccy_out = sr_risk_sharpe_ccy_out.reset_index().drop(columns=["abs", "_plot_link", "Names"])
    sr_risk_sharpe_ccy_out = sr_risk_sharpe_ccy_out.set_index("Ticker")
    sr_risk_sharpe_ccy_out.index.name = ""
    sr_risk_sharpe_cmds_out = sr_risk_sharpe_cmds_out.reset_index().drop(columns=["abs", "_plot_link", "Names"])
    sr_risk_sharpe_cmds_out = sr_risk_sharpe_cmds_out.set_index("Ticker")
    sr_risk_sharpe_cmds_out.index.name = ""
    sr_risk_sharpe_trsy_out = sr_risk_sharpe_trsy_out.reset_index().drop(columns=["abs", "_plot_link", "Names"])
    sr_risk_sharpe_trsy_out = sr_risk_sharpe_trsy_out.set_index("Ticker")
    sr_risk_sharpe_trsy_out.index.name = ""
    per_component_kpis = pd.concat([sr_risk_sharpe_eqt_out, sr_risk_sharpe_ccy_out, sr_risk_sharpe_cmds_out, sr_risk_sharpe_trsy_out], axis=0)
    sr_risk_sharpe = sr_risk_sharpe.set_index("Names").loc[risk_sharpe_ticker_weights.keys()].reset_index()
    sr_risk_sharpe.drop(columns=["Names"], inplace=True)
    sr_risk_sharpe["rank1"] = get_rank(sr_risk_sharpe["Shrp 3D COB 1"])
    sr_risk_sharpe["rank2"] = get_rank(sr_risk_sharpe["Shrp 3D COB 2"])
    sr_risk_sharpe["rank3"] = get_rank(sr_risk_sharpe["Shrp 3D COB 3"])
    sr_risk_sharpe["new"] = 0
    for idx, row in sr_risk_sharpe.iterrows():
        if (
            ((row["rank1"] > 10) or (row["rank1"] <= 10 and np.sign(row["Shrp 3D COB"]) != np.sign(row["Shrp 3D COB 1"])))
            and ((row["rank2"] > 10) or (row["rank2"] <= 10 and np.sign(row["Shrp 3D COB"]) != np.sign(row["Shrp 3D COB 2"])))
            and ((row["rank3"] > 10) or (row["rank3"] <= 10 and np.sign(row["Shrp 3D COB"]) != np.sign(row["Shrp 3D COB 3"])))
        ):
            sr_risk_sharpe.loc[idx, "new"] = 1

    sr_risk_sharpe = sr_risk_sharpe.set_index("Ticker")[["Shrp 3D Live", "Shrp 3D COB", "Shrp 22D", "Signal", "KPI", "KPI t-1", "KPI Index", "Index", "new", "Shrp 3D COB 1", "Shrp 3D COB 2", "Shrp 3D COB 3"]]

    if len(sr_risk_sharpe) > 0:
        sr_risk_html_component = format_table_html_component(sr_risk_sharpe, "Risk Index 3d COB Sharpe Ratio")
        sr_risk_html = table.html_format(
            df=weighted_sr_risk,
            hide_cols=["Shrp 3D COB 1", "Shrp 3D COB 2", "Shrp 3D COB 3"],
            format_column={
                "Shrp 3 Live": {"width": "80px", "text-align": "center", "format": "{:,.2f}"},
                "Shrp 3": {"width": "80px", "text-align": "center", "format": "{:,.2f}"},
                "Shrp 22": {"width": "80px", "text-align": "center", "format": "{:,.2f}"},
                "KPI": {"width": "80px", "text-align": "center", "format": "{:,.0f}%"},
                "KPI Index": {"width": "80px", "text-align": "center", "format": "{:,.0f}%"},
                "KPI t-1": {"width": "80px", "text-align": "center", "format": "{:,.0f}%"},
                "Index": {"width": "80px", "text-align": "center", "format": "{:,.2f}"},
            },
        )

        sr_risk_html_component = risk_index_def_link + sr_risk_html_component
        table.to_html([sr_risk_html_component], c_path)
    else:
        sr_risk_html = "No DATA: Risk Index 3d COB Sharpe Ratio<br>"
    return sr_risk_html


def send_sharpe_ratio_rank(send_to=None):

    sr = pd.read_csv(convert_path_to_linux(f"{folder_path}sharpe_ratio.csv"))
    sr.set_index("Unnamed: 0", inplace=True)
    sr.dropna(axis=0, inplace=True)
    sr["abs"] = np.abs(sr["Shrp 3D COB"])
    sr = sr.sort_values("abs", ascending=False)
    sr["Ticker"] = [x for x in sr["_plot_link"]]
    sr.drop(columns=["_plot_link"])

    def get_rank(data):
        data_ = np.abs(data)
        rank_ = data_.rank(ascending=False)
        return rank_

    def table_format(df, header):
        return table.html_format(
            df=df,
            header=header,
            hide_cols=["abs", "new"],
            format_column={
                "Instr": {"width": "40px", "text-align": "center"},
                "Ticker": {"width": "350px", "text-align": "center"},
                "Shrp 3D Live": {"width": "80px", "text-align": "center"},
                "Shrp 3D COB": {"width": "80px", "text-align": "center"},
                "Shrp 22D": {"width": "80px", "text-align": "center"},
                "Signal": {"width": "80px", "text-align": "center"},
                "new": {"highlight": ["Ticker", "new"]},
            },
        )

    sr_risk_html = get_risk_index_html()

    sr_macro_ = sr.loc[sr["Instr"].isin(["EQT", "FI", "CCY", "VOL"]), :]
    sr_macro_.reset_index(drop=True, inplace=True)
    sr_macro_["rank1"] = get_rank(sr_macro_["Shrp 3D COB 1"])
    sr_macro_["rank2"] = get_rank(sr_macro_["Shrp 3D COB 2"])
    sr_macro_["rank3"] = get_rank(sr_macro_["Shrp 3D COB 3"])
    sr_macro = sr_macro_.iloc[:10, :]
    sr_macro["new"] = 0
    for idx, row in sr_macro.iterrows():
        if (
            ((row["rank1"] > 10) or (row["rank1"] <= 10 and np.sign(row["Shrp 3D COB"]) != np.sign(row["Shrp 3D COB 1"])))
            and ((row["rank2"] > 10) or (row["rank2"] <= 10 and np.sign(row["Shrp 3D COB"]) != np.sign(row["Shrp 3D COB 2"])))
            and ((row["rank3"] > 10) or (row["rank3"] <= 10 and np.sign(row["Shrp 3D COB"]) != np.sign(row["Shrp 3D COB 3"])))
        ):
            sr_macro.loc[idx, "new"] = 1
    sr_macro = sr_macro[["Instr", "Ticker", "Shrp 3D Live", "Shrp 3D COB", "Shrp 22D", "Signal", "new", "abs"]]
    figs = []
    figs_macro = []
    figs_oil = []
    figs_metals_ags = []
    figs_gas = []
    figs.append("<br>")
    figs.append("<div style='font-family:Calibri;' >")
    figs.append("<br>")
    figs.append("<b>Risk Index</b>")
    figs.append("<br>")
    figs.append("<a href=https://elementcapital.atlassian.net/wiki/spaces/LO25/pages/2176090113>Risk Index Definition</a><br>")
    figs.append(sr_risk_html)
    figs.append("<br>Green ticker is new signal.")
    figs_oil.append("<br>")
    figs_oil.append("<div style='font-family:Calibri;' >")
    figs_oil.append("<b>Risk Index</b>")
    figs_oil.append("<br>")
    figs_oil.append("<a href=https://elementcapital.atlassian.net/wiki/spaces/LO25/pages/2176090113>Risk Index Definition</a><br>")
    figs_oil.append(sr_risk_html)
    figs_oil.append("Green ticker is new signal.")
    figs_macro.append("<div style='font-family:Calibri;' >")
    figs_macro.append("<b>Risk Index</b>")
    figs_macro.append("<br>")
    figs_macro.append("<a href=https://elementcapital.atlassian.net/wiki/spaces/LO25/pages/2176090113>Risk Index Definition</a><br>")
    figs_macro.append(sr_risk_html)
    figs_macro.append("Green ticker is new signal.")
    figs_metals_ags.append("<br>")
    figs_metals_ags.append("<div style='font-family:Calibri;' >")
    figs_metals_ags.append("<b>Risk Index</b>")
    figs_metals_ags.append("<br>")
    figs_metals_ags.append("<a href=https://elementcapital.atlassian.net/wiki/spaces/LO25/pages/2176090113>Risk Index Definition</a><br>")
    figs_metals_ags.append(sr_risk_html)
    figs_metals_ags.append("Green ticker is new signal.")
    figs_gas.append("<br>")
    figs_gas.append("<div style='font-family:Calibri;' >")
    figs_gas.append("<b>Risk Index</b>")
    figs_gas.append("<br>")
    figs_gas.append("<a href=https://elementcapital.atlassian.net/wiki/spaces/LO25/pages/2176090113>Risk Index Definition</a><br>")
    figs_gas.append(sr_risk_html)
    figs_gas.append("Green ticker is new signal.")
    if len(sr_macro) > 0:
        sr_macro_html = table_format(sr_macro, "Macro 3d COB Sharpe Ratio")
        sr_macro_html = table.figs_to_grid([sr_macro_html, "", ""], columns=3, email=True)
        figs.append(sr_macro_html)
        figs.append("<br>")
    else:
        figs.append("No DATA: Macro 3d COB Sharpe Ratio")
        figs.append("<br>")

    sr_oil_ = sr.loc[sr["Instr"].isin(["BRT", "WTI", "DIST", "GASOLINE"]), :].copy()
    sr_oil_flat_ = sr_oil_.loc[sr_oil_["Type"].isin(["FLAT"]), :]
    sr_oil_sprd_ = sr_oil_.loc[sr_oil_["Type"].isin(["SPRD"]), :]
    sr_oil_flat_.reset_index(drop=True, inplace=True)
    sr_oil_flat_["rank1"] = get_rank(sr_oil_flat_["Shrp 3D COB 1"])
    sr_oil_flat_["rank2"] = get_rank(sr_oil_flat_["Shrp 3D COB 2"])
    sr_oil_flat_["rank3"] = get_rank(sr_oil_flat_["Shrp 3D COB 3"])
    sr_oil_flat = sr_oil_flat_.sort_values(by=["abs"], ascending=False).iloc[:10, :]
    sr_oil_flat["new"] = 0
    for idx, row in sr_oil_flat.iterrows():
        if (
            ((row["rank1"] > 10) or (row["rank1"] <= 10 and np.sign(row["Shrp 3D COB"]) != np.sign(row["Shrp 3D COB 1"])))
            and ((row["rank2"] > 10) or (row["rank2"] <= 10 and np.sign(row["Shrp 3D COB"]) != np.sign(row["Shrp 3D COB 2"])))
            and ((row["rank3"] > 10) or (row["rank3"] <= 10 and np.sign(row["Shrp 3D COB"]) != np.sign(row["Shrp 3D COB 3"])))
        ):
            sr_oil_flat.loc[idx, "new"] = 1
    sr_oil_flat = sr_oil_flat[["Instr", "Ticker", "Shrp 3D Live", "Shrp 3D COB", "Shrp 22D", "Signal", "new", "abs"]]
    if len(sr_oil_flat) > 0:
        sr_oil_flat_html = table_format(sr_oil_flat.drop(columns=["abs"]), "Oil Flat 3d COB Sharpe Ratio")
    else:
        sr_oil_flat_html = "Oil Flat 3d COB Sharpe Ratio"
    sr_oil_sprd_.reset_index(drop=True, inplace=True)
    sr_oil_sprd_["rank1"] = get_rank(sr_oil_sprd_["Shrp 3D COB 1"])
    sr_oil_sprd_["rank2"] = get_rank(sr_oil_sprd_["Shrp 3D COB 2"])
    sr_oil_sprd_["rank3"] = get_rank(sr_oil_sprd_["Shrp 3D COB 3"])
    sr_oil_sprd = sr_oil_sprd_.sort_values(by=["abs"], ascending=False).iloc[:10, :]
    sr_oil_sprd["new"] = 0
    for idx, row in sr_oil_sprd.iterrows():
        if (
            ((row["rank1"] > 10) or (row["rank1"] <= 10 and np.sign(row["Shrp 3D COB"]) != np.sign(row["Shrp 3D COB 1"])))
            and ((row["rank2"] > 10) or (row["rank2"] <= 10 and np.sign(row["Shrp 3D COB"]) != np.sign(row["Shrp 3D COB 2"])))
            and ((row["rank3"] > 10) or (row["rank3"] <= 10 and np.sign(row["Shrp 3D COB"]) != np.sign(row["Shrp 3D COB 3"])))
        ):
            sr_oil_sprd.loc[idx, "new"] = 1
    sr_oil_sprd = sr_oil_sprd[["Instr", "Ticker", "Shrp 3D Live", "Shrp 3D COB", "Shrp 22D", "Signal", "new", "abs"]]
    if len(sr_oil_sprd) > 0:
        sr_oil_spread_html = table_format(sr_oil_sprd.drop(columns=["abs"]), "Oil Spread 3d COB Sharpe Ratio")
    else:
        sr_oil_spread_html = "No Oil Spread 3d COB Sharpe Ratio Alert"
    sr_oil_individual_html = table.figs_to_grid([sr_oil_flat_html, sr_oil_spread_html, ""], columns=3, email=True)
    figs_oil.append(sr_oil_individual_html)
    figs_oil.append("<br>")

    sr_oil = pd.concat([sr_oil_flat, sr_oil_sprd], axis=0, ignore_index=True)
    sr_oil = sr_oil.loc[sr_oil["abs"] > 1, :].iloc[:10, :]
    sr_oil_sprd = sr_oil_sprd.loc[sr_oil_sprd["abs"] > 1, :].iloc[:10, :]
    sr_oil_flt = sr_oil_flat.loc[sr_oil_flat["abs"] > 1, :].iloc[:10, :]
    if len(sr_oil) > 0:
        if sr_oil_flt.empty:
            sr_oil_html_flt = "No Oil Flat 3d COB Sharpe has an abs value > 1"
        else:
            sr_oil_html_flt = table_format(sr_oil_flt, "Oil Flat 3d COB Sharpe Ratio")
        if sr_oil_sprd.empty:
            sr_oil_html_sprd = "No Oil Spread 3d COB Sharpe has an abs value > 1"
        else:
            sr_oil_html_sprd = table_format(sr_oil_sprd, "Oil Spread 3d COB Sharpe Ratio")
        sr_oil_html = table.figs_to_grid([sr_oil_html_flt, sr_oil_html_sprd, ""], columns=3, email=True)
        figs.append(sr_oil_html)
        figs.append("<br>")
    else:
        figs.append("No Oil 3d COB Sharpe Ratio Alert")
        figs.append("<br>")

    sr_gas_ = sr.loc[sr["Instr"].isin(["NG", "TTF", "HH"]), :].copy()

    sr_gas_flat_ = sr_gas_.loc[sr_gas_["Type"].isin(["FLAT"]), :]
    sr_gas_flat_.reset_index(drop=True, inplace=True)
    sr_gas_flat_["rank1"] = get_rank(sr_gas_flat_["Shrp 3D COB 1"])
    sr_gas_flat_["rank2"] = get_rank(sr_gas_flat_["Shrp 3D COB 2"])
    sr_gas_flat_["rank3"] = get_rank(sr_gas_flat_["Shrp 3D COB 3"])
    sr_gas_flat = sr_gas_flat_.sort_values(by=["abs"], ascending=False).iloc[:10, :]
    sr_gas_flat["new"] = 0
    for idx, row in sr_gas_flat.iterrows():
        if (
            ((row["rank1"] > 10) or (row["rank1"] <= 10 and np.sign(row["Shrp 3D COB"]) != np.sign(row["Shrp 3D COB 1"])))
            and ((row["rank2"] > 10) or (row["rank2"] <= 10 and np.sign(row["Shrp 3D COB"]) != np.sign(row["Shrp 3D COB 2"])))
            and ((row["rank3"] > 10) or (row["rank3"] <= 10 and np.sign(row["Shrp 3D COB"]) != np.sign(row["Shrp 3D COB 3"])))
        ):
            sr_gas_flat.loc[idx, "new"] = 1
    sr_gas_flat = sr_gas_flat[["Instr", "Ticker", "Shrp 3D Live", "Shrp 3D COB", "Shrp 22D", "Signal", "new", "abs"]]
    if len(sr_gas_flat) > 0:
        sr_gas_flat_html = table_format(sr_gas_flat.drop(columns=["abs"]), "Gas Flat 3d COB Sharpe Ratio Alert")
    else:
        sr_gas_flat_html = "No Gas Flat 3d COB Sharpe Ratio Alert"

    sr_gas_sprd_ = sr_gas_.loc[sr_gas_["Type"].isin(["SPRD"]), :]
    sr_gas_sprd_.reset_index(drop=True, inplace=True)
    sr_gas_sprd_["rank1"] = get_rank(sr_gas_sprd_["Shrp 3D COB 1"])
    sr_gas_sprd_["rank2"] = get_rank(sr_gas_sprd_["Shrp 3D COB 2"])
    sr_gas_sprd_["rank3"] = get_rank(sr_gas_sprd_["Shrp 3D COB 3"])
    sr_gas_sprd = sr_gas_sprd_.sort_values(by=["abs"], ascending=False).iloc[:10, :]
    sr_gas_sprd["new"] = 0
    for idx, row in sr_gas_sprd.iterrows():
        if (
            ((row["rank1"] > 10) or (row["rank1"] <= 10 and np.sign(row["Shrp 3D COB"]) != np.sign(row["Shrp 3D COB 1"])))
            and ((row["rank2"] > 10) or (row["rank2"] <= 10 and np.sign(row["Shrp 3D COB"]) != np.sign(row["Shrp 3D COB 2"])))
            and ((row["rank3"] > 10) or (row["rank3"] <= 10 and np.sign(row["Shrp 3D COB"]) != np.sign(row["Shrp 3D COB 3"])))
        ):
            sr_gas_sprd.loc[idx, "new"] = 1
    sr_gas_sprd = sr_gas_sprd[["Instr", "Ticker", "Shrp 3D Live", "Shrp 3D COB", "Shrp 22D", "Signal", "new", "abs"]]
    if len(sr_gas_sprd) > 0:
        sr_gas_spread_html = table_format(sr_gas_sprd.drop(columns=["abs"]), "Gas Spread 3d COB Sharpe Ratio Alert")
    else:
        sr_gas_spread_html = "No Gas Spread 3d COB Sharpe Ratio Alert"
    sr_gas_individual_html = table.figs_to_grid([sr_gas_flat_html, sr_gas_spread_html, ""], columns=3, email=True)
    figs_gas.append(sr_gas_individual_html)
    figs_gas.append("<br>")

    sr_gas = pd.concat([sr_gas_flat, sr_gas_sprd], axis=0, ignore_index=True)
    sr_gas = sr_gas.loc[sr_gas["abs"] > 1, :].iloc[:10, :]
    sr_gas_sprd = sr_gas_sprd.loc[sr_gas_sprd["abs"] > 1, :].iloc[:10, :]
    sr_gas_flt = sr_gas_flat.loc[sr_gas_flat["abs"] > 1, :].iloc[:10, :]
    if len(sr_gas) > 0:
        if sr_gas_flt.empty:
            sr_gas_html_flt = "No Gas Flat 3d COB Sharpe has an abs value >1"
        else:
            sr_gas_html_flt = table_format(sr_gas_flt, "Gas Flat 3d COB Sharpe Ratio")
        if sr_gas_sprd.empty:
            sr_gas_html_sprd = "No Gas Spread 3d COB Sharpe has an abs value >1"
        else:
            sr_gas_html_sprd = table_format(sr_gas_sprd, "Gas Spread 3d COB Sharpe Ratio")
        sr_gas_html = table.figs_to_grid([sr_gas_html_flt, sr_gas_html_sprd, ""], columns=3, email=True)
        figs.append(sr_gas_html)
        figs.append("<br>")
    else:
        figs.append("No Gas 3d COB Sharpe Ratio Alert")
        figs.append("<br>")

    sr_metal_ = sr.loc[sr["Instr"].isin(["PM", "BM"]), :].copy()

    sr_metal_flat_ = sr_metal_.loc[sr_metal_["Type"].isin(["FLAT"]), :]
    sr_metal_flat_.reset_index(drop=True, inplace=True)
    sr_metal_flat_["rank1"] = get_rank(sr_metal_flat_["Shrp 3D COB 1"])
    sr_metal_flat_["rank2"] = get_rank(sr_metal_flat_["Shrp 3D COB 2"])
    sr_metal_flat_["rank3"] = get_rank(sr_metal_flat_["Shrp 3D COB 3"])
    sr_metal_flat = sr_metal_flat_.sort_values(by=["abs"], ascending=False).iloc[:10, :]
    sr_metal_flat["new"] = 0
    for idx, row in sr_metal_flat.iterrows():
        if (
            ((row["rank1"] > 10) or (row["rank1"] <= 10 and np.sign(row["Shrp 3D COB"]) != np.sign(row["Shrp 3D COB 1"])))
            and ((row["rank2"] > 10) or (row["rank2"] <= 10 and np.sign(row["Shrp 3D COB"]) != np.sign(row["Shrp 3D COB 2"])))
            and ((row["rank3"] > 10) or (row["rank3"] <= 10 and np.sign(row["Shrp 3D COB"]) != np.sign(row["Shrp 3D COB 3"])))
        ):
            sr_metal_flat.loc[idx, "new"] = 1
    sr_metal_flat = sr_metal_flat[["Instr", "Ticker", "Shrp 3D Live", "Shrp 3D COB", "Shrp 22D", "Signal", "new", "abs"]]
    if len(sr_metal_flat) > 0:
        sr_metal_flat_html = table_format(sr_metal_flat.drop(columns=["abs"]), "Metal Flat 3d COB Sharpe Ratio Alert")
        figs_metals_ags.append(sr_metal_flat_html)
        figs_metals_ags.append("<br>")
    else:
        sr_metal_flat_html = "No Metal Flat 3d COB Sharpe Ratio Alert"
        figs_metals_ags.append(sr_metal_flat_html)
        figs_metals_ags.append("<br>")

    sr_metal_sprd_ = sr_metal_.loc[sr_metal_["Type"].isin(["SPRD"]), :]
    sr_metal_sprd_.reset_index(drop=True, inplace=True)
    sr_metal_sprd_["rank1"] = get_rank(sr_metal_sprd_["Shrp 3D COB 1"])
    sr_metal_sprd_["rank2"] = get_rank(sr_metal_sprd_["Shrp 3D COB 2"])
    sr_metal_sprd_["rank3"] = get_rank(sr_metal_sprd_["Shrp 3D COB 3"])
    sr_metal_sprd = sr_metal_sprd_.sort_values(by=["abs"], ascending=False).iloc[:10, :]
    sr_metal_sprd["new"] = 0
    for idx, row in sr_metal_sprd.iterrows():
        if (
            ((row["rank1"] > 10) or (row["rank1"] <= 10 and np.sign(row["Shrp 3D COB"]) != np.sign(row["Shrp 3D COB 1"])))
            and ((row["rank2"] > 10) or (row["rank2"] <= 10 and np.sign(row["Shrp 3D COB"]) != np.sign(row["Shrp 3D COB 2"])))
            and ((row["rank3"] > 10) or (row["rank3"] <= 10 and np.sign(row["Shrp 3D COB"]) != np.sign(row["Shrp 3D COB 3"])))
        ):
            sr_metal_sprd.loc[idx, "new"] = 1
    sr_metal_sprd = sr_metal_sprd[["Instr", "Ticker", "Shrp 3D Live", "Shrp 3D COB", "Shrp 22D", "Signal", "new", "abs"]]
    if len(sr_metal_sprd) > 0:
        sr_metal_sprd_html = table_format(sr_metal_sprd.drop(columns=["abs"]), "Metal Spread 3d COB Sharpe Ratio Alert")
        figs_metals_ags.append(sr_metal_sprd_html)
        figs_metals_ags.append("<br>")
    else:
        sr_metal_sprd_html = "No Metal Spread 3d COB Sharpe Ratio Alert"
        figs_metals_ags.append(sr_metal_sprd_html)
        figs_metals_ags.append("<br>")

    sr_metal = pd.concat([sr_metal_flat, sr_metal_sprd], axis=0, ignore_index=True)
    sr_metal = sr_metal.loc[sr_metal["abs"] > 1, :].iloc[:10, :]
    if len(sr_metal) > 0:
        sr_metal_html = table_format(sr_metal, "Metals 3d COB Sharpe Ratio Alert")
        sr_metal_html = table.figs_to_grid([sr_metal_flat_html, "", ""], columns=3, email=True)
        figs.append(sr_metal_html)
        figs.append("<br>")
    else:
        figs.append("No Metals 3d COB Sharpe Ratio Alert")
        figs.append("<br>")

    sr_ags_ = sr.loc[sr["Instr"].isin(["Bulks", "GR", "SFT"]), :]

    sr_ags_flat_ = sr_ags_.loc[sr_ags_["Type"].isin(["FLAT"]), :]
    sr_ags_flat_.reset_index(drop=True, inplace=True)
    sr_ags_flat_["rank1"] = get_rank(sr_ags_flat_["Shrp 3D COB 1"])
    sr_ags_flat_["rank2"] = get_rank(sr_ags_flat_["Shrp 3D COB 2"])
    sr_ags_flat_["rank3"] = get_rank(sr_ags_flat_["Shrp 3D COB 3"])
    sr_ags_flat = sr_ags_flat_.sort_values(by=["abs"], ascending=False).iloc[:10, :]
    sr_ags_flat["new"] = 0
    for idx, row in sr_ags_flat.iterrows():
        if (
            ((row["rank1"] > 10) or (row["rank1"] <= 10 and np.sign(row["Shrp 3D COB"]) != np.sign(row["Shrp 3D COB 1"])))
            and ((row["rank2"] > 10) or (row["rank2"] <= 10 and np.sign(row["Shrp 3D COB"]) != np.sign(row["Shrp 3D COB 2"])))
            and ((row["rank3"] > 10) or (row["rank3"] <= 10 and np.sign(row["Shrp 3D COB"]) != np.sign(row["Shrp 3D COB 3"])))
        ):
            sr_ags_flat.loc[idx, "new"] = 1
    sr_ags_flat = sr_ags_flat[["Instr", "Ticker", "Shrp 3D Live", "Shrp 3D COB", "Shrp 22D", "Signal", "new", "abs"]]
    if len(sr_ags_flat) > 0:
        sr_ags_flat_html = table_format(sr_ags_flat.drop(columns=["abs"]), "Ags Flat 3d COB Sharpe Ratio Alert")
        figs_metals_ags.append(sr_ags_flat_html)
        figs_metals_ags.append("<br>")
    else:
        sr_metal_ags_flat_html = "No Ags Flat 3d COB Sharpe Ratio Alert"
        figs_metals_ags.append(sr_metal_ags_flat_html)
        figs_metals_ags.append("<br>")

    sr_ags_sprd_ = sr_ags_.loc[sr_ags_["Type"].isin(["SPRD"]), :]
    sr_ags_sprd_.reset_index(drop=True, inplace=True)
    sr_ags_sprd_["rank1"] = get_rank(sr_ags_sprd_["Shrp 3D COB 1"])
    sr_ags_sprd_["rank2"] = get_rank(sr_ags_sprd_["Shrp 3D COB 2"])
    sr_ags_sprd_["rank3"] = get_rank(sr_ags_sprd_["Shrp 3D COB 3"])
    sr_ags_sprd = sr_ags_sprd_.sort_values(by=["abs"], ascending=False).iloc[:10, :]
    sr_ags_sprd["new"] = 0
    for idx, row in sr_ags_sprd.iterrows():
        if (
            ((row["rank1"] > 10) or (row["rank1"] <= 10 and np.sign(row["Shrp 3D COB"]) != np.sign(row["Shrp 3D COB 1"])))
            and ((row["rank2"] > 10) or (row["rank2"] <= 10 and np.sign(row["Shrp 3D COB"]) != np.sign(row["Shrp 3D COB 2"])))
            and ((row["rank3"] > 10) or (row["rank3"] <= 10 and np.sign(row["Shrp 3D COB"]) != np.sign(row["Shrp 3D COB 3"])))
        ):
            sr_ags_sprd.loc[idx, "new"] = 1
    sr_ags_sprd = sr_ags_sprd[["Instr", "Ticker", "Shrp 3D Live", "Shrp 3D COB", "Shrp 22D", "Signal", "new", "abs"]]
    if len(sr_ags_sprd) > 0:
        sr_ags_sprd_html = table_format(sr_ags_sprd.drop(columns=["abs"]), "Ags Spread 3d COB Sharpe Ratio Alert")
        figs_metals_ags.append(sr_ags_sprd_html)
        figs_metals_ags.append("<br>")
    else:
        sr_ags_sprd_html = "No Ags Spread 3d COB Sharpe Ratio Alert"
        figs_metals_ags.append(sr_ags_sprd_html)
        figs_metals_ags.append("<br>")

    sr_ags = pd.concat([sr_ags_flat, sr_ags_sprd], axis=0, ignore_index=True)
    sr_ags = sr_ags.loc[sr_ags["abs"] > 1, :].iloc[:10, :]
    if len(sr_ags) > 0:
        sr_metal_ags_html = table_format(sr_ags.drop(columns=["abs"]), "Ags 3d COB Sharpe Ratio Alert")
        sr_metal_ags_html = table.figs_to_grid([sr_ags_flat_html, "", ""], columns=3, email=True)
        figs.append(sr_metal_ags_html)
        figs.append("<br>")
    else:
        figs.append("No Ags 3d COB Sharpe Ratio Alert")
        figs.append("<br>")

    table.figures_to_html(figs_oil, filename=convert_path_to_linux(f"{html_path}\\cross_cmds\\sharpe_ratio_oil.html"), task_name="market_scan")
    table.figures_to_html(figs_gas, filename=convert_path_to_linux(f"{html_path}\\cross_cmds\\sharpe_ratio_gas.html"), task_name="market_scan")
    table.figures_to_html(figs_metals_ags, filename=convert_path_to_linux(f"{html_path}\\cross_cmds\\sharpe_ratio_cross_cmds.html"), task_name="market_scan")
    table.figures_to_html(figs_macro, filename=convert_path_to_linux(f"{html_path}\\cross_cmds\\sharpe_ratio_macro.html"), task_name="market_scan")
    table.figures_to_html(figs, filename=convert_path_to_linux(f"{html_path}\\cross_cmds\\sharpe_ratio.html"), task_name="market_scan")
    if send_to:
        send_email(
            send_to=send_to,
            subject="Sharpe Ratio Rank",
            body=figs,
            html_path=f"{html_path}\\cross_cmds\\sharpe_ratio.html",
        )
        send_email(
            send_to=config.gas_group_no_mk,
            subject="Sharpe Ratio Rank: GAS",
            body=figs_gas,
            html_path=f"{html_path}\\cross_cmds\\sharpe_ratio_gas.html",
        )
        send_email(
            send_to=config.oil_group_no_mk,
            subject="Sharpe Ratio Rank: OIL",
            body=figs_oil,
            html_path=f"{html_path}\\cross_cmds\\sharpe_ratio_oil.html",
        )


def get_extras_by_asset_class(asset_class=None):
    extras = []
    if asset_class.upper() in ["OIL", "DIST", "GASOLINE", "WTI", "BRT"]:
        diverge_html_path = convert_path_to_linux(f"{folder_path}\\divergence_only.html")
        if os.path.exists(diverge_html_path):
            with open(diverge_html_path, mode="r") as file:
                html_diverge = file.read()
        else:
            df_diverge = pd.read_csv(ut.convert_path_to_linux(f"{folder_path}\\diverge_signal.csv"))
            df_diverge.set_index("Unnamed: 0", inplace=True)
            if len(df_diverge) > 0:
                html_diverge = send_diverge_emails(df_diverge)
            else:
                html_diverge = "No Diverge signal today"
            table.to_html([html_diverge], diverge_html_path, add_home=False, add_generation_timestamp=False)
        fly_swith_html_path = convert_path_to_linux(f"{folder_path}\\fly_switch_only.html")
        if os.path.exists(fly_swith_html_path):
            with open(fly_swith_html_path, mode="r") as file:
                fly_swith_html = file.read()
        else:
            fly_switch_signal = pd.read_csv(ut.convert_path_to_linux(f"{folder_path}\\fly_switch_signal.csv"))
            fly_switch_signal.set_index("Unnamed: 0", inplace=True)
            if len(fly_switch_signal) > 0:
                fly_switch_signal.index.name = "Name"
                fly_switch_signal.reset_index(inplace=True)
                fly_swith_html = table.html_format(
                    fly_switch_signal,
                    precision=2,
                    format_column={
                        "Name": {"width": "160px", "text-align": "center", "right_border": True},
                        tuple(fly_switch_signal.columns[1:]): {"width": "100px", "text-align": "center"},
                    },
                )
            else:
                fly_swith_html = "No fly switch alert today"
            table.to_html([fly_swith_html], fly_swith_html_path, add_home=False, add_generation_timestamp=False)

        alert_spot_html_path = convert_path_to_linux(f"{folder_path}\\spot_only.html")
        if os.path.exists(alert_spot_html_path):
            with open(alert_spot_html_path, mode="r") as file:
                html_spot = file.read()
        else:
            alert_spot = spot_price_alert()
            if len(alert_spot) > 0:
                html_spot = table.html_format(
                    alert_spot,
                    precision=2,
                    format_column={
                        "Spot": {"width": "160px", "text-align": "center", "right_border": True},
                        "Market": {"width": "160px", "text-align": "center", "right_border": True},
                        "Price change": {"width": "100px", "text-align": "center"},
                        "Linked sprd chg": {"width": "100px", "text-align": "center"},
                        "Price": {"width": "100px", "text-align": "center"},
                        "Price 20d MA": {"width": "100px", "text-align": "center"},
                        "20d z-score": {"width": "100px", "text-align": "center"},
                        "Latest Date": {"width": "100px", "text-align": "center"},
                    },
                )
                spot_alert_outpath = convert_path_to_linux(
                    f"{folder_path}\\physical_oil\\spot_alert_{today().strftime('%Y%m%d')}.txt",
                )
                with open(spot_alert_outpath, mode="w") as file:
                    file.write(html_spot)
            else:
                html_spot = "No physical spot price alert today"
            table.to_html([html_spot], alert_spot_html_path, add_home=False, add_generation_timestamp=False)

        alert_phys_html_path = convert_path_to_linux(f"{folder_path}\\physical_only.html")
        if os.path.exists(alert_phys_html_path):
            with open(alert_phys_html_path, mode="r") as file:
                html_phys = file.read()
        else:
            alert_phys, _ = physical_price_alert()
            if len(alert_phys) > 0:
                html_phys = table.html_format(
                    alert_phys,
                    precision=2,
                    format_column={
                        "Crude": {"width": "160px", "text-align": "center", "right_border": True},
                        "Market": {"width": "160px", "text-align": "center", "right_border": True},
                        "Price change": {"width": "100px", "text-align": "center"},
                        "Linked sprd chg": {"width": "100px", "text-align": "center"},
                        "Price": {"width": "100px", "text-align": "center"},
                        "Price 20d MA": {"width": "100px", "text-align": "center"},
                        "20d z-score": {"width": "100px", "text-align": "center"},
                        "Latest Date": {"width": "100px", "text-align": "center"},
                    },
                )
            else:
                html_phys = "No physical crude alert today"
            table.to_html([html_phys], alert_phys_html_path, add_home=False, add_generation_timestamp=False)

        product_alert_phys_html_path = convert_path_to_linux(f"{folder_path}\\physical_product_only.html")
        if os.path.exists(product_alert_phys_html_path):
            with open(product_alert_phys_html_path, mode="r") as file:
                html_phys_product = file.read()
        else:
            _, alert_phys_product = physical_price_alert()
            if len(alert_phys) > 0:
                html_phys_product = table.html_format(
                    alert_phys_product,
                    precision=2,
                    format_column={
                        "Product": {"width": "160px", "text-align": "center", "right_border": True},
                        "Market": {"width": "160px", "text-align": "center", "right_border": True},
                        "Price change": {"width": "100px", "text-align": "center"},
                        "Linked sprd chg": {"width": "100px", "text-align": "center"},
                        "Price": {"width": "100px", "text-align": "center"},
                        "Price 20d MA": {"width": "100px", "text-align": "center"},
                        "20d z-score": {"width": "100px", "text-align": "center"},
                        "Latest Date": {"width": "100px", "text-align": "center"},
                    },
                )
            else:
                html_phys_product = "No physical product alert today"
            table.to_html([html_phys_product], product_alert_phys_html_path, add_home=False, add_generation_timestamp=False)

        alert_crack_html_path = convert_path_to_linux(f"{folder_path}\\crack_only.html")
        if os.path.exists(alert_crack_html_path):
            with open(alert_crack_html_path, mode="r") as file:
                html_crack = file.read()
        else:
            alert_crack = crack_price_alert()
            if len(alert_crack) > 0:
                html_crack = table.html_format(
                    alert_crack,
                    precision=2,
                    format_column={
                        "Crack": {"width": "160px", "text-align": "center", "right_border": True},
                        "Market": {"width": "160px", "text-align": "center", "right_border": True},
                        "Price change": {"width": "100px", "text-align": "center"},
                        "Linked sprd chg": {"width": "100px", "text-align": "center"},
                        "Price": {"width": "100px", "text-align": "center"},
                        "Price 20d MA": {"width": "100px", "text-align": "center"},
                        "20d z-score": {"width": "100px", "text-align": "center"},
                        "Latest Date": {"width": "100px", "text-align": "center"},
                    },
                )
            else:
                html_crack = "No crack alert today"
            table.to_html([html_crack], alert_crack_html_path, add_home=False, add_generation_timestamp=False)

        swap_alert_html_path = convert_path_to_linux(f"{folder_path}\\swap_only.html")
        if os.path.exists(swap_alert_html_path):
            with open(swap_alert_html_path, mode="r") as file:
                html_swap = file.read()
        else:
            alert_swap, _ = swap_price_alert()
            if len(alert_swap) > 0:
                html_swap = table.html_format(
                    alert_swap,
                    precision=2,
                    format_column={
                        "Crude": {"width": "160px", "text-align": "center", "right_border": True},
                        "Market": {"width": "160px", "text-align": "center", "right_border": True},
                        "Price change": {"width": "100px", "text-align": "center"},
                        "Linked sprd chg": {"width": "100px", "text-align": "center"},
                        "Price": {"width": "100px", "text-align": "center"},
                        "Price 20d MA": {"width": "100px", "text-align": "center"},
                        "20d z-score": {"width": "100px", "text-align": "center"},
                        "Latest Date": {"width": "100px", "text-align": "center"},
                    },
                )
            else:
                html_swap = "No crude swap alert today"
            table.to_html([html_swap], swap_alert_html_path, add_home=False, add_generation_timestamp=False)

        product_swap_alert_html_path = convert_path_to_linux(f"{folder_path}\\product_swap_only.html")
        if os.path.exists(product_swap_alert_html_path):
            with open(product_swap_alert_html_path, mode="r") as file:
                html_product_swap = file.read()
        else:
            _, alert_product_swap = swap_price_alert()
            if len(alert_product_swap) > 0:
                html_product_swap = table.html_format(
                    alert_product_swap,
                    precision=2,
                    format_column={
                        "Product": {"width": "160px", "text-align": "center", "right_border": True},
                        "Market": {"width": "160px", "text-align": "center", "right_border": True},
                        "Price change": {"width": "100px", "text-align": "center"},
                        "Linked sprd chg": {"width": "100px", "text-align": "center"},
                        "Price": {"width": "100px", "text-align": "center"},
                        "Price 20d MA": {"width": "100px", "text-align": "center"},
                        "20d z-score": {"width": "100px", "text-align": "center"},
                        "Latest Date": {"width": "100px", "text-align": "center"},
                    },
                )
            else:
                html_product_swap = "No product swap alert today"
            table.to_html([html_product_swap], product_swap_alert_html_path, add_home=False, add_generation_timestamp=False)

        extras = extras + [
            "<br><b>Spread and flat price divergence alert:<br>",
            html_diverge,
            "<br><b>Fly switch alert:<br>",
            fly_swith_html,
            "<br><b>Physical price alert:<br>",
            html_phys,
            html_phys_product,
            "<br><b>Crack price alert:<br>",
            html_crack,
            "<br><b>Swap price alert:<br>",
            html_swap,
            html_product_swap,
        ]
    elif asset_class.upper() in ["TTF", "HH", "NG", "GAS"]:

        eu_gas_html_path = convert_path_to_linux(f"{folder_path}\\ttf_only.html")
        if os.path.exists(eu_gas_html_path):
            with open(eu_gas_html_path, mode="r") as file:
                html_eugas = file.read()
        else:
            eu_gas = pd.read_csv(ut.convert_path_to_linux(f"{output_path}\\csvs\\ttf\\ttf_alert.csv"))
            eu_gas.set_index("Unnamed: 0", inplace=True)
            if len(eu_gas) > 0:
                html_eugas = table.html_format(
                    eu_gas,
                    precision=2,
                    format_column={
                        "Hub": {"width": "160px", "text-align": "center"},
                        "Tenor": {"width": "100px", "text-align": "center"},
                        "Price change": {"width": "100px", "text-align": "center"},
                        "Price": {"width": "100px", "text-align": "center"},
                        "Price 20d MA": {"width": "100px", "text-align": "center"},
                        "3m z-score": {"width": "100px", "text-align": "center"},
                    },
                )
            else:
                html_eugas = "No EU gas spread alert today"
            table.to_html([html_eugas], eu_gas_html_path, add_home=False, add_generation_timestamp=False)

        us_gas_html_path = convert_path_to_linux(f"{folder_path}\\usgas_only.html")
        if os.path.exists(us_gas_html_path):
            with open(us_gas_html_path, mode="r") as file:
                html_usgas = file.read()
        else:
            us_gas = pd.read_csv(ut.convert_path_to_linux(f"{output_path}\\csvs\\gas\\basis\\usgas_alert.csv"))
            us_gas.set_index("Unnamed: 0", inplace=True)
            if len(us_gas) > 0:
                html_usgas = table.html_format(
                    us_gas,
                    precision=2,
                    format_column={
                        "Hub": {"width": "160px", "text-align": "center"},
                        "Tenor": {"width": "100px", "text-align": "center"},
                        "Price change": {"width": "100px", "text-align": "center"},
                        "Price": {"width": "100px", "text-align": "center"},
                        "Price 20d MA": {"width": "100px", "text-align": "center"},
                        "3m z-score": {"width": "100px", "text-align": "center"},
                        "Latest Date": {"width": "100px", "text-align": "center"},
                    },
                )
            else:
                html_usgas = "No US gas spread alert today"
            table.to_html([html_usgas], us_gas_html_path, add_home=False, add_generation_timestamp=False)

        diverge_html_path = convert_path_to_linux(f"{folder_path}\\diverge_only.html")
        if os.path.exists(diverge_html_path):
            with open(diverge_html_path, mode="r") as file:
                html_diverge = file.read()
        else:
            df_diverge = pd.read_csv(ut.convert_path_to_linux(f"{folder_path}\\diverge_signal.csv"))
            df_diverge.set_index("Unnamed: 0", inplace=True)
            if len(df_diverge) > 0:
                html_diverge = send_diverge_emails(df_diverge)
            else:
                html_diverge = "No Diverge signal today"
            table.to_html([html_diverge], diverge_html_path, add_home=False, add_generation_timestamp=False)
        extras = extras + [
            "<br><b>Spread and flat price divergence alert:<br>",
            html_diverge,
            "<br><b>EU Gas Spread Alert:<br>",
            html_eugas,
            "<br><b>US Gas Spread Alert:<br>",
            html_usgas,
        ]
    elif asset_class.upper() in ["METAL"]:

        bm_html_path = convert_path_to_linux(f"{folder_path}\\bm_only.html")
        if os.path.exists(bm_html_path):
            with open(bm_html_path, mode="r") as file:
                html_bm = file.read()
        else:
            df_bm = pd.read_csv(ut.convert_path_to_linux(f"{folder_path}bm_alert_am.csv"))
            if len(df_bm) > 0:
                df_bm.set_index("Unnamed: 0", inplace=True)
                df_bm.reset_index(drop=True, inplace=True)
                df_bm["Ticker"] = [x.rpartition(" ")[0] for x in df_bm["Ticker"]]
                pct_format_rows = df_bm.loc[df_bm["chg type"] == "FLAT", :].index
                dec_format_rows = df_bm.loc[df_bm["chg type"] == "SPRD", :].index
                html_bm = table.html_format(
                    df_bm,
                    precision=2,
                    hide_cols=["chg type"],
                    format_column={
                        "Ticker": {"width": "160px", "text-align": "center"},
                        "Price change": {"width": "100px", "text-align": "center"},
                        "Price": {"width": "100px", "text-align": "center"},
                        "Price 20d MA": {"width": "100px", "text-align": "center"},
                        "20d z-score": {"width": "100px", "text-align": "center"},
                        "Date": {"width": "100px", "text-align": "center"},
                    },
                    format_row={
                        tuple(pct_format_rows): {"format": "{:.2%}", "columns": ["Price change"]},
                        tuple(dec_format_rows): {"format": "{:.2f}", "columns": ["Price change"]},
                    },
                )
            else:
                html_bm = "No base metal alert"
            table.to_html([html_bm], bm_html_path, add_home=False, add_generation_timestamp=False)
        extras = extras + [
            "<br>Base Metal Alert:<br>",
            html_bm,
        ]
    return extras


def generate_range_vol_page_by_asset_class(asset_class=None):
    asset_class_dict = {
        "Macro": ["EQT", "FI", "CCY", "VOL"],
        "Oil": ["WTI", "BRT", "DIST", "GASOLINE", "NSea"],
        "GAS": ["HH", "TTF"],
        "Metal": ["PM", "BM", "Bulks"],
        "AGS": ["GR", "SFT"],
    }
    if asset_class in asset_class_dict.keys():
        aset_class_list = asset_class_dict[asset_class]
    else:
        raise ValueError("Please input correct asset class: Macro, Oil, GAS, Metal, AGS")
    if asset_class == "GAS":
        alert_csv_path_ = ut.convert_path_to_linux(f"{folder_path}\\gas\\trend_alert_{asset_class.lower()}.csv")
        alert_csv_path = ut.convert_path_to_linux(f"{folder_path}\\trend_alert_ng.csv")
        alert_df = pd.read_csv(alert_csv_path_, index_col=0)
    elif asset_class == "Metal":
        alert_csv_path = ut.convert_path_to_linux(f"{folder_path}\\metal\\trend_alert_all.csv")
        alert_df = pd.DataFrame()
    else:
        alert_csv_path = ut.convert_path_to_linux(f"{folder_path}trend_alert_{asset_class.lower()}.csv")
        alert_df = pd.DataFrame()
    if os.path.exists(alert_csv_path):
        alert_df_ = pd.read_csv(alert_csv_path, index_col=0)
        alert_df_ = alert_df_.loc[alert_df_["Instr"].isin(aset_class_list), :]
        if not alert_df.empty:
            alert_df = pd.concat([alert_df, alert_df_], axis=0, ignore_index=True)
        else:
            alert_df = alert_df_
        if len(alert_df) > 0:
            alert_html = get_alert_email_html(alert_df)
        else:
            alert_html = f"No {asset_class} Trend Alert"
    else:
        alert_html = f"Trend Alert file {alert_csv_path} does not exist"
    risk_index_html = get_risk_index_html()
    sr_file = convert_path_to_linux(f"{html_path}\\cross_cmds\\sharpe_ratio_{asset_class.lower()}.html")
    if os.path.exists(sr_file):
        with open(sr_file, "r") as f:
            sr_html = f.read()
        split_idx = sr_html.find("Green")
        if split_idx != -1:
            sr_html = sr_html[split_idx:]
        else:
            split_idx = sr_html.find("</table>")
            sr_html = sr_html[split_idx + len("</table>"):]
    else:
        sr_html = f"File Missing for Sharpe Ratio - {asset_class}"
    extras = get_extras_by_asset_class(asset_class=asset_class)
    all_html = [
        "<div style='font-family:Calibri;' >",
        f"<h2>Range Vol {asset_class.capitalize()} Alert</h2><br>",
        f"Update time {now_ldn().strftime('%Y-%m-%d %H:%M')} <br>",
        "<b>Risk Index</b><br>",
        risk_index_html,
        f"<br><b>{asset_class} Sharpe Ratio Rank</b><br>",
        sr_html,
        f"<br><b>{asset_class} Range-Vol alert: </b><br>",
        '<a href="https://.elementcapital.corp/display/LO25/Range-Vol+and+Divergence">Description</a><br>',
        alert_html,
    ] + extras + ["</div>"]
    out_html_path = convert_path_to_linux(f"{html_path}\\cross_cmds\\range_vol_{asset_class.lower()}.html")
    table.to_html(all_html, out_html_path, add_home=True)
    log.info(f"Range vol page for {asset_class} generated")


def update():
    market_scan()
    send_all_tables(send_to=["ltrindade", "rzhao"], asset_class=None)
    send_alert_email_am(send_to=send_to)
    send_sharpe_ratio_rank(send_to=send_to)


if __name__ == "__main__":
    update()
